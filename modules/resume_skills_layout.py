"""JD-first skill expansion and measured, fixed-line layout; no provider calls."""
from collections import Counter
import re
import unicodedata


def settings(config):
    enabled = config.get('skills_expand_to_fill', False)
    target = config.get('skills_target_fill_ratio', .97)
    minimum = config.get('skills_min_fill_ratio', .92)
    if type(enabled) is not bool:
        raise ValueError('skills_expand_to_fill must be a boolean')
    if any(type(v) not in (int, float) or not .5 <= v < 1 for v in (target, minimum)) or minimum > target:
        raise ValueError('Skills fill ratios must satisfy 0.5 <= minimum <= target < 1')
    limit = config.get('max_skills', 65)
    if type(limit) is not int or limit < 1:
        raise ValueError('max_skills must be a positive integer')
    return dict(enabled=enabled, target=target, minimum=minimum, priority_limit=limit)


def skill_key(value):
    text = re.sub(r'[^a-z0-9+#.]', '', value.casefold())
    aliases = {'golang':'go', 'largelanguagemodelsllms':'llms',
               'largelanguagemodels':'llms', 'llm':'llms',
               'modelcontextprotocolmcp':'mcp', 'modelcontextprotocol':'mcp',
               'retrievalaugmentedgenerationrag':'rag', 'retrievalaugmentedgeneration':'rag',
               'nodejs':'node.js', 'reactjs':'react', 'react.js':'react'}
    return aliases.get(text, text)


def is_skill(value):
    if not isinstance(value, str) or not value.strip() or len(value) > 100 or '\n' in value:
        return False
    if re.search(r'(?i)\b(?:citizen|citizenship|authorized|authorization|visa|sponsorship|salary|compensation|benefits|united states|eastern time|remote.first|work.location)\b|\d+\+?\s*years?\b', value):
        return False
    if re.search(r'(?i)^(?:(?:senior|staff|principal|lead|founding)\s+)?(?:software\s+|backend\s+|platform\s+|ai\s+)*(?:engineer|developer|architect|manager)\b', value):
        return False
    return bool(re.search(r'[A-Za-z]', value))


FAMILIES = (
    r'backend|api|rest|grpc|microservice|distributed|go(?:lang)?\b|python|java\b|node|concurren|async|performance|latency|cach|message|queue|event.driven|linux|testing|tdd|system.design',
    r'agent|llm|language.model|prompt|retrieval|\brag\b|lang|model.context|\bmcp\b|evaluation|guardrail|context.engineering|tool.call|structured.output',
    r'cloud|\baws\b|azure|\bgcp\b|terraform|kubernetes|docker|infrastructure|ci/cd|observab|monitor|deployment|reliab|sre|devops|on.call|linux|ansible|bash|ec2|lambda|cloudwatch',
    r'data|sql|postgres|mysql|redis|kafka|mongo|snowflake|etl|airflow|spark|vector|storage|schema|cassandra|warehouse|rdbms',
    r'voip|sip|\brtp\b|webrtc|sms|telecom|communication.platform|routing|messaging',
    r'frontend|front.end|react|typescript|javascript|next.js|\bux\b|\bui\b|web.application',
    r'leadership|mentor|technical.strategy|roadmap|architecture|design.review|code.review|cross.team',
    r'security|secure|hipaa|healthcare|rbac|authentication|authorization|oauth|compliance|edi',
)


def candidate_pool(jd, direct, expanded, original, bank, role=''):
    """Keep provenance and priority separate from the direct-JD score denominator."""
    context = jd+' '+role
    families = {i for i, pattern in enumerate(FAMILIES) if re.search(pattern, context, re.I)}
    jd_words = set(re.findall(r'[a-z0-9+#]{3,}', context.casefold()))
    def relevance(term):
        words = set(re.findall(r'[a-z0-9+#]{3,}', term.casefold()))
        overlap = len(words & jd_words)
        related = bool(families & {i for i,p in enumerate(FAMILIES) if re.search(p,term,re.I)})
        return overlap, related
    candidates, seen = [], set()
    def add(values, origin, tier, require_related=False):
        cleaned = []
        for value in values:
            if not is_skill(value):
                continue
            value = re.sub(r'\s+', ' ', value).strip(' ,;:')
            if require_related and not any(relevance(value)):
                continue
            cleaned.append(value)
        if require_related:
            cleaned.sort(key=lambda value: -relevance(value)[0])
        for value in cleaned:
            key = skill_key(value)
            if key and key not in seen:
                seen.add(key)
                candidates.append(dict(term=value, origin=origin, priority=tier))
    add(direct, 'jd_direct', 0)
    add(expanded, 'jd_expanded', 1)
    add(original, 'source_resume', 2, True)
    avoid = {skill_key(v) for v in bank.get('avoid', []) if isinstance(v,str)}
    bank_terms = [row['keyword'] for row in bank.get('keywords', []) if isinstance(row,dict)
                  and isinstance(row.get('keyword'),str) and 'avoid' not in row.get('tier','').casefold()
                  and skill_key(row['keyword']) not in avoid]
    add(bank_terms, 'skill_bank', 3, True)
    return candidates


def pack_rows(candidates, capacities, width, *, target=.97, minimum=.92):
    """Fill each measured row with whole terms, keeping highest priorities early."""
    pending = list(candidates)
    rows = []
    for index, capacity in enumerate(capacities):
        chosen = []
        suffix = ',' if index < len(capacities)-1 else ''
        ceiling = capacity-1.5  # Word/PDF kerning tolerance, without shrinking text.
        while pending:
            fit = next((i for i,item in enumerate(pending)
                        if width(', '.join([*(v['term'] for v in chosen),item['term']])+suffix) <= ceiling), None)
            if fit is None:
                break
            chosen.append(pending.pop(fit))
            if width(', '.join(item['term'] for item in chosen)+suffix) >= capacity*target:
                break
        occupied = width(', '.join(item['term'] for item in chosen)+suffix)
        if not chosen or occupied < capacity*minimum:
            raise ValueError(f'Skills layout cannot fill line {index+1}: {occupied/capacity:.1%}; need {minimum:.0%}')
        rows.append(chosen)
    return rows


def pack_document(doc, source_docx, source_pdf, optimizer, candidates, config):
    import pdfplumber
    from docx import Document
    from modules.resume_layout import paragraph_density
    opts = settings(config)
    source = Document(source_docx)
    start, end = optimizer._find_skills_section(source)
    baselines = [r for r in paragraph_density(source_pdf,source_docx,(start,end)) if r['kind']=='skills']
    if not baselines or any(not r['matched'] for r in baselines):
        raise ValueError('Skills source lines could not be mapped')
    compact = lambda value: re.sub(r'\s+','',unicodedata.normalize('NFKC',value).casefold()).replace('\u00ad','')
    chars = []
    with pdfplumber.open(source_pdf) as pdf:
        for page in pdf.pages:
            chars.extend(page.chars)
    stream, owners = '', []
    for char in chars:
        text = compact(char['text']);stream += text;owners.extend([char]*len(text))
    capacities, measures, groups = [], [], []
    for row in baselines:
        paragraph = source.paragraphs[row['paragraph_index']]
        position = stream.find(compact(paragraph.text))
        if position < 0:
            raise ValueError('Skills source font could not be mapped')
        font = Counter((c.get('fontname'),round(c.get('size',0),2))
            for c in owners[position:position+len(compact(paragraph.text))]).most_common(1)[0][0]
        glyphs = {}
        for char in chars:
            if (char.get('fontname'),round(char.get('size',0),2))==font:
                glyphs.setdefault(char['text'],[]).append(float(char['x1'])-float(char['x0']))
        typical = {c:sorted(values)[len(values)//2] for c,values in glyphs.items()}
        measure = lambda value, table=typical: sum(table.get(c,7.0) for c in value)
        section = source.sections[0]
        right = paragraph.paragraph_format.right_indent or paragraph.style.paragraph_format.right_indent
        capacity = section.page_width.pt-section.right_margin.pt-(right.pt if right else 0)-row['lines'][0]['x0']
        capacities.extend([capacity]*row['line_count'])
        measures.extend([measure]*row['line_count'])
        groups.append((row['paragraph_index'],row['line_count']))
    # The template can contain multiple Skills paragraphs/fonts. Dispatch rows
    # in reading order while carrying the still-unselected ranked pool forward.
    selected, replacements, line_reports = [], {}, []
    pending = list(candidates)
    for paragraph_index, count in groups:
        offset = len(line_reports)
        # This paragraph's style/font is constant in the measured source.
        rows = pack_rows(pending,capacities[offset:offset+count],measures[offset],
                         target=opts['target'],minimum=opts['minimum'])
        text = ',\n'.join(', '.join(item['term'] for item in row) for row in rows)
        optimizer._replace_paragraph_text(doc.paragraphs[paragraph_index],text)
        replacements[paragraph_index] = text
        for row, capacity, measure in zip(rows,capacities[offset:offset+count],measures[offset:offset+count]):
            selected.extend(row)
            line_reports.append(dict(estimated_fill_ratio=round(measure(', '.join(item['term'] for item in row))/capacity,4),usable_width=capacity))
        used = {skill_key(item['term']) for row in rows for item in row}
        pending = [item for item in pending if skill_key(item['term']) not in used]
    return dict(method='jd-first-source-bank-measured-width-v1',source_line_count=len(capacities),
                target_fill_ratio=opts['target'],minimum_fill_ratio=opts['minimum'],
                selected=selected,output_skill_count=len(selected),priority_limit=opts['priority_limit'],
                expanded_beyond_max_skills=len(selected)>opts['priority_limit'],lines=line_reports), replacements


def verify_fill(density, layout):
    rows = [r for r in density['output'] if r['kind']=='skills']
    lines = [line for row in rows for line in row.get('lines',[])]
    widths = [line['usable_width'] for line in layout['lines']]
    ratios = [line['width']/width for line,width in zip(lines,widths)]
    issues = []
    if len(lines) != layout['source_line_count']:
        issues.append('Skills source line count changed')
    for index, ratio in enumerate(ratios):
        if ratio < layout['minimum_fill_ratio'] or ratio > 1.002:
            issues.append(f'Skills line {index+1} width fill {ratio:.1%} outside bounds')
    return dict(passed=not issues,line_count=len(lines),fill_ratios=ratios,issues=issues)
