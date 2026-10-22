"""Source-unconstrained test mode, isolated from actual application eligibility."""
import json
import socket
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import Mock, patch

from docx import Document
from docx.shared import Pt
from modules import resume_all_bullets as m
from modules import resume_jd_experiment as experiment
from modules.ai.resume_optimizer import ResumeOptimizer, run_resume_optimization
from test_resume_all_bullets import inv, response
import optimize_resume as cli


class SyntheticJDTest(unittest.TestCase):
    def test_payload_has_no_source_content_and_no_source_keyword_restrictions(self):
        inventory=inv(3,1); group=inventory['groups'][0]
        targets=experiment.alignment('Build Rust and AG-UI systems',['Rust','AG-UI'])
        payload=m.compact_payload(inventory,'Build Rust and AG-UI systems','Engineer',experiment.POLICY,
            group,group['bullets'],alignment=targets)
        encoded=json.dumps(payload)
        for source in inventory['bullets']:
            self.assertNotIn(source['source_text'],encoded)
        self.assertNotIn('Source company',encoded)
        self.assertNotIn('source_experience',payload)
        self.assertTrue(payload['fact_policy']['invented_claims_allowed'])
        self.assertEqual(payload['excluded_jd_terms'],[])
        for slot in payload['requested_bullets']:
            self.assertNotIn('source_text',slot)
            self.assertNotIn('protected_source_terms',slot)
            self.assertEqual(slot['fact_anchors'],[])
            self.assertIn('Rust',slot['jd_keyword_opportunities'])

    def test_only_synthetic_mode_accepts_fictional_new_claims(self):
        inventory=inv(1,1);group=inventory['groups'][0]
        raw=response({'experience_id':group['experience_id'],'requested_bullets':group['bullets']})
        raw['results'][0].update(rewritten_text='Built Rust and AG-UI agents, reducing latency by 99%.',preserved_fact_anchors=[])
        for policy in ('grounded','review_draft'):
            self.assertEqual(m.validate_response(raw,group,group['bullets'],content_policy=policy)['0']['status'],'blocked')
        with patch.object(m,'source_claim_issue',side_effect=AssertionError('source content gate called')), \
             patch('modules.ai.resume_optimizer._source_fact_issue',side_effect=AssertionError('source facts called')):
            result=m.validate_response(raw,group,group['bullets'],forbidden_keywords=['Rust'],content_policy=experiment.POLICY)['0']
        self.assertEqual(result['status'],'candidate')
        self.assertEqual(result['rewritten_text'],raw['results'][0]['rewritten_text'])

    def test_protocol_and_measured_fit_still_apply(self):
        inventory=inv(1,1);group=inventory['groups'][0]
        raw=response({'experience_id':group['experience_id'],'requested_bullets':group['bullets']})
        raw['results'][0]['source_text_hash']='wrong'
        self.assertEqual(m.validate_response(raw,group,group['bullets'],content_policy=experiment.POLICY)['0']['reason'],'source hash mismatch')
        raw['results'][0].update(source_text_hash=group['bullets'][0]['source_text_hash'],
            rewritten_text='Built Rust agents.',alternative_texts=['Implemented AG-UI services.'])
        out=m.validate_response(raw,group,group['bullets'],content_policy=experiment.POLICY,
            fit_check=lambda _: 'too wide')['0']
        self.assertEqual(out['status'],'blocked')
        self.assertEqual(out['reason'],'too wide')

    def test_schema_and_actual_provider_use_the_synthetic_override(self):
        inventory=inv(1,1);group=inventory['groups'][0]
        payload=m.compact_payload(inventory,'Rust agents','Engineer',experiment.POLICY,group,group['bullets'])
        schema=m.provider_response_format('bullets',payload)['json_schema']['schema']
        self.assertIn('anyOf',schema['properties']['results']['items'])
        optimizer=Mock()
        def exercise(*args,provider,**kwargs):provider('bullets',payload,7)
        with patch.object(m,'run_all_bullets',side_effect=exercise):
            run_resume_optimization('Rust agents','source.docx','test','Engineer',
                rewrite_all_bullets=True,content_policy=experiment.POLICY,optimizer_instance=optimizer,
                runtime_config_override={'optimization_mode':'full','rewrite_all_bullets':True})
        actual=optimizer.client.with_options.return_value.chat.completions.create.call_args.kwargs['messages'][0]['content']
        self.assertEqual(actual,m.SYSTEM+experiment.PROMPT_OVERRIDE)
        self.assertEqual(m.generation_system('grounded'),m.SYSTEM)
        self.assertEqual(m.generation_system('review_draft'),m.SYSTEM)

    def test_cli_requires_explicit_all_bullets_and_excludes_other_policy(self):
        base=['--resume','source.docx','--jd-text','Rust','--output','new']
        args=cli.parser().parse_args(base+['--synthetic-jd-test','--rewrite-all-bullets'])
        self.assertTrue(args.synthetic_jd_test)
        self.assertFalse(args.review_draft)
        with self.assertRaises(SystemExit):
            cli.parser().parse_args(base+['--synthetic-jd-test','--review-draft'])

    def test_real_render_free_rewrites_skills_and_blocks_attachment(self):
        with TemporaryDirectory() as tmp:
            root=Path(tmp);source=root/'original.docx';doc=Document()
            doc.styles['Normal'].font.name='Arial';doc.styles['Normal'].font.size=Pt(10)
            doc.add_heading('SKILLS',level=1);doc.add_paragraph('Python, AWS, React')
            doc.add_heading('PROFESSIONAL EXPERIENCE',level=1);doc.add_paragraph('Original company')
            doc.add_paragraph('Engineer\tJanuary 2020 - Present')
            doc.add_paragraph('Built Python APIs with no public ingress for 20 customers',style='List Bullet')
            doc.add_heading('Academic Distinctions:',level=1)
            doc.add_paragraph('Earned a university award with a 3.9 GPA',style='List Bullet')
            doc.add_heading('PROJECTS',level=1);doc.add_paragraph('Original project')
            doc.add_paragraph('Co-founder planning a launch within 2 months',style='List Bullet')
            doc.save(source);before=source.read_bytes()
            optimizer=ResumeOptimizer.__new__(ResumeOptimizer)
            self.assertTrue(optimizer.convert_docx_to_pdf(str(source),str(source.with_suffix('.pdf'))))
            def provider(stage,payload,timeout):
                if stage=='keywords':return {'keywords':['Rust','AG-UI','MCP']}
                return {'experience_id':payload['experience_id'],'results':[
                    {'bullet_id':b['bullet_id'],'source_text_hash':b['source_text_hash'],'status':'rewritten',
                     'rewritten_text':'Built Rust agents with AG-UI and MCP.' if b['ordinal']==0 else 'Shipped MCP tools through Rust services.',
                     'matched_jd_spans':[],'preserved_fact_anchors':[],'change_reason':'Synthetic JD example','blocked_reason':None}
                    for b in payload['requested_bullets']]}
            config=dict(m.DEFAULTS,rewrite_all_bullets=True,expected_page_count=1,
                rewrite_requests_per_minute=1000000,rewrite_tokens_per_minute=1000000000)
            with patch.object(socket.socket,'connect',side_effect=AssertionError('network forbidden')):
                result=m.run_all_bullets('Required: Rust, AG-UI, MCP.',source,'synthetic-test','Staff Engineer',config,
                    policy=experiment.POLICY,execution_kind='offline_fixture',provider=provider,output_root=root/'review')
            report=json.loads((root/'review/match-report.json').read_text())
            self.assertEqual(report['status'],'complete',str(report.get('errors'))+str(report.get('format_verification',{}).get('checks')))
            self.assertEqual(report['coverage']['accepted'],2)
            self.assertEqual(report['preserved_sections']['Academic Distinctions']['bullet_count'],1)
            self.assertTrue(report['format_verification']['checks']['academic_distinctions_preserved'])
            self.assertIn('Earned a university award with a 3.9 GPA',
                '\n'.join(p.text for p in Document(root/'review/optimized.docx').paragraphs))
            self.assertFalse(report['source_content_checks_enabled'])
            self.assertEqual(report['skills_optimization']['output'],['Rust','AG-UI','MCP'])
            self.assertFalse(report['strict_attachment_eligible']);self.assertIsNone(result)
            self.assertFalse(m.verify_strict_attachment(root/'review/optimized.pdf'))
            self.assertEqual(source.read_bytes(),before)
            self.assertIn(experiment.LABEL,(root/'review/comparison.html').read_text())
            m.verify_binding(root/'review')


if __name__=='__main__':unittest.main()
