import unittest
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
import tempfile
import shutil
from copy import deepcopy
from unittest.mock import patch, Mock
import importlib
import sys
import subprocess
import json
from types import SimpleNamespace

from modules.ai.resume_optimizer import (
    ResumeOptimizer,
    build_guided_keywords,
    calculate_keyword_score,
    infer_role_profile,
    load_bullet_bank,
    load_project_bullet_banks,
    load_skill_bank,
    select_bullet_bank_bullets,
)


class ResumeOptimizerHelpersTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.bank = load_skill_bank()

    def test_skill_bank_and_profiles_are_loaded(self):
        self.assertGreater(len(self.bank["keywords"]), 100)
        self.assertIn("Staff / Principal AI Platform", self.bank["profiles"])
        self.assertIn("Healthcare AI", self.bank["profiles"])

    def test_profile_inference_and_acronym_guidance(self):
        profile = infer_role_profile(
            "Staff AI Platform Engineer",
            "Build MCP services, Kubernetes infrastructure, evals and observability",
            self.bank["profiles"],
        )
        self.assertEqual(profile, "Staff / Principal AI Platform")
        keywords = build_guided_keywords(
            "Model Context Protocol (MCP), Kubernetes and Terraform",
            ["MCP", "Kubernetes"],
            "Staff AI Platform Engineer",
            ["Python"],
            self.bank,
            profile,
        )
        self.assertIn("Model Context Protocol (MCP)", keywords)

    def test_score_is_deterministic_and_acronym_aware(self):
        report = calculate_keyword_score(
            "Built MCP services with Kubernetes",
            ["Model Context Protocol (MCP)", "Kubernetes", "Terraform"],
            self.bank,
        )
        self.assertEqual(report["matched"], ["Model Context Protocol (MCP)", "Kubernetes"])
        self.assertEqual(report["missing"], ["Terraform"])

    def test_optimizer_preserves_skill_section_density(self):
        optimizer = ResumeOptimizer.__new__(ResumeOptimizer)
        result = optimizer.optimize_skills(
            ["Python", "Legacy Tool", "React"],
            ["React", "Kubernetes"],
            "Kubernetes React",
        )
        self.assertEqual(len(result), 3)
        self.assertEqual(result[0], "React")
        self.assertIn("Kubernetes", result)

    def test_skill_acronym_aliases_are_deduplicated_with_full_jd_display(self):
        optimizer = ResumeOptimizer.__new__(ResumeOptimizer)
        result = optimizer.optimize_skills(
            ["MCP", "Model Context Protocol (MCP)", "RAG", "Retrieval-Augmented Generation (RAG)", "Python"],
            ["MCP", "RAG", "Model Context Protocol (MCP)", "Retrieval-Augmented Generation (RAG)", "Python"],
            "MCP and RAG Python",
        )
        self.assertEqual(result, ["Model Context Protocol (MCP)", "Retrieval-Augmented Generation (RAG)", "Python"])

    def test_proven_plural_acronym_expansion_deduplicates_singular(self):
        optimizer = ResumeOptimizer.__new__(ResumeOptimizer)
        result = optimizer.optimize_skills(["LLM", "Python"], ["Large Language Models (LLMs)", "LLM", "Python"], "LLM Python")
        self.assertEqual(result, ["Large Language Models (LLMs)", "Python"])

    def test_parenthetical_context_is_not_treated_as_an_acronym_expansion(self):
        optimizer = ResumeOptimizer.__new__(ResumeOptimizer)
        result = optimizer.optimize_skills(["Infrastructure (AWS)", "AWS", "Python"], ["AWS", "Infrastructure (AWS)", "Python"], "AWS")
        self.assertIn("AWS", result)
        self.assertIn("Infrastructure (AWS)", result)
        self.assertEqual(len(result), 3)

    def test_bullet_rewrite_receives_shared_jd_context(self):
        optimizer = ResumeOptimizer.__new__(ResumeOptimizer)
        message = optimizer._generate_user_message(
            [], "Principal Engineer", "Healthcare agent platform JD", ["Python", "RAG"], ["agent runtime", "evals"]
        )
        self.assertIn('"Python"', message)
        self.assertIn('"agent runtime"', message)


def synthetic_resume(include_projects=False, headings=None):
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10)
    doc.styles["List Paragraph"].font.size = Pt(10)
    for name in headings or ["SKILLS", "EDUCATION", "PROFESSIONAL EXPERIENCE"]:
        doc.add_heading(name, level=1)
        if name == "SKILLS":
            paragraph = doc.add_paragraph("Python, Legacy Tool, AWS")
            paragraph.runs[0].font.size = Pt(10)
        elif name == "PROFESSIONAL EXPERIENCE":
            for company, count in [("TREK HEALTH", 8), ("OTHER COMPANY", 2), ("RIVIGO", 3)]:
                doc.add_paragraph(company + "\tRemote")
                paragraph = doc.add_paragraph("Software Engineer\tJanuary 2020 - Present")
                paragraph.runs[0].italic = True
                paragraph.runs[0].font.size = Pt(10)
                for index in range(count):
                    paragraph = doc.add_paragraph(f"Built Python service {index + 1} with reliable production APIs", style="List Paragraph")
                    paragraph.runs[0].font.size = Pt(10)
        else:
            doc.add_paragraph("Example university")
    if include_projects:
        doc.add_heading("ENTREPRENEURIAL VENTURES & PROJECTS", level=1)
        for company, count in [("CODEX SELF-HOSTED", 5), ("SOHAMHAMSO", 3), ("SPONSORSKIP", 2), ("ARICEUM", 2)]:
            doc.add_paragraph(company + "\tRemote")
            doc.add_paragraph("Founder\tJanuary 2020 - Present")
            for index in range(count):
                paragraph = doc.add_paragraph(f"Built project service {index + 1}", style="List Paragraph")
                paragraph.runs[0].font.size = Pt(10)
        doc.add_heading("ACADEMIC DISTINCTIONS", level=1)
        doc.add_paragraph("Earned an academic award", style="List Paragraph")
    return doc


def write_test_pdf(path, lines=None, pages=1, extra_positions=None):
    writer = canvas.Canvas(str(path), pagesize=letter)
    for page in range(pages):
        writer.setFont("Helvetica", 10)
        for index, text in enumerate(lines or ["SKILLS", "Python AWS", "EDUCATION", "Example university", "PROFESSIONAL EXPERIENCE", "Built Python services"]):
            writer.drawString(72, 730 - index * 20, text)
        if extra_positions:
            for text, x, y in extra_positions:
                writer.drawString(x, y, text)
        writer.showPage()
    writer.save()


class ResumeFormatVerifierTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source_docx = self.root / "source.docx"
        self.output_docx = self.root / "output.docx"
        self.source_pdf = self.root / "source.pdf"
        self.output_pdf = self.root / "output.pdf"
        synthetic_resume().save(self.source_docx)
        shutil.copyfile(self.source_docx, self.output_docx)
        self.lines = [paragraph.text.replace("\t", " ") for paragraph in Document(self.source_docx).paragraphs if paragraph.text]
        write_test_pdf(self.source_pdf, self.lines)
        shutil.copyfile(self.source_pdf, self.output_pdf)
        self.optimizer = ResumeOptimizer.__new__(ResumeOptimizer)

    def verify(self, expected_pages=1):
        return self.optimizer.verify_format(str(self.source_docx), str(self.output_docx), str(self.source_pdf), str(self.output_pdf), expected_pages)

    def test_source_relative_order_passes_even_when_different_from_old_global_order(self):
        self.assertTrue(self.verify()["passed"])

    def test_wrong_page_count_rejected(self):
        write_test_pdf(self.output_pdf, self.lines, pages=2)
        self.assertFalse(self.verify()["checks"]["page_count"])
        self.assertFalse(self.verify()["passed"])

    def test_text_outside_each_page_edge_is_rejected(self):
        cases = [("left", -30, 400, "no_horizontal_clipping"), ("right", 610, 400, "no_horizontal_clipping"),
                 ("top", 72, 800, "no_vertical_clipping"), ("bottom", 72, -10, "no_vertical_clipping")]
        for name, x, y, check in cases:
            with self.subTest(edge=name):
                write_test_pdf(self.output_pdf, self.lines, extra_positions=[("clipped text", x, y)])
                report = self.verify()
                self.assertFalse(report["checks"][check], report)
                self.assertFalse(report["passed"])

    def test_section_reordering_and_missing_section_are_rejected(self):
        for headings in (["EDUCATION", "SKILLS", "PROFESSIONAL EXPERIENCE"], ["SKILLS", "PROFESSIONAL EXPERIENCE"]):
            with self.subTest(headings=headings):
                synthetic_resume(headings=headings).save(self.output_docx)
                self.assertFalse(self.verify()["passed"])

    def test_missing_rendered_heading_rejected(self):
        write_test_pdf(self.output_pdf, ["SKILLS", "Python", "PROFESSIONAL EXPERIENCE", "Built services"])
        report = self.verify()
        self.assertFalse(report["checks"]["rendered_headings_present"])

    def test_rendered_section_order_must_match_docx_source_order(self):
        source_lines = list(self.lines)
        reordered = ["EDUCATION", "Example university", "SKILLS", "Python, Legacy Tool, AWS"] + source_lines[4:]
        write_test_pdf(self.output_pdf, reordered)
        report = self.verify()
        self.assertTrue(report["checks"]["docx_pdf_content_consistent"])
        self.assertFalse(report["checks"]["rendered_section_order_preserved"])

    def test_margin_change_rejected(self):
        doc = Document(self.output_docx)
        doc.sections[0].left_margin = Inches(0.1)
        doc.save(self.output_docx)
        self.assertFalse(self.verify()["checks"]["margins_similar"])

    def test_font_maximum_change_rejected_even_with_unchanged_minimum(self):
        doc = Document(self.output_docx)
        bullet = next(paragraph for paragraph in doc.paragraphs if paragraph.style.name == "List Paragraph")
        bullet.runs[0].font.size = Pt(14)
        doc.save(self.output_docx)
        self.assertFalse(self.verify()["checks"]["font_sizes_similar"])

    def test_new_orphaned_heading_rejected(self):
        write_test_pdf(self.output_pdf, ["SKILLS", "Python", "EDUCATION", "Example university", "PROFESSIONAL EXPERIENCE"])
        self.assertFalse(self.verify()["checks"]["no_new_orphaned_headings"])

    def test_existing_source_orphan_is_not_falsely_rejected(self):
        doc = Document()
        for text in ["SKILLS", "Python", "EDUCATION", "Example university", "PROFESSIONAL EXPERIENCE"]:
            if text in {"SKILLS", "EDUCATION", "PROFESSIONAL EXPERIENCE"}:
                doc.add_heading(text, level=1)
            else:
                doc.add_paragraph(text)
        doc.save(self.source_docx)
        shutil.copyfile(self.source_docx, self.output_docx)
        write_test_pdf(self.source_pdf, ["SKILLS", "Python", "EDUCATION", "Example university", "PROFESSIONAL EXPERIENCE"])
        shutil.copyfile(self.source_pdf, self.output_pdf)
        self.assertTrue(self.verify()["passed"])

    def test_extreme_body_font_does_not_disappear_from_signature(self):
        doc = Document(self.output_docx)
        bullet = next(paragraph for paragraph in doc.paragraphs if paragraph.style.name == "List Paragraph")
        bullet.runs[0].font.size = Pt(20)
        doc.save(self.output_docx)
        self.assertFalse(self.verify()["checks"]["font_sizes_similar"])

    def test_unrelated_pdf_content_is_rejected_even_when_source_pdf_matches_it(self):
        write_test_pdf(self.output_pdf, ["Unrelated document with text"])
        shutil.copyfile(self.output_pdf, self.source_pdf)
        report = self.verify()
        self.assertFalse(report["checks"]["docx_pdf_content_consistent"])
        self.assertFalse(report["checks"]["rendered_headings_present"])
        self.assertFalse(report["passed"])

    def test_invalid_pdf_rejected(self):
        self.output_pdf.write_text("invalid PDF")
        report = self.verify()
        self.assertFalse(report["passed"])
        self.assertIn("verifier_error", report["checks"])

    def test_wrapping_is_recorded_as_evidence_without_false_rejection(self):
        doc = Document()
        paragraph = doc.add_paragraph("Built Python services", style="List Paragraph")
        paragraph.runs[0].font.size = Pt(10)
        doc.save(self.source_docx)
        shutil.copyfile(self.source_docx, self.output_docx)
        write_test_pdf(self.source_pdf, ["Built Python services"])
        write_test_pdf(self.output_pdf, ["Built Python", "services"])
        report = self.verify()
        self.assertTrue(report["passed"], report)
        self.assertEqual(report["details"]["optimized_layout"]["max_bullet_lines"], 2)
        self.assertTrue(report["warnings"])


class ResumeRewritingRegressionTest(unittest.TestCase):
    def setUp(self):
        self.optimizer = ResumeOptimizer.__new__(ResumeOptimizer)

    def test_extraction_stops_before_projects_and_awards(self):
        doc = synthetic_resume(include_projects=True)
        entries = self.optimizer._extract_experience_bullets(doc)
        self.assertEqual([entry["company"] for entry in entries], ["TREK HEALTH", "OTHER COMPANY", "RIVIGO"])
        self.assertEqual(len(entries[-1]["bullets"]), 3)
        serialized = json.dumps(entries)
        self.assertNotIn("ENTREPRENEURIAL", serialized)
        self.assertNotIn("ACADEMIC", serialized)
        self.assertNotIn("Earned an academic award", serialized)

    def test_llm_update_does_not_rewrite_section_headings_or_project_awards(self):
        doc = synthetic_resume(include_projects=True)
        entries = self.optimizer._extract_experience_bullets(doc)
        rewritten = deepcopy(entries)
        for entry in rewritten:
            entry["bullets"] = ["Updated professional bullet"] * len(entry["bullets"])
        original_headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
        self.optimizer._update_document_with_rewritten_bullets(doc, entries, rewritten)
        self.assertEqual(original_headings, [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")])
        self.assertIn("Earned an academic award", [p.text for p in doc.paragraphs])
        self.assertEqual(next(p for p in doc.paragraphs if p.text == "Updated professional bullet").runs[0].font.size.pt, 10)

    def test_bank_replacement_retains_font_and_numbering_when_adding_bullets(self):
        doc = synthetic_resume()
        for original in doc.paragraphs:
            if original.style.name == "List Paragraph":
                original.paragraph_format.left_indent = Inches(0.2)
                original.runs[0].font.name = "Arial"
        self.assertTrue(self.optimizer._replace_first_company_bullets(doc, "Trek Health", ["Replacement"] * 9))
        replacements = [p for p in doc.paragraphs if p.text == "Replacement"]
        self.assertEqual(len(replacements), 9)
        for paragraph in replacements:
            self.assertEqual(paragraph.style.name, "List Paragraph")
            self.assertEqual(paragraph.runs[0].font.size.pt, 10)
        self.assertEqual(replacements[-1].runs[0].font.name, "Arial")

    def test_text_replacement_keeps_all_source_run_spacing_and_fonts(self):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        doc = Document()
        paragraph = doc.add_paragraph()
        first = paragraph.add_run("Lead ")
        first.font.size = Pt(10)
        tail = paragraph.add_run("compressed source text")
        tail.font.size = Pt(10)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:val"), "-2")
        tail._r.get_or_add_rPr().append(spacing)
        original_properties = [run._r.rPr.xml for run in paragraph.runs]
        self.optimizer._replace_paragraph_text(paragraph, "New compressed replacement")
        self.assertEqual(paragraph.text, "New compressed replacement")
        self.assertEqual([run._r.rPr.xml for run in paragraph.runs], original_properties)
        before = paragraph._p.xml
        self.optimizer._replace_paragraph_text(paragraph, paragraph.text)
        self.assertEqual(paragraph._p.xml, before)

    def test_text_replacement_updates_hyperlink_runs_without_stale_text(self):
        from docx.oxml import OxmlElement
        doc = Document()
        paragraph = doc.add_paragraph("Old lead ")
        hyperlink = OxmlElement("w:hyperlink")
        linked_run = OxmlElement("w:r")
        linked_text = OxmlElement("w:t")
        linked_text.text = "old linked words"
        linked_run.append(linked_text)
        hyperlink.append(linked_run)
        paragraph._p.append(hyperlink)
        self.optimizer._replace_paragraph_text(paragraph, "New sentence")
        self.assertEqual(paragraph.text, "New sentence")
        self.assertNotIn("old linked words", paragraph._p.xml)

    def test_full_title_replacement_stops_before_projects(self):
        doc = synthetic_resume(include_projects=True)
        self.optimizer.replace_job_titles(doc, "Principal Engineer")
        self.assertTrue(any(p.text.startswith("Principal Engineer\t") for p in doc.paragraphs))
        self.assertEqual(sum(p.text.startswith("Founder\t") for p in doc.paragraphs), 4)

    def test_title_replacement_supports_each_recognized_experience_heading(self):
        for heading in ("PROFESSIONAL EXPERIENCE", "WORK EXPERIENCE", "EXPERIENCE"):
            with self.subTest(heading=heading):
                doc = Document()
                doc.add_heading(heading, level=1)
                doc.add_paragraph("EXAMPLE COMPANY\tRemote")
                doc.add_paragraph("Software Engineer\tJanuary 2020 - Present")
                self.optimizer.replace_job_titles(doc, "Principal Engineer")
                self.assertTrue(any(paragraph.text.startswith("Principal Engineer\t") for paragraph in doc.paragraphs))

    def test_empty_explicit_project_bank_list_disables_banks(self):
        self.assertEqual(load_project_bullet_banks([]), [])
        doc = synthetic_resume(include_projects=True)
        self.optimizer._call_openai_to_rewrite_bullets = lambda entries, *args: entries
        self.optimizer.rewrite_experience_bullets(doc, "Engineer", "Python", [], bullet_bank=load_bullet_bank(), project_bullet_banks=[])
        self.assertEqual(self.optimizer.last_bullet_bank_selections, {})

    def test_import_performs_no_cloud_key_lookup(self):
        import modules.ai.resume_optimizer as module
        with patch("modules.ai.openai_key_loader.load_openai_key") as key_loader, patch("requests.get") as request:
            importlib.reload(module)
        key_loader.assert_not_called()
        request.assert_not_called()

class ResumeApiBoundaryTest(unittest.TestCase):
    def setUp(self):
        self.optimizer = ResumeOptimizer.__new__(ResumeOptimizer)
        self.optimizer.client = Mock()
        self.entries = [{"company": "Example", "job_title": "Engineer", "bullets": ["Built reliable Python APIs"]}]

    def response(self, content):
        return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=content))])

    def test_keyword_response_rejects_wrong_json_shapes(self):
        for value in ({"keywords": "Python"}, {"keywords": ["Python", 3]}, {"keywords": [""]}, ["Python"]):
            with self.subTest(value=value):
                self.optimizer.client.chat.completions.create.return_value = self.response(json.dumps(value))
                self.assertEqual(self.optimizer.extract_keywords("Python"), [])
                self.assertTrue(self.optimizer.last_keyword_error)

    def test_keyword_response_normalizes_and_deduplicates_strings(self):
        self.optimizer.client.chat.completions.create.return_value = self.response('{"keywords": [" Python ", "Python", "AWS"]}')
        self.assertEqual(self.optimizer.extract_keywords("Python AWS"), ["Python", "AWS"])

    def test_invalid_rewrites_preserve_all_original_entries(self):
        cases = [
            {"rewritten_experience": "invalid"},
            {"rewritten_experience": []},
            {"rewritten_experience": [{"company": "Other", "job_title": "Engineer", "bullets": ["Built APIs"]}]},
            {"rewritten_experience": [{"company": "Example", "job_title": "Engineer", "bullets": ["", "extra"]}]},
            {"rewritten_experience": [{"company": "Example", "job_title": "Engineer", "bullets": [42]}]},
            {"rewritten_experience": [{"company": "Example", "job_title": "Engineer", "bullets": ["much longer new bullet " * 20]}]},
        ]
        for value in cases:
            with self.subTest(value=value):
                result = self.optimizer._parse_openai_response(self.response(json.dumps(value)), self.entries)
                self.assertEqual(result, self.entries)
                self.assertTrue(self.optimizer.last_bullet_rewrite_error)

    def test_malformed_json_response_preserves_source(self):
        result = self.optimizer._parse_openai_response(self.response("{malformed"), self.entries)
        self.assertEqual(result, self.entries)
        self.assertTrue(self.optimizer.last_bullet_rewrite_error)

    def test_valid_rewrite_preserves_entry_identity_and_formatting_contract(self):
        result = self.optimizer._parse_openai_response(self.response('```json\n{"rewritten_experience": [{"company": "Example", "job_title": "Engineer", "bullets": ["Built Python APIs"]}]}\n```'), self.entries)
        self.assertEqual(result[0]["bullets"], ["Built Python APIs"])
        self.assertEqual(result[0]["company"], "Example")

    def test_rewrite_api_exception_preserves_source(self):
        self.optimizer.client.chat.completions.create.side_effect = RuntimeError("API unavailable")
        result = self.optimizer._call_openai_to_rewrite_bullets(self.entries, "Engineer", "Python")
        self.assertEqual(result, self.entries)
        self.assertIn("API unavailable", self.optimizer.last_bullet_rewrite_error)


class BankResourceResolutionTest(unittest.TestCase):
    def test_relative_banks_resolve_independently_of_working_directory(self):
        previous = Path.cwd()
        with tempfile.TemporaryDirectory() as directory:
            try:
                import os
                os.chdir(directory)
                self.assertGreater(len(load_skill_bank("config/skills/Skills_Keyword_Bank_Staff_Principal_2026-10.md")["keywords"]), 100)
                self.assertGreater(len(load_bullet_bank("config/bullets/Trek_Health_Bullet_Bank.md")["bullets"]), 0)
                self.assertEqual(len(load_project_bullet_banks()), 4)
            finally:
                os.chdir(previous)

    def test_bundled_bank_defaults_and_relative_paths_use_resource_directory(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            resources = root / "JobRobo.app" / "Contents" / "Resources"
            bank_folder = resources / "config" / "bullets"
            bank_folder.mkdir(parents=True)
            bank_folder.joinpath("Trek_Health_Bullet_Bank.md").write_text("# Custom Bullet Bank\n")
            skill_folder = resources / "config" / "skills"
            skill_folder.mkdir()
            skill_folder.joinpath("Skills_Keyword_Bank_Staff_Principal_2026-10.md").write_text("# Skills Bank\n")
            executable = root / "JobRobo.app" / "Contents" / "MacOS" / "JobRobo"
            with patch.object(sys, "frozen", True, create=True), patch.object(sys, "executable", str(executable)):
                self.assertEqual(Path(load_bullet_bank()["path"]), bank_folder / "Trek_Health_Bullet_Bank.md")
                self.assertEqual(Path(load_bullet_bank("config/bullets/Trek_Health_Bullet_Bank.md")["path"]), bank_folder / "Trek_Health_Bullet_Bank.md")
                self.assertEqual(Path(load_skill_bank()["path"]), skill_folder / "Skills_Keyword_Bank_Staff_Principal_2026-10.md")

class TrekBulletBankTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.bank = load_bullet_bank()

    def test_trek_bank_selects_safe_jd_relevant_bullets(self):
        selected = select_bullet_bank_bullets(
            "Principal Engineer",
            "Healthcare contract intelligence, agent platform, evals, security, Snowflake and price transparency",
            self.bank,
        )
        self.assertEqual(self.bank["company"], "Trek Health")
        self.assertEqual(len(selected), 7)
        self.assertIn("A", {bullet["section"] for bullet in selected})
        self.assertIn("N", {bullet["section"] for bullet in selected})
        self.assertTrue(all(110 <= bullet["length"] <= 126 for bullet in selected))
        metric_ids = [metric for bullet in selected for metric in bullet["metrics"]]
        self.assertEqual(len(metric_ids), len(set(metric_ids)))
        self.assertTrue(all(not bullet["unverified"] for bullet in selected))


class ProjectBulletBankTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.root = Path(__file__).resolve().parents[1]
        cls.specs = load_project_bullet_banks()

    def test_all_project_banks_parse_with_their_rules(self):
        self.assertEqual(
            [spec["company"] for spec in self.specs],
            ["Trek Health", "CODEX SELF-HOSTED", "SOHAMHAMSO", "SPONSORSKIP"],
        )
        for spec in self.specs:
            bank = spec["bank"]
            self.assertGreater(len(bank["bullets"]), 0)
            selected = select_bullet_bank_bullets(
                "Staff AI Platform Engineer",
                "LLM inference, agents, evaluation, search, cloud infrastructure and security",
                bank,
                spec["max_bullets"],
                spec["allow_unverified"],
                spec["required_sections"],
                spec["min_bullets"],
            )
            self.assertGreaterEqual(len(selected), spec["min_bullets"], spec["company"])
            self.assertLessEqual(len(selected), spec["max_bullets"], spec["company"])
            self.assertTrue(set(spec["required_sections"]).issubset({item["section"] for item in selected}))
            self.assertTrue(all(108 <= item["length"] <= 126 for item in selected), spec["company"])
            self.assertTrue(all(not item["unverified"] for item in selected), spec["company"])
            metrics = [metric for item in selected for metric in item["metrics"]]
            self.assertEqual(len(metrics), len(set(metrics)), spec["company"])

    def test_resume_projects_receive_exact_bank_bullets(self):
        doc = synthetic_resume(include_projects=True)
        optimizer = ResumeOptimizer.__new__(ResumeOptimizer)
        optimizer._call_openai_to_rewrite_bullets = lambda entries, title, description, skills, keywords: entries
        optimizer.rewrite_experience_bullets(
            doc,
            "Staff AI Platform Engineer",
            "LLM inference, agents, evaluation, search and cloud infrastructure",
            ["Python", "AWS"],
            project_bullet_banks=self.specs,
        )
        self.assertEqual(
            set(optimizer.last_bullet_bank_selections),
            {"Trek Health", "CODEX SELF-HOSTED", "SOHAMHAMSO", "SPONSORSKIP"},
        )
        for company, selected in optimizer.last_bullet_bank_selections.items():
            paragraphs = [
                paragraph.text
                for paragraph in doc.paragraphs
                if paragraph.style.name == "List Paragraph"
            ]
            for item in selected:
                self.assertIn(item["text"], paragraphs, company)


if __name__ == "__main__":
    unittest.main()
