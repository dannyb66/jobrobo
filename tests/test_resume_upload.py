import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from modules.resume_upload import ResumeUploadError, upload_verified_resume, verify_source_resume


class VerifiedResumeUploadTest(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.source = Path(self.directory.name) / "resume.docx"
        self.source.write_bytes(b"source docx")
        self.pdf = self.source.with_suffix(".pdf")
        self.pdf.write_bytes(b"source pdf")
        self.optimized = Path(self.directory.name) / "optimized.pdf"
        self.optimized.write_bytes(b"verified output pdf")
        self.file_input = Mock()
        self.modal = Mock()
        self.modal.find_element.return_value = self.file_input
        self.verify = Mock(return_value=True)

    def upload(self, optimize):
        return upload_verified_resume(
            self.modal, str(self.source), "JD", "123", "Staff Engineer",
            optimize=optimize, verify_source=self.verify,
        )

    def test_verified_output_attached_with_job_context(self):
        optimize = Mock(return_value=str(self.optimized))
        self.assertEqual(self.upload(optimize), (True, "optimized.pdf"))
        self.file_input.send_keys.assert_called_once_with(str(self.optimized.resolve()))
        self.assertEqual(optimize.call_args.kwargs["job_title"], "Staff Engineer")
        self.verify.assert_not_called()

    def test_drafting_receives_exact_successfully_attached_pdf_path(self):
        attached = Mock()
        result = upload_verified_resume(
            self.modal, str(self.source), 'JD', '123',
            optimize=Mock(return_value=str(self.optimized)), verify_source=self.verify,
            on_attached=attached)
        self.assertTrue(result[0])
        attached.assert_called_once_with(str(self.optimized.resolve()))

    def test_failed_optimized_upload_reports_only_source_path_after_success(self):
        attached = Mock()
        self.file_input.send_keys.side_effect = [RuntimeError('upload failed'), None]
        result = upload_verified_resume(
            self.modal, str(self.source), 'JD', '123',
            optimize=Mock(return_value=str(self.optimized)), verify_source=self.verify,
            on_attached=attached)
        self.assertTrue(result[0])
        attached.assert_called_once_with(str(self.pdf.resolve()))

    def test_rejection_attaches_independently_verified_source(self):
        self.assertEqual(self.upload(Mock(return_value=None)), (True, "resume.pdf"))
        self.verify.assert_called_once_with(self.source, self.pdf)
        self.file_input.send_keys.assert_called_once_with(str(self.pdf.resolve()))

    def test_optimizer_exception_uses_verified_source(self):
        self.assertEqual(self.upload(Mock(side_effect=RuntimeError("conversion failed"))), (True, "resume.pdf"))

    def test_nonexistent_output_uses_verified_source(self):
        self.assertEqual(self.upload(Mock(return_value=str(self.optimized.with_name("missing.pdf")))), (True, "resume.pdf"))

    def test_unverified_source_never_uploaded(self):
        self.verify.return_value = False
        self.assertEqual(self.upload(Mock(return_value=None)), (False, "No verified resume"))
        self.file_input.send_keys.assert_not_called()

    def test_missing_source_pdf_never_uploaded(self):
        self.pdf.unlink()
        self.assertEqual(self.upload(Mock(return_value=None)), (False, "No verified resume"))
        self.verify.assert_not_called()
        self.file_input.send_keys.assert_not_called()

    def test_failed_optimized_attachment_retries_source(self):
        self.file_input.send_keys.side_effect = [RuntimeError("upload failed"), None]
        self.assertEqual(self.upload(Mock(return_value=str(self.optimized))), (True, "resume.pdf"))
        self.assertEqual(self.file_input.send_keys.call_count, 2)

    def test_failed_source_attachment_never_reports_previous_resume_success(self):
        self.file_input.send_keys.side_effect = RuntimeError("upload failed")
        self.assertEqual(self.upload(Mock(return_value=None)), (False, "No verified resume"))

    def bot_preparation(self, uploaded):
        from test_search_runner import extract_functions
        modal = Mock()
        modal.find_elements.return_value = [self.file_input]
        def navigate(browser, fill, **kwargs):
            fill(modal)
            return modal
        namespace = {
            "use_AI_for_resume": True, "ResumeUploadError": ResumeUploadError,
            "driver": Mock(), "By": SimpleNamespace(NAME="name"),
            "wait_for_application_modal": Mock(return_value=modal),
            "advance_to_review": navigate, "print_lg": Mock(),
            "answer_questions": Mock(return_value=set()),
            "upload_optimized_resume": Mock(return_value=(uploaded, "verified.pdf")),
            "default_resume_docx_path": str(self.source),
            "first_name": "Test", "last_name": "Candidate",
            "resume_upload_target": lambda browser, modal: modal,
        }
        extract_functions(namespace, "prepare_easy_apply")
        return namespace

    def test_bot_preparation_blocks_when_attachment_is_missing(self):
        namespace = self.bot_preparation(False)
        with self.assertRaises(ResumeUploadError):
            namespace["prepare_easy_apply"]("123", "Staff Engineer", "JD", "Bay Area")
        namespace["upload_optimized_resume"].assert_called_once()

    def test_bot_preparation_returns_only_after_verified_attachment(self):
        namespace = self.bot_preparation(True)
        modal, questions, uploaded, resume = namespace["prepare_easy_apply"](
            "123", "Staff Engineer", "JD", "Bay Area")
        self.assertTrue(uploaded)
        self.assertEqual(resume, "verified.pdf")
        namespace["upload_optimized_resume"].assert_called_once_with(
            modal, str(self.source), "JD", "123", "Staff Engineer", "Test", "Candidate")

    def test_later_job_reattaches_configured_pdf_before_ai_prose(self):
        namespace = self.bot_preparation(True)
        upload = Mock(return_value=(True, "configured.pdf"))
        namespace.update(use_AI_for_resume=False, useNewResume=False,
                         ai_descriptive_answers_enabled=True,
                         default_resume_path=str(self.pdf), upload_resume=upload)
        _, _, uploaded, resume = namespace["prepare_easy_apply"](
            "456", "Staff Engineer", "JD", "Bay Area")
        self.assertTrue(uploaded)
        self.assertEqual(resume, "configured.pdf")
        upload.assert_called_once_with(namespace["wait_for_application_modal"].return_value, str(self.pdf))

    def test_source_fallback_requires_configured_pdf_page_count(self):
        from reportlab.pdfgen import canvas

        for pages, expected in ((2, True), (1, False), (3, False)):
            pdf = Path(self.directory.name) / f"source-{pages}.pdf"
            document = canvas.Canvas(str(pdf))
            for page in range(pages):
                document.drawString(50, 700, f"Source resume page {page + 1}")
                document.showPage()
            document.save()
            with patch("modules.config_loader.load_runtime_config", return_value={"expected_page_count": 2}):
                self.assertEqual(verify_source_resume(pdf, pdf), expected)

    def test_unrelated_companion_pdf_cannot_be_used_as_source_fallback(self):
        from docx import Document
        from reportlab.pdfgen import canvas

        doc = Document()
        doc.add_heading("SKILLS", level=1)
        doc.add_paragraph("Python, AWS")
        doc.add_heading("PROFESSIONAL EXPERIENCE", level=1)
        doc.add_paragraph("Designed a production Python platform", style="List Paragraph")
        doc.save(self.source)
        pdf = canvas.Canvas(str(self.pdf))
        for page in range(2):
            pdf.drawString(50, 700, "Unrelated document text")
            pdf.showPage()
        pdf.save()
        with patch("modules.config_loader.load_runtime_config", return_value={"expected_page_count": 2}):
            self.assertFalse(verify_source_resume(self.source, self.pdf))


if __name__ == "__main__":
    unittest.main()
