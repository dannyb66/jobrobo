"""Experience-backed JD targeting and intact ranked Skills, without API calls."""
import hashlib
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from docx import Document
from modules import resume_all_bullets as m
from modules.ai.resume_optimizer import ResumeOptimizer, split_skills,skill_paragraph_text


class JDAlignmentTest(unittest.TestCase):
    def group(self, text):
        from test_resume_all_bullets import inv
        inventory=inv(1,1)
        bullet=inventory['bullets'][0]
        bullet.update(source_text=text,source_text_hash=hashlib.sha256(text.encode()).hexdigest(),fact_anchors=[])
        return inventory,inventory['groups'][0],bullet

    def test_targets_have_same_experience_evidence_and_exclude_unsupported_tools(self):
        inventory,group,bullet=self.group('Designed agentic AI platform for healthcare contract intelligence, serving clients')
        jd='Build agentic systems with Python and .NET. Set technical direction.'
        plan=m.jd_alignment_targets(jd,group,['agentic systems','Python','.NET','technical direction'])
        selected={target['keyword']:target for target in plan['supported_jd_targets']}
        self.assertIn('agentic systems',selected)
        evidence=selected['agentic systems']['source_evidence'][0]
        self.assertEqual(evidence['bullet_id'],bullet['bullet_id'])
        self.assertEqual(evidence['source_text'],bullet['source_text'])
        for term in ('Python','.NET','technical direction'):
            self.assertNotIn(term,selected)
            self.assertIn(term,plan['excluded_jd_terms'])
        payload=m.compact_payload(inventory,jd,'Engineer','grounded',group,[bullet],alignment=plan)
        self.assertEqual(payload['supported_jd_targets'],plan['supported_jd_targets'])
        self.assertIn('JD',payload['requested_bullets'][0]['rewrite_focus'])
        self.assertEqual(payload['requested_bullets'][0]['jd_keyword_opportunities'],['agentic systems'])

    def test_react_ui_framework_is_not_inferred_from_react_reasoning(self):
        _,group,_=self.group('Built ReAct orchestration for AI agents')
        plan=m.jd_alignment_targets('React front-end applications',group,['React'])
        self.assertNotIn('React',[target['keyword'] for target in plan['supported_jd_targets']])
        self.assertIn('React',plan['excluded_jd_terms'])
        self.assertTrue(m.literal_keyword_mention('Built React frontends','React'))

    def test_review_draft_allows_paraphrases_but_keeps_fact_checks(self):
        _,group,bullet=self.group('Built customer workflow services using Python')
        rewritten='Engineered Python services for client workflows'
        item=dict(bullet_id=bullet['bullet_id'],source_text_hash=bullet['source_text_hash'],status='rewritten',
            rewritten_text=rewritten,matched_jd_spans=[],preserved_fact_anchors=[],
            change_reason='Concise equivalent customer/client wording',blocked_reason=None)
        raw={'experience_id':group['experience_id'],'results':[item]}
        self.assertEqual(m.validate_response(raw,group,[bullet])[bullet['bullet_id']]['status'],'blocked')
        draft=m.validate_response(raw,group,[bullet],content_policy='review_draft')[bullet['bullet_id']]
        self.assertEqual(draft['status'],'candidate',draft.get('reason'))
        self.assertIn('customer',draft['source_terms_for_review'])
        item['rewritten_text']+=' with 20% higher revenue'
        self.assertEqual(m.validate_response(raw,group,[bullet],content_policy='review_draft')[bullet['bullet_id']]['status'],'blocked')
        self.assertIsNotNone(m.source_claim_issue('Sole developer of self-hosted inference using Python',
            'Developer of inference using Python',allow_paraphrase=True))
        self.assertIsNotNone(m.source_claim_issue('Built Python and Rust services','Built Python services',allow_paraphrase=True))

    def test_supported_keyword_gain_can_be_meaningful_without_four_changed_words(self):
        source='Established AI operating model with incident response and rollback'
        rewritten='Established AI operating model for reliability with incident response and rollback'
        self.assertFalse(m.substantive(source,rewritten))
        self.assertTrue(m.meaningful_jd_edit(source,rewritten,['reliability']))
        self.assertFalse(m.meaningful_jd_edit('Built Python services','Built Python services with reliability',['reliability']))

    def test_review_draft_cli_is_explicit_and_preflight_stays_provider_free(self):
        import optimize_resume as cli
        from test_resume_pipeline import synthetic_source,mocked_convert
        with TemporaryDirectory() as tmp:
            root=Path(tmp);source=root/'original.docx';synthetic_source(source);mocked_convert(str(source),str(source.with_suffix('.pdf')))
            args=cli.parser().parse_args(['--resume',str(source),'--jd-text','Build Python services',
                '--job-title','Engineer','--output',str(root/'draft'),'--rewrite-all-bullets','--review-draft','--preflight'])
            _,_,_,config,details=cli.plan(args,runtime_config_override={})
            self.assertEqual(config['content_policy'],'review_draft')
            self.assertEqual(details['content_policy'],'review_draft')
            self.assertFalse(details['employer_actions'])
            args.rewrite_all_bullets=False
            with self.assertRaises(ValueError):cli.plan(args,runtime_config_override={})

    def test_keyword_aligned_valid_alternative_wins_over_clause_rotation(self):
        _,group,bullet=self.group('Designed agentic AI platform for healthcare contract intelligence, serving clients')
        first='Serving clients, designed agentic AI platform for healthcare contract intelligence'
        aligned='For healthcare contract intelligence, designed agentic systems as an AI platform serving clients'
        item=dict(bullet_id=bullet['bullet_id'],source_text_hash=bullet['source_text_hash'],status='rewritten',
            rewritten_text=first,alternative_texts=[aligned],matched_jd_spans=[],preserved_fact_anchors=[],
            change_reason='Reframe supported agentic platform contribution for the target role',blocked_reason=None)
        value=m.validate_response({'experience_id':group['experience_id'],'results':[item]},group,[bullet],
            ranking_keywords=['agentic systems'])[bullet['bullet_id']]
        self.assertEqual(value['status'],'candidate',value.get('reason'))
        self.assertEqual(value['rewritten_text'],aligned)
        # Keyword gains must never override actual layout acceptance.
        value=m.validate_response({'experience_id':group['experience_id'],'results':[item]},group,[bullet],
            ranking_keywords=['agentic systems'],fit_check=lambda row:'overflow' if row['rewritten_text']==aligned else None)[bullet['bullet_id']]
        self.assertEqual(value['rewritten_text'],first)

    def test_excluded_jd_tool_is_rejected_even_when_missing_from_the_skill_bank(self):
        _,group,bullet=self.group('Designed agentic AI platform for healthcare contract intelligence, serving clients')
        clean='Serving clients, designed agentic AI platform for healthcare contract intelligence'
        item=dict(bullet_id=bullet['bullet_id'],source_text_hash=bullet['source_text_hash'],status='rewritten',
            rewritten_text=clean+' with .NET and QuantumSDK',alternative_texts=[clean],matched_jd_spans=[],
            preserved_fact_anchors=[],change_reason='Reframe contribution for target role',blocked_reason=None)
        value=m.validate_response({'experience_id':group['experience_id'],'results':[item]},group,[bullet],
            forbidden_keywords=['.NET','QuantumSDK'])[bullet['bullet_id']]
        self.assertEqual(value['rewritten_text'],clean)
        self.assertIn('unsupported excluded JD term',value['candidate_evaluations'][0]['reason'])

    def test_separate_skills_paragraphs_do_not_merge_distinct_entries(self):
        optimizer=ResumeOptimizer.__new__(ResumeOptimizer)
        with TemporaryDirectory() as tmp:
            path=Path(tmp)/'original.docx';doc=Document();doc.add_heading('SKILLS',1)
            doc.add_paragraph('Python, AWS');doc.add_paragraph('React, Docker')
            doc.add_heading('EXPERIENCE',1);doc.save(path)
            document,skills=optimizer.read_docx_resume(str(path))
            self.assertEqual(skills,['Python','AWS','React','Docker'])
            start,end=optimizer._find_skills_section(document)
            optimizer._update_skills_section(document,start,end,['React','Python','AWS','Docker'])
            actual=[entry for paragraph in document.paragraphs[start:end] for entry in split_skills(paragraph.text)]
            self.assertEqual(actual,['React','Python','AWS','Docker'])

    def test_skills_parser_and_document_edit_keep_compound_entries_whole(self):
        source='Snowflake (Snowpark, Cortex AI), Python, Model Context Protocol (MCP), React'
        self.assertEqual(split_skills(source),['Snowflake (Snowpark, Cortex AI)','Python','Model Context Protocol (MCP)','React'])
        self.assertEqual(split_skills('Cloud (AWS, tools [A, B]),\nPython'),['Cloud (AWS, tools [A, B])','Python'])
        optimizer=ResumeOptimizer.__new__(ResumeOptimizer)
        with TemporaryDirectory() as tmp:
            path=Path(tmp)/'original.docx'
            doc=Document();doc.add_heading('SKILLS',1);doc.add_paragraph(source);doc.add_heading('EXPERIENCE',1);doc.save(path)
            document,skills=optimizer.read_docx_resume(str(path))
            self.assertEqual(len(skills),4)
            start,end=optimizer._find_skills_section(document)
            optimizer._update_skills_section(document,start,end,['React','Python','Snowflake (Snowpark, Cortex AI)'])
            actual=split_skills(' '.join(p.text for p in document.paragraphs[start:end]))
            self.assertEqual(actual,['React','Python','Snowflake (Snowpark, Cortex AI)','Model Context Protocol (MCP)'])

    def test_layout_partition_preserves_relevance_and_complete_entries(self):
        terms=['Python','Snowflake (Snowpark, Cortex AI)','React','LangGraph','Docker']
        rows=m._ranked_skill_lines(terms,3,36,len)
        self.assertIsNotNone(rows)
        self.assertEqual([term for row in rows for term in row],terms)
        self.assertEqual(len(rows),3)
        self.assertTrue(all(len(', '.join(row)+(',' if i<2 else ''))<=36 for i,row in enumerate(rows)))
        self.assertIsNone(m._ranked_skill_lines(terms,1,10,len))
        # Physical wrapping within a compound entry must preserve the logical
        # entry when the text is read back from one DOCX paragraph.
        words=', '.join(terms).split()
        rows=m._ranked_skill_lines(words,4,28,len,separator=' ',continuation='')
        self.assertIsNotNone(rows)
        wrapped='\n'.join(' '.join(row) for row in rows)
        self.assertEqual(split_skills(wrapped.replace('\n',' ')),terms)

    def test_hyphen_slash_continuations_preserve_skill_text_on_readback(self):
        self.assertEqual(skill_paragraph_text('Tool/\nFunction Calling, Human-in-\nthe-Loop'),
            'Tool/Function Calling, Human-in-the-Loop')
        self.assertEqual(skill_paragraph_text('Reasoning / \nPlanning'), 'Reasoning /  Planning')
        terms=['Tool/Function Calling','Human-in-the-Loop','Reasoning / Planning']
        wrapped='Tool/\nFunction Calling, Human-in-\nthe-Loop, Reasoning / \nPlanning'
        self.assertEqual(split_skills(skill_paragraph_text(wrapped)),terms)
        optimizer=ResumeOptimizer.__new__(ResumeOptimizer)
        with TemporaryDirectory() as tmp:
            path=Path(tmp)/'wrapped.docx';doc=Document();doc.add_heading('SKILLS',1)
            doc.add_paragraph(wrapped);doc.add_heading('EXPERIENCE',1);doc.save(path)
            self.assertEqual(optimizer.read_docx_resume(str(path))[1],terms)


if __name__=='__main__':
    unittest.main()
