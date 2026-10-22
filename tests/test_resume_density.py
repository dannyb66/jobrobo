"""Source-relative density regressions; actual PDF coordinates, no network."""
import json
import os
import shutil
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from docx import Document
from reportlab.pdfgen import canvas

from modules.ai.resume_optimizer import ResumeOptimizer
from modules.resume_layout import compare_density, paragraph_density
import test_resume_pipeline as pipeline


def response(entries):
    return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(
        content=json.dumps({"rewritten_experience": entries})))])


class ResumeDensityTest(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.optimizer = ResumeOptimizer.__new__(ResumeOptimizer)

    def artifact(self, name, text, lines, style="List Paragraph"):
        docx, pdf = self.root / (name + ".docx"), self.root / (name + ".pdf")
        doc = Document()
        doc.add_paragraph(text, style=style)
        doc.save(docx)
        writer = canvas.Canvas(str(pdf))
        writer.setFont("Helvetica", 10)
        for i, line in enumerate(lines):
            writer.drawString(72, 730 - i * 15, line)
        writer.save()
        skills = (0, 1) if style != "List Paragraph" else (None, 1)
        return paragraph_density(str(pdf), str(docx), skills)

    def test_same_character_count_cannot_hide_sparse_glyph_width(self):
        source = self.artifact("source", "WWWW WWWW", ["WWWW WWWW"])
        output = self.artifact("output", "iiii iiii", ["iiii iiii"])
        result = compare_density(source, output)
        self.assertFalse(result["passed"])
        self.assertIn("occupied_width_outside_92_to_108_percent", result["violations"][0]["reasons"])

    def test_same_total_width_cannot_hide_changed_line_usage(self):
        source = self.artifact("source", "Built Python services", ["Built Python", "services"])
        output = self.artifact("output", "Built Python services", ["Built Python services"])
        self.assertFalse(compare_density(source, output)["passed"])

    def test_skills_must_preserve_number_and_fill_of_rendered_lines(self):
        source = self.artifact("source", "Python, AWS, REST APIs", ["Python, AWS,", "REST APIs"], "Normal")
        output = self.artifact("output", "Python, AWS", ["Python, AWS"], "Normal")
        result = compare_density(source, output)
        self.assertFalse(result["passed"])
        self.assertEqual(result["violations"][0]["kind"], "skills")

    def test_unmatched_source_is_explicit_failure(self):
        source = self.artifact("source", "Missing paragraph", ["Different PDF"])
        self.assertFalse(compare_density(source, source)["passed"])

    def test_list_styled_skills_remain_skills_in_density_and_inventory(self):
        from modules.resume_all_bullets import inventory,strict_density
        for style in ('List Paragraph', 'List Bullet'):
            with self.subTest(style=style):
                source=self.root/'list-skills.docx'
                pipeline.synthetic_source(source)
                doc=Document(source)
                doc.paragraphs[1].style=style
                doc.save(source)
                pdf=source.with_suffix('.pdf')
                pipeline.mocked_convert(str(source),str(pdf))
                skills_range=self.optimizer._find_skills_section(Document(source))
                records=paragraph_density(pdf,source,skills_range)
                self.assertEqual([r['kind'] for r in records],['skills','bullet','bullet'])
                self.assertTrue(all(r['matched'] for r in records))
                inventoried=inventory(source,self.optimizer,pdf)
                self.assertEqual(len(inventoried['bullets']),2)
                self.assertTrue(strict_density(records,records,inventoried)['passed'])

    def test_repeated_identical_bullets_match_in_reading_order(self):
        docx, pdf = self.root / "repeated.docx", self.root / "repeated.pdf"
        doc = Document()
        for _ in range(2):
            doc.add_paragraph("Built APIs", style="List Paragraph")
        doc.save(docx)
        writer = canvas.Canvas(str(pdf))
        for _ in range(2):
            writer.drawString(72, 730, "Built APIs")
            writer.showPage()
        writer.save()
        records = paragraph_density(str(pdf), str(docx), (None, 2))
        self.assertEqual([r["lines"][0]["page"] for r in records], [1, 2])

    def test_bounded_api_retry_rejects_short_bullet_and_retains_source(self):
        source = [{"company": "Example", "job_title": "Engineer", "bullets": [
            "Built reliable production APIs using Python and AWS, improving service latency by 30% for enterprise clients"]}]
        self.optimizer.client = Mock()
        self.optimizer.rewrite_response_attempts = 2
        self.optimizer.client.chat.completions.create.return_value = response([
            dict(source[0], bullets=["Built Python APIs"])])
        result = self.optimizer._call_openai_to_rewrite_bullets(source, "Engineer", "Production API reliability")
        self.assertEqual(result, source)
        self.assertEqual(self.optimizer.client.chat.completions.create.call_count, 2)
        self.assertIn("minimum", self.optimizer.last_bullet_rewrite_error)
        retry = self.optimizer.client.chat.completions.create.call_args.kwargs["messages"][-1]["content"]
        self.assertIn("minimum and maximum", retry)
        self.assertIn("never add filler", retry)

    def test_prompt_includes_actual_source_line_targets(self):
        self.optimizer.source_bullet_density = {"Built APIs": {"line_count": 2, "width": 500}}
        message = self.optimizer._generate_user_message([
            {"company": "Example", "job_title": "Engineer", "bullets": ["Built APIs"]}], "Engineer", "API JD")
        self.assertIn('"line_count": 2', message)
        self.assertIn('"minimum"', message)
        self.assertIn('"maximum"', message)

    def test_skill_cap_does_not_erase_source_density_and_order_is_tailored(self):
        doc = Document()
        doc.add_heading("SKILLS", 1)
        source = ["Tool " + str(i) for i in range(65)] + ["Python", "AWS"]
        doc.add_paragraph(", ".join(source))
        selected = self.optimizer.optimize_skills(source, ["AWS", "Python", "Scala"], "Required: AWS Python", max_count=50)
        self.optimizer._update_skills_section(doc, *self.optimizer._find_skills_section(doc), selected)
        output = doc.paragraphs[1].text.split(", ")
        self.assertEqual(output[:2], ["AWS", "Python"])
        self.assertCountEqual(output, source)
        self.assertNotIn("Scala", output)


class DensityRecoveryPipelineTest(unittest.TestCase):
    def setUp(self):
        self.runner = pipeline.ResumePipelineTest(methodName="test_all_modes_change_only_intended_sections_and_record_reports")
        self.runner.setUp()
        self.addCleanup(self.runner.doCleanups)

    def test_real_pdf_recovery_retains_only_sparse_bullet_and_records_failure(self):
        def rewrite(entries, *args):
            return [dict(entry, bullets=["Built API", "Enhanced API reliability on AWS"]) for entry in entries]
        self.runner.rewrite.side_effect = rewrite
        result = self.runner.run_pipeline("bullets")
        self.assertIsNotNone(result)
        report = self.runner.report(result)
        attempt = report["attempts"][0]
        self.assertFalse(attempt["initial_format_verification"]["passed"])
        self.assertTrue(attempt["format_verification"]["passed"])
        self.assertEqual(len(attempt["density_source_retention"]), 1)
        text = "\n".join(p.text for p in Document(Path(result).with_suffix(".docx")).paragraphs)
        self.assertIn("Built production services using Python", text)
        self.assertIn("Enhanced API reliability on AWS", text)
        self.assertEqual(self.runner.rewrite.call_count, 1)
        self.assertTrue(Path(result).with_suffix(".before_density.pdf").exists())

    def test_missing_key_with_previous_artifact_returns_failure_without_reusing_it(self):
        previous = self.runner.run_pipeline("skills")
        self.assertIsNotNone(previous)
        with patch.dict(os.environ, {"OPENAI_API_KEY": ""}):
            self.assertIsNone(self.runner.run_pipeline("skills"))
        self.assertEqual(self.runner.report()["status"], "failed")


if __name__ == "__main__":
    unittest.main()
