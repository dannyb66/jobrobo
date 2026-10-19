"""Per-bullet source fallback preserves identity, reporting and attachment checks."""
import json
import socket
import threading
import unittest
from copy import deepcopy
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import Mock,patch
from docx import Document
from modules import resume_all_bullets as m
from modules.ai import resume_optimizer as engine
from modules.resume_upload import upload_verified_resume
from modules.resume_progress import ResumeConsoleProgress
from test_resume_pipeline import synthetic_source,mocked_convert


class SourceFallbackTest(unittest.TestCase):
    def setUp(self):
        self.tmp=TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name);self.source=self.root/'source.docx'
        synthetic_source(self.source);mocked_convert(str(self.source),str(self.source.with_suffix('.pdf')))
        self.config={**m.DEFAULTS,'rewrite_all_bullets':True,'rewrite_source_fallback_enabled':True,
            'source_checks_enabled':False,'content_policy':'synthetic_jd_test','optimization_mode':'full',
            'replace_job_title':False,'expected_page_count':1,'rewrite_bullets_per_request':1,
            'rewrite_requests_per_minute':1000000,'rewrite_tokens_per_minute':1000000000}
        self.jd='Build Python backend APIs with AWS.'

    def run_resume(self,mode='timeout',*,enabled=True,execution='provider',bad_layout=False):
        config={**self.config,'rewrite_source_fallback_enabled':enabled}
        cancel=threading.Event();events=[]
        def provider(stage,payload,timeout):
            if stage=='keywords':data={'keywords':['Python','AWS']}
            else:
                bullet=payload['requested_bullets'][0]
                fail=bullet['ordinal']==1 or mode=='all_failed'
                if fail and mode=='cancel':cancel.set()
                if fail and mode!='width':raise TimeoutError('fixture timeout')
                text='Built reliable backend APIs using Python and AWS.' if not fail else 'Overlong candidate '*100
                data={'experience_id':payload['experience_id'],'results':[{
                    'bullet_id':bullet['bullet_id'],'source_text_hash':bullet['source_text_hash'],
                    'status':'rewritten','rewritten_text':text,'alternative_texts':[],
                    'matched_jd_spans':[],'preserved_fact_anchors':[],
                    'change_reason':'Fixture JD rewrite','blocked_reason':None}]}
            return SimpleNamespace(id='mock-provider',usage=SimpleNamespace(model_dump=lambda:
                {'prompt_tokens':0,'completion_tokens':0}),choices=[SimpleNamespace(message=
                SimpleNamespace(content=json.dumps(data)))])
        # Another legacy test reloads this module; resolve the current class
        # when patching so conversion/layout effects cannot escape the fixture.
        optimizer_class=engine.ResumeOptimizer
        original_verify=optimizer_class.verify_format
        def verify(instance,*args):
            result=original_verify(instance,*args)
            if bad_layout:result['passed']=False;result['checks']['no_horizontal_clipping']=False
            return result
        output=self.root/(mode+'-'+str(enabled)+'-'+execution+'-'+str(bad_layout))
        with patch.object(socket.socket,'connect',side_effect=AssertionError('NETWORK FORBIDDEN')),\
             patch.object(optimizer_class,'convert_docx_to_pdf',side_effect=mocked_convert),\
             patch.object(optimizer_class,'verify_format',new=verify):
            result=m.run_all_bullets(self.jd,self.source,'123','Engineer',config,
                provider=provider,execution_kind=execution,output_root=output,cancel=cancel,progress=events.append)
        report=json.loads((output/'match-report.json').read_text())
        coverage=json.loads((output/'rewrite-coverage.json').read_text())
        return result,report,coverage,config,events

    def test_timeout_and_width_failures_retain_exact_source_and_allow_verified_upload(self):
        for mode in ('timeout','width'):
            with self.subTest(mode=mode):
                result,report,coverage,config,events=self.run_resume(mode)
                self.assertIsNotNone(result,str(report.get('errors')))
                self.assertEqual(report['status'],'complete')
                self.assertTrue(report['document_complete']);self.assertFalse(report['rewrite_completed'])
                self.assertEqual(report['completion_mode'],'source_fallback')
                self.assertEqual(report['coverage']['accepted'],1)
                self.assertEqual(report['coverage']['source_fallback'],1)
                self.assertEqual(report['coverage']['failed'],1)
                self.assertEqual(report['coverage']['unresolved'],0)
                row=next(r for r in coverage['bullets'] if r['status']=='source_fallback')
                self.assertEqual(row['source_text'],row['rewritten_text']);self.assertTrue(row['fallback_reason'])
                original=Document(self.source).paragraphs[row['paragraph_index']]
                attached=Document(Path(result).with_suffix('.docx')).paragraphs[row['paragraph_index']]
                self.assertEqual(original._p.xml,attached._p.xml)
                expected=dict(expected_jd=self.jd,expected_job_id='123',expected_source_docx=self.source,
                    expected_source_checks_enabled=False,expected_config=config)
                self.assertTrue(m.verify_strict_attachment(result,**expected))
                self.assertFalse(m.verify_strict_attachment(result,**{**expected,'expected_job_id':'wrong'}))
                self.assertFalse(m.verify_strict_attachment(result,**{**expected,'expected_config':{**config,'rewrite_source_fallback_enabled':False}}))
                modal=Mock()
                self.assertTrue(upload_verified_resume(modal,str(self.source),self.jd,'123',
                    optimize=Mock(return_value=result),runtime_config_override=config)[0])
                modal.find_element.return_value.send_keys.assert_called_once_with(str(Path(result).resolve()))
                messages=[]
                with ResumeConsoleProgress(messages.append,'123') as progress:progress(events[-1])
                self.assertIn('1 bullets rewritten; 1 original bullets retained',messages[-1])
                self.assertIn('upload eligible: True',messages[-1])

    def test_disabled_fallback_preserves_strict_rewrite_gate(self):
        result,report,_,_,_=self.run_resume(enabled=False)
        self.assertIsNone(result);self.assertEqual(report['status'],'partial')
        self.assertFalse(report['strict_attachment_eligible'])

    def test_cancellation_total_failure_bad_layout_and_offline_results_do_not_upload(self):
        for mode,execution,bad_layout in [('cancel','provider',False),('all_failed','provider',False),
                ('timeout','provider',True),('timeout','offline_fixture',False)]:
            with self.subTest(mode=mode,execution=execution,bad_layout=bad_layout):
                result,report,_,_,_=self.run_resume(mode,execution=execution,bad_layout=bad_layout)
                self.assertIsNone(result);self.assertFalse(report['strict_attachment_eligible'])

    def test_fallback_flag_and_source_identity_are_validated(self):
        from test_resume_all_bullets import inv
        for value in (1,'true',None):
            with self.assertRaises(ValueError):m.validate_options({'rewrite_source_fallback_enabled':value},'full')
        inventory=inv(2,1)
        rows=[dict(b,status='accepted',rewritten_text='Rewritten') for b in inventory['bullets']]
        rows[1].update(status='blocked',reason='fixture failure')
        m.retain_failed_source_bullets(rows)
        self.assertTrue(m.source_fallback_coverage_valid(rows,inventory))
        rows[1]['rewritten_text']='Different original text'
        self.assertFalse(m.source_fallback_coverage_valid(rows,inventory))
        rows[1]['rewritten_text']=rows[1]['source_text'];rows[1]['source_text_hash']='stale'
        self.assertFalse(m.source_fallback_coverage_valid(rows,inventory))


if __name__=='__main__':unittest.main()
