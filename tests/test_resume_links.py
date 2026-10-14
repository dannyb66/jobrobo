import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Pt
from reportlab.pdfgen import canvas
from reportlab.pdfbase.pdfmetrics import stringWidth

from modules.ai.resume_optimizer import ResumeOptimizer
from modules.resume_links import read_docx_links, read_pdf_links
from modules.resume_upload import verify_source_resume


TARGET = "https://example.test/work?a=1&b=2"
LABEL = "Portfolio & work"


def docx_with_link(path, target=TARGET):
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10)
    paragraph = doc.add_paragraph("Contact ")
    relationship_id = doc.part.relate_to(target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = LABEL
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    doc.add_heading("SKILLS", level=1)
    doc.add_paragraph("Python, AWS")
    doc.add_heading("PROFESSIONAL EXPERIENCE", level=1)
    doc.add_paragraph("Built production APIs with Python", style="List Paragraph")
    doc.save(path)


def pdf_with_links(path, target=TARGET, linked=True, zero_area=False, duplicate=False, mail=False):
    writer = canvas.Canvas(str(path))
    writer.setFont("Helvetica", 10)
    writer.drawString(72, 730, "Contact " + LABEL)
    left = 72 + stringWidth("Contact ", "Helvetica", 10)
    right = left + stringWidth(LABEL, "Helvetica", 10)
    if linked:
        writer.linkURL(target, (left, 728, left if zero_area else right, 740), relative=0)
    if duplicate:
        writer.drawString(72, 710, LABEL)
        writer.linkURL(target, (72, 708, 72 + stringWidth(LABEL, "Helvetica", 10), 720), relative=0)
    if mail:
        label = "test@example.test"
        writer.drawString(72, 690, label)
        writer.linkURL("mailto:" + label, (72, 688, 72 + stringWidth(label, "Helvetica", 10), 700), relative=0)
    for y, line in zip((650, 630, 590, 570), (
            "SKILLS", "Python, AWS", "PROFESSIONAL EXPERIENCE", "Built production APIs with Python")):
        writer.drawString(72, y, line)
    writer.showPage()
    writer.save()


class ResumeHyperlinkTest(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.docx = self.root / "source.docx"
        self.source = self.root / "source.pdf"
        self.output = self.root / "output.pdf"
        docx_with_link(self.docx)
        pdf_with_links(self.source)
        shutil.copy2(self.source, self.output)
        self.optimizer = ResumeOptimizer.__new__(ResumeOptimizer)

    def report(self):
        return self.optimizer.verify_format(str(self.docx), str(self.docx), str(self.source), str(self.output), 1)

    def test_link_targets_and_click_areas_pass(self):
        report = self.report()
        self.assertTrue(report["passed"], report)
        self.assertEqual(report["details"]["hyperlinks"]["expected_count"], 1)
        self.assertEqual(report["details"]["hyperlinks"]["clickable_count"], 1)

    def test_blue_looking_text_without_annotation_is_rejected(self):
        pdf_with_links(self.output, linked=False)
        report = self.report()
        self.assertFalse(report["passed"])
        self.assertFalse(report["checks"]["hyperlinks_preserved"])
        self.assertEqual(report["details"]["hyperlinks"]["missing"][0]["target"], TARGET)

    def test_equal_link_count_with_wrong_target_is_rejected(self):
        pdf_with_links(self.output, target="https://example.test/wrong")
        self.assertFalse(self.report()["checks"]["hyperlinks_preserved"])

    def test_zero_size_click_area_is_not_a_preserved_link(self):
        pdf_with_links(self.output, zero_area=True)
        report = self.report()
        self.assertFalse(report["checks"]["hyperlinks_preserved"])
        self.assertEqual(len(report["details"]["hyperlinks"]["invalid_click_areas"]), 1)

    def test_export_added_trailing_slash_is_allowed_but_query_changes_are_not(self):
        docx_with_link(self.docx, "https://example.test/work")
        pdf_with_links(self.output, target="https://example.test/work/")
        self.assertTrue(self.report()["checks"]["hyperlinks_preserved"])
        pdf_with_links(self.output, target="https://example.test/work/?changed=1")
        self.assertFalse(self.report()["checks"]["hyperlinks_preserved"])

    def test_linkless_source_fallback_is_rejected_using_docx_targets(self):
        pdf_with_links(self.source, linked=False)
        with patch("modules.config_loader.load_runtime_config", return_value={"expected_page_count": 1}):
            self.assertFalse(verify_source_resume(self.docx, self.source))

    def test_verified_source_fallback_preserves_clickable_links(self):
        with patch("modules.config_loader.load_runtime_config", return_value={"expected_page_count": 1}):
            self.assertTrue(verify_source_resume(self.docx, self.source))

    def test_pdf_input_retains_targets_labels_and_duplicate_occurrences(self):
        pdf_with_links(self.source, duplicate=True, mail=True)
        self.optimizer.save_optimized_pdf(str(self.source), ["Python", "Kubernetes"], {}, str(self.output))
        links = read_pdf_links(self.output, include_text=True)
        self.assertEqual([link["target"] for link in links].count(TARGET), 2)
        self.assertIn("mailto:test@example.test", [link["target"] for link in links])
        self.assertTrue(all(link["clickable"] for link in links))
        self.assertEqual([link["label"] for link in links].count(LABEL), 2)
        self.assertTrue(self.optimizer.verify_format("", "", str(self.source), str(self.output), 1)["passed"])

    def test_pdf_input_without_annotations_still_works(self):
        pdf_with_links(self.source, linked=False)
        self.optimizer.save_optimized_pdf(str(self.source), ["Python"], {}, str(self.output))
        self.assertEqual(read_pdf_links(self.output), [])

    def test_plain_earlier_label_does_not_steal_later_linked_occurrence(self):
        writer = canvas.Canvas(str(self.source))
        writer.setFont("Helvetica", 10)
        writer.drawString(72, 730, "Plain Portfolio")
        writer.drawString(72, 710, "Actual Portfolio")
        left = 72 + stringWidth("Actual ", "Helvetica", 10)
        writer.linkURL(TARGET, (left, 708, left + stringWidth("Portfolio", "Helvetica", 10), 720))
        writer.drawString(72, 670, "SKILLS")
        writer.drawString(72, 650, "Python, AWS")
        writer.drawString(72, 610, "PROFESSIONAL EXPERIENCE")
        writer.drawString(72, 590, "Built production APIs")
        writer.save()
        self.optimizer.save_optimized_pdf(str(self.source), ["Python", "AWS"], {}, str(self.output))
        links = read_pdf_links(self.output, include_text=True)
        self.assertEqual(len(links), 1)
        self.assertEqual(links[0]["context"], "Actual Portfolio")
        self.assertEqual(links[0]["target"], TARGET)
        self.assertTrue(self.optimizer.verify_format("", "", str(self.source), str(self.output), 1)["passed"])

    def test_identical_plain_and_linked_rows_fail_closed_instead_of_relocating_link(self):
        writer = canvas.Canvas(str(self.source))
        writer.setFont("Helvetica", 10)
        for y in (730, 710):
            writer.drawString(72, y, "Portfolio")
        writer.linkURL(TARGET, (72, 708, 72 + stringWidth("Portfolio", "Helvetica", 10), 720))
        writer.drawString(72, 670, "SKILLS")
        writer.drawString(72, 650, "Python, AWS")
        writer.drawString(72, 610, "PROFESSIONAL EXPERIENCE")
        writer.drawString(72, 590, "Built production APIs")
        writer.save()
        self.output.unlink()
        with self.assertRaisesRegex(ValueError, "ambiguous; use DOCX input"):
            self.optimizer.save_optimized_pdf(str(self.source), ["Python", "AWS"], {}, str(self.output))
        self.assertFalse(self.output.exists())

    def test_repeated_label_in_same_row_fails_closed_instead_of_linking_first_word(self):
        writer = canvas.Canvas(str(self.source))
        writer.setFont("Helvetica", 10)
        writer.drawString(72, 730, "Portfolio Portfolio")
        left = 72 + stringWidth("Portfolio ", "Helvetica", 10)
        writer.linkURL(TARGET, (left, 728, left + stringWidth("Portfolio", "Helvetica", 10), 740))
        writer.drawString(72, 670, "SKILLS")
        writer.drawString(72, 650, "Python, AWS")
        writer.drawString(72, 610, "PROFESSIONAL EXPERIENCE")
        writer.drawString(72, 590, "Built production APIs")
        writer.save()
        self.output.unlink()
        with self.assertRaisesRegex(ValueError, "ambiguous; use DOCX input"):
            self.optimizer.save_optimized_pdf(str(self.source), ["Python", "AWS"], {}, str(self.output))
        self.assertFalse(self.output.exists())

    def test_pdf_input_without_visible_link_label_fails_instead_of_dropping_link(self):
        writer = canvas.Canvas(str(self.source))
        writer.drawString(72, 700, "SKILLS")
        writer.drawString(72, 680, "Python")
        writer.drawString(72, 640, "PROFESSIONAL EXPERIENCE")
        writer.drawString(72, 620, "Built APIs")
        writer.linkURL(TARGET, (400, 400, 450, 450), relative=0)
        writer.save()
        self.output.unlink()
        with self.assertRaisesRegex(ValueError, "use DOCX input"):
            self.optimizer.save_optimized_pdf(str(self.source), ["Python"], {}, str(self.output))
        self.assertFalse(self.output.exists())

    def test_pdf_only_verification_detects_missing_source_links(self):
        pdf_with_links(self.output, linked=False)
        report = self.optimizer.verify_format("", "", str(self.source), str(self.output), 1)
        self.assertFalse(report["checks"]["hyperlinks_preserved"])

    def test_docx_links_inside_headers_are_checked(self):
        doc = Document(self.docx)
        paragraph = doc.sections[0].header.paragraphs[0]
        relationship_id = paragraph.part.relate_to("https://example.test/header", RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), relationship_id)
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "Header link"
        run.append(text)
        link.append(run)
        paragraph._p.append(link)
        doc.save(self.docx)
        self.assertEqual(len(read_docx_links(self.docx)), 2)
        self.assertFalse(self.report()["checks"]["hyperlinks_preserved"])


if __name__ == "__main__":
    unittest.main()
