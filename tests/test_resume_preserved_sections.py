"""Academic distinctions never enter a rewrite, even when old results are reused."""
import json
import unittest
from copy import deepcopy
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from docx import Document
from modules import resume_all_bullets as m, resume_review_draft as drafts
from modules.resume_sections import academic_paragraph_indexes, academic_sections_preserved, exclude_academic_inventory
from modules.ai.resume_optimizer import ResumeOptimizer
from test_resume_all_bullets import inv, response, config


class AcademicPreservationTest(unittest.TestCase):
    def setUp(self):self.optimizer=ResumeOptimizer.__new__(ResumeOptimizer)

    def document(self):
        doc=Document();doc.add_heading('SKILLS',1);doc.add_paragraph('Python, AWS')
        doc.add_heading('PROFESSIONAL EXPERIENCE',1)
        doc.add_paragraph('Example company\tRemote');doc.add_paragraph('Engineer\tJanuary 2020 - Present')
        doc.add_paragraph('Built Python APIs',style='List Paragraph')
        doc.add_paragraph('  Academic   Distinctions:  ')
        p=doc.add_paragraph(style='List Paragraph');p.add_run('Earned award ').bold=True;p.add_run('with 3.9 GPA').italic=True
        doc.add_heading('PROJECTS',1);doc.add_paragraph('Next project')
        doc.add_paragraph('Built AWS services',style='List Paragraph')
        return doc

    def test_heading_variants_and_boundary_keep_next_section_editable(self):
        doc=self.document()
        self.assertEqual(academic_paragraph_indexes(doc,self.optimizer._is_resume_section),{6,7})
        self.assertTrue(self.optimizer._is_resume_section(doc.paragraphs[6]))
        clone=deepcopy(doc)
        clone.paragraphs[-1].runs[0].text='Changed project'
        self.assertTrue(academic_sections_preserved(doc,clone,self.optimizer._is_resume_section))
        clone.paragraphs[7].runs[0].bold=False
        self.assertFalse(academic_sections_preserved(doc,clone,self.optimizer._is_resume_section))
        clone=deepcopy(doc);clone.paragraphs[7].runs[0].text='Invented award'
        self.assertFalse(academic_sections_preserved(doc,clone,self.optimizer._is_resume_section))

    def test_inventory_preserves_awards_and_keeps_other_ids_and_ordinals_stable(self):
        with TemporaryDirectory() as tmp:
            source=Path(tmp)/'source.docx';self.document().save(source)
            measurements=[{'paragraph_index':i,'matched':True,'line_count':1,
                'lines':[{'x0':72,'x1':250,'width':178,'page':1}]} for i in (5,7,10)]
            with patch.object(m,'paragraph_density',return_value=measurements):
                inventory=m.inventory(source,self.optimizer,'unused.pdf')
            self.assertEqual([b['ordinal'] for b in inventory['bullets']],[0,2])
            self.assertEqual([b['ordinal'] for b in inventory['preserved_bullets']],[1])
            self.assertEqual(inventory['preserved_bullets'][0]['source_text'],'Earned award with 3.9 GPA')
            self.assertEqual(inventory['bullets'][-1]['bullet_id'],m.sha(f"{inventory['source_sha256']}:10:2"))

    def test_ordinary_title_and_bank_paths_leave_awards_alone(self):
        doc=self.document();before=deepcopy(doc)
        self.assertEqual(len(self.optimizer._extract_experience_bullets(doc)),1)
        self.optimizer.replace_job_titles(doc,'Staff Engineer')
        self.assertTrue(academic_sections_preserved(before,doc,self.optimizer._is_resume_section))
        # A configured bank cannot use the protected heading as a company target.
        self.assertFalse(self.optimizer._replace_first_company_bullets(doc,'Academic Distinctions:',['Invented award']))
        self.assertTrue(academic_sections_preserved(before,doc,self.optimizer._is_resume_section))

    def test_old_inventories_cannot_send_academic_group_to_any_provider(self):
        inventory=inv(2,2);inventory['groups'][1]['section']='ACADEMIC DISTINCTIONS'
        for policy in ('grounded','review_draft','synthetic_jd_test','unchecked_review'):
            captured=[]
            def provider(stage,payload,timeout):
                captured.extend(payload['requested_bullet_ids']);return response(payload)
            scheduler=m.Scheduler(config(),2)
            if policy=='unchecked_review':
                result=drafts.collect_proposals(inventory,'Python','Engineer',scheduler,provider)
            else:result=m.rewrite_inventory(inventory,'Python','Engineer',policy,scheduler,provider)
            self.assertEqual(captured,['0'])
            self.assertEqual([b['bullet_id'] for b in result],['0'])
        self.assertEqual(len(inventory['bullets']),2,'Filtering must not mutate archived inventory')

    def test_unchecked_writer_ignores_old_academic_proposals(self):
        with TemporaryDirectory() as tmp:
            root=Path(tmp);source=root/'original.docx';doc=self.document();doc.save(source)
            source.with_suffix('.pdf').write_bytes(b'non-rendering test fixture')
            measurements=[{'paragraph_index':i,'matched':True,'line_count':1,
                'lines':[{'x0':72,'x1':250,'width':178,'page':1}]} for i in (5,7,10)]
            with patch.object(m,'paragraph_density',return_value=measurements):
                inventory=m.inventory(source,self.optimizer,'unused.pdf')
            academic=inventory['preserved_bullets'][0]
            inventory['bullets'].append(academic)
            inventory['groups'].append({'experience_id':academic['experience_id'],
                'section':'ACADEMIC DISTINCTIONS','bullets':[academic]})
            records=[dict(b,rewritten_text='Rewritten '+b['source_text'],status='unchecked_proposal',
                strict_validation_status='blocked',requested=True,response_received=True) for b in inventory['bullets']]
            # Other configuration tests reload this module; patch its current
            # class rather than an alias captured during test discovery.
            with patch('modules.ai.resume_optimizer.ResumeOptimizer.convert_docx_to_pdf',return_value=False):
                report=drafts.write_review(root/'review',source,'Python','Engineer','test',inventory,
                    records,m.Scheduler(config(),3),config(),{},execution_kind='saved_provider_response')
            output=Document(root/'review/optimized.docx')
            self.assertTrue(academic_sections_preserved(doc,output,self.optimizer._is_resume_section))
            self.assertEqual(report['coverage']['inventoried'],2)
            self.assertEqual(report['preserved_sections']['Academic Distinctions']['bullet_count'],1)
            self.assertNotEqual(doc.paragraphs[5].text,output.paragraphs[5].text)


if __name__=='__main__':unittest.main()
