"""Deterministic contract tests; no provider traffic."""
import json
import threading
import time
import unittest
from copy import deepcopy
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import patch
from docx import Document

from modules import resume_all_bullets as m
from modules.ai.resume_optimizer import ResumeOptimizer


def config(**values):
    return {**m.DEFAULTS, 'rewrite_all_bullets':True, 'rewrite_requests_per_minute':1000000,
            'rewrite_tokens_per_minute':1000000000, **values}


def bullet(i,exp='e'):
    source=f'Built Python API services for customer workflows number {i}'
    return dict(bullet_id=str(i),experience_id=exp,source_text=source,source_text_hash=m.sha(source),
                fact_anchors=[str(i),'Python','API'],paragraph_index=i,ordinal=i,
                usable_line_width=500,right_boundary=550,text_left=50, max_rendered_lines=1,
                font_runs=[],soft_character_estimate=100)


def inv(count=4,groups=2):
    bs=[bullet(i,str(i%groups)) for i in range(count)]
    return dict(source_sha256='s',bullets=bs,groups=[dict(experience_id=str(g),context=['original title and dates'],section='Experience',company='Source company',historical_title_dates=['Source title'],
                bullets=[b for b in bs if b['experience_id']==str(g)]) for g in range(groups)])


def response(payload):
    return dict(experience_id=payload['experience_id'],results=[dict(bullet_id=b['bullet_id'],
        source_text_hash=b['source_text_hash'],status='rewritten',
        rewritten_text=f"Built customer workflow services number {b['bullet_id']} through Python API development",
        matched_jd_spans=['Python'],preserved_fact_anchors=b['fact_anchors'],
        change_reason='Emphasizes API development and service scope',blocked_reason=None) for b in payload['requested_bullets']])


class AllBulletsTest(unittest.TestCase):
    def test_config_independence_and_validation(self):
        self.assertFalse(m.validate_options({},'full')['rewrite_all_bullets'])
        for bad in ('true',1,None):
            with self.assertRaises(ValueError):m.validate_options({'rewrite_all_bullets':bad},'full')
        for bad in (0,5,2.0,True):
            with self.assertRaises(ValueError):m.validate_options({'rewrite_experience_concurrency':bad},'full')
        with self.assertRaises(ValueError):m.validate_options(config(),'skills')
        for policy in ('grounded','review_draft'):
            self.assertTrue(m.validate_options(config(content_policy=policy),'bullets')['rewrite_all_bullets'])

    def test_context_and_parallel_order_and_sequential_equality(self):
        source=inv(13,3); jd='Python\nFULL EXACT JD '+('details '*500)
        outputs=[]
        for concurrency in (1,2,4):
            lock=threading.Lock(); active=0; maximum=0; captured=[]
            def provider(stage,payload,timeout):
                nonlocal active, maximum
                with lock:active+=1;maximum=max(maximum,active);captured.append(deepcopy(payload))
                self.assertEqual(payload['job_description'],jd)
                group=next(g for g in source['groups'] if g['experience_id']==payload['experience_id'])
                self.assertEqual(payload['source_experience']['context'],group['context'])
                self.assertEqual(payload['source_experience']['historical_title_dates'],group['historical_title_dates'])
                for original,compact in zip(group['bullets'],payload['source_experience']['bullets']):
                    for key in ('bullet_id','ordinal','source_text','source_text_hash','fact_anchors'):self.assertEqual(original[key],compact[key])
                time.sleep(.012 if payload['experience_id']=='0' else .001)
                value=response(payload);value['results'].reverse()
                with lock:active-=1
                return value
            scheduler=m.Scheduler(config(rewrite_experience_concurrency=concurrency),3)
            out=m.rewrite_inventory(source,jd,'Engineer','grounded',scheduler,provider)
            self.assertEqual([r['bullet_id'] for r in out],[str(i) for i in range(13)])
            self.assertTrue(all(r['status']=='candidate' for r in out))
            self.assertLessEqual(maximum,concurrency);outputs.append(out)
        self.assertEqual(outputs[0],outputs[1]);self.assertEqual(outputs[1],outputs[2])

    def test_53_of_54_and_13_of_54_are_partial(self):
        for successful in (13,53):
            source=inv(54,6)
            def provider(stage,payload,timeout):
                value=response(payload)
                for r in value['results']:
                    if int(r['bullet_id'])>=successful:r.update(status='blocked',blocked_reason='facts cannot fit')
                return value
            out=m.rewrite_inventory(source,'Python','role','review_draft',m.Scheduler(config(),6),provider)
            self.assertEqual(sum(r['status']=='candidate' for r in out),successful)
            self.assertEqual(len(out),54)
            self.assertFalse(all(r['status']=='candidate' for r in out))

    def test_retry_only_failed_ids_with_full_experience(self):
        source=inv(4,1); calls=[]
        def provider(stage,payload,timeout):
            calls.append(deepcopy(payload));value=response(payload)
            if len(calls)==1:value['results']=value['results'][:-1]
            return value
        out=m.rewrite_inventory(source,'Python','role','grounded',m.Scheduler(config(),1),provider)
        self.assertEqual(calls[1]['requested_bullet_ids'],['3'])
        self.assertEqual(calls[1]['source_experience'],calls[0]['source_experience'])
        self.assertTrue(all(r['status']=='candidate' for r in out))

    def test_chunk_budget_counts_experience_boundaries(self):
        source=inv(26,2);calls=[]
        def provider(stage,payload,timeout):calls.append(payload);return response(payload)
        scheduler=m.Scheduler(config(),4)
        m.rewrite_inventory(source,'Python','role','grounded',scheduler,provider)
        self.assertEqual([len(p['requested_bullet_ids']) for p in calls].count(12),2)
        self.assertEqual(len(calls),4);self.assertEqual(scheduler.limit,9)

    def test_duplicate_cross_experience_hash_and_fact_guards(self):
        source=inv(2,1);group=source['groups'][0];requested=source['bullets']
        payload=dict(experience_id='0',requested_bullets=requested)
        for mutate in (lambda r:r.update(experience_id='other'),
                       lambda r:r['results'].append(deepcopy(r['results'][0])),
                       lambda r:r['results'][0].update(bullet_id='foreign')):
            value=response(payload);mutate(value)
            with self.assertRaises(ValueError):m.validate_response(value,group,requested)
        for text in (requested[0]['source_text'],requested[0]['source_text'].replace('Built','Made'),
                     'Built Python API services with 999 customers'):
            value=response(payload);value['results'][0]['rewritten_text']=text
            self.assertEqual(m.validate_response(value,group,requested)['0']['status'],'blocked')
        value=response(payload);value['results'][0]['source_text_hash']='stale'
        self.assertEqual(m.validate_response(value,group,requested)['0']['status'],'blocked')

    def test_cost_context_and_global_request_limits(self):
        for changes in ({'rewrite_max_cost_usd':.000001},{'rewrite_context_tokens':50},{'rewrite_max_requests':1}):
            scheduler=m.Scheduler(config(**changes),1);called=[]
            def provider(stage,payload,timeout):called.append(payload);return response(payload)
            out=m.rewrite_inventory(inv(4,2),'Python','role','grounded',scheduler,provider)
            self.assertLessEqual(len(called),1)
            self.assertTrue(any(r['status']=='blocked' for r in out))
            self.assertLessEqual(scheduler.reserved,scheduler.options['rewrite_max_cost_usd'])

    def test_auth_cancels_queued_and_late_responses(self):
        class AuthenticationError(Exception):status_code=401
        scheduler=m.Scheduler(config(rewrite_experience_concurrency=1),3)
        def provider(*args):raise AuthenticationError('secret must not be persisted')
        out=m.rewrite_inventory(inv(6,3),'Python','role','grounded',scheduler,provider)
        self.assertEqual(len(scheduler.events),1);self.assertTrue(scheduler.cancel.is_set())
        self.assertNotIn('secret',json.dumps(out))
        scheduler=m.Scheduler(config(),1)
        def late(stage,payload,timeout):scheduler.cancel.set();return response(payload)
        out=m.rewrite_inventory(inv(2,1),'Python','role','grounded',scheduler,late)
        self.assertTrue(scheduler.events[0]['late_response_ignored'])
        self.assertTrue(all(r['status']!='candidate' for r in out))

    def test_rate_limit_retry_after_shared_cooldown(self):
        class RateLimitError(Exception):
            status_code=429
            response=SimpleNamespace(headers={'retry-after':'5'})
        scheduler=m.Scheduler(config(rewrite_run_timeout=.2),1);calls=[]
        def provider(stage,payload,timeout):calls.append(payload);raise RateLimitError()
        m.rewrite_inventory(inv(3,1),'Python','role','grounded',scheduler,provider)
        self.assertEqual(len(calls),1);self.assertGreater(scheduler.cooldown,time.monotonic())

    def test_partial_experience_failure_preserves_sibling(self):
        def provider(stage,payload,timeout):
            if payload['experience_id']=='1':raise TimeoutError('timeout')
            return response(payload)
        out=m.rewrite_inventory(inv(4,2),'Python','role','grounded',m.Scheduler(config(),2),provider)
        self.assertEqual([r['status'] for r in out],['candidate','blocked','candidate','blocked'])

    def test_one_line_no_minimum_width_and_real_bounds(self):
        source=[dict(paragraph_index=0,kind='bullet',matched=True,line_count=2,lines=[dict(page=1,width=400),dict(page=1,width=400)],width=800)]
        output=[dict(paragraph_index=0,kind='bullet',matched=True,line_count=1,lines=[dict(page=1,width=100,x0=50,x1=150)],width=100)]
        inventory={'bullets':[bullet(0)]}
        self.assertTrue(m.strict_density(source,output,inventory)['passed'])
        self.assertFalse(m.compare_density(source,output)['passed'])
        output[0]['lines'][0]['x1']=600
        self.assertFalse(m.strict_density(source,output,inventory)['passed'])
        output[0]['line_count']=2
        self.assertFalse(m.strict_density(source,output,inventory)['passed'])

    def test_html_escaping(self):
        with TemporaryDirectory() as tmp:
            root=Path(tmp);(root/'exact-jd.txt').write_text('<script>JD</script>')
            m._comparison(root,dict(status='partial',strict_attachment_eligible=False,run_id='test'),[])
            value=(root/'comparison.html').read_text()
            self.assertNotIn('<script>JD',value);self.assertIn('&lt;script&gt;JD',value)


if __name__=='__main__':unittest.main()


class IntegrationContractTest(unittest.TestCase):
    def test_per_call_false_overrides_saved_true_without_mutation(self):
        import modules.ai.resume_optimizer as optimizer
        saved={'rewrite_all_bullets':True,'optimization_mode':'skills'}
        # Legacy run can fail for missing input; it must not enter strict mode.
        with patch.object(optimizer,'load_runtime_config',return_value=deepcopy(saved)), \
             patch.object(m,'run_all_bullets') as strict, patch.object(optimizer,'load_dotenv'), \
             TemporaryDirectory() as tmp:
            optimizer.run_resume_optimization('JD',str(Path(tmp)/'missing.docx'),'id',rewrite_all_bullets=False)
            strict.assert_not_called()
        self.assertTrue(saved['rewrite_all_bullets'])

    def test_per_call_true_full_context_and_mode_override(self):
        import modules.ai.resume_optimizer as optimizer
        with patch.object(optimizer,'load_runtime_config',return_value={'rewrite_all_bullets':False,'optimization_mode':'skills'}), \
             patch.object(m,'run_all_bullets',return_value=None) as strict:
            optimizer.run_resume_optimization('FULL JD','source.docx','id',optimization_mode='bullets',
                                              rewrite_all_bullets=True,rewrite_experience_concurrency=4)
        self.assertEqual(strict.call_args.args[0],'FULL JD')
        self.assertEqual(strict.call_args.args[4]['optimization_mode'],'bullets')
        self.assertEqual(strict.call_args.args[4]['rewrite_experience_concurrency'],4)

    def test_source_technical_transfer_and_metric_subject_rejected(self):
        self.assertIn('unsupported',m.source_claim_issue('Built Python services','Built Kubernetes services'))
        self.assertIn('metric subject',m.source_claim_issue('Cut latency by 50%','Raised reliability by 50%'))

    def test_layout_retry_reuses_siblings_and_two_attempt_budget(self):
        source=inv(3,1);calls=[];scheduler=m.Scheduler(config(),1)
        def provider(stage,payload,timeout):calls.append(payload);return response(payload)
        first=m.rewrite_inventory(source,'Python','role','grounded',scheduler,provider)
        first[1].update(status='blocked',reason='final PDF overflow')
        second=m.rewrite_inventory(source,'Python','role','grounded',scheduler,provider,initial_records=first)
        self.assertEqual(calls[1]['requested_bullet_ids'],['1'])
        second[1].update(status='blocked',reason='overflow persists')
        third=m.rewrite_inventory(source,'Python','role','grounded',scheduler,provider,initial_records=second)
        self.assertEqual(len(calls),2);self.assertEqual(third[1]['status'],'blocked')

    def test_config_flags_persist_and_false_survives_merge(self):
        from modules import config_loader as loader
        with TemporaryDirectory() as tmp:
            root=Path(tmp);defaults=root/'defaults.json';runtime=root/'runtime/config.json'
            defaults.write_text(json.dumps(config(rewrite_all_bullets=False)))
            with patch.object(loader,'APP_BUNDLE_CONFIG_FILE',defaults),patch.object(loader,'RUNTIME_CONFIG_DIR',runtime.parent),patch.object(loader,'RUNTIME_CONFIG_FILE',runtime):
                loader.save_runtime_config({'rewrite_all_bullets':True,'rewrite_experience_concurrency':4})
                self.assertTrue(loader.load_runtime_config()['rewrite_all_bullets'])
                loader.save_runtime_config({'rewrite_all_bullets':False})
                self.assertFalse(loader.load_runtime_config()['rewrite_all_bullets'])
                self.assertEqual(loader.load_runtime_config()['rewrite_experience_concurrency'],4)

    def test_strict_upload_refuses_source_fallback(self):
        from modules.resume_upload import upload_verified_resume
        from unittest.mock import Mock
        modal=Mock();verify=Mock(return_value=True)
        result=upload_verified_resume(modal,'source.docx','JD','id',optimize=Mock(return_value=None),verify_source=verify,rewrite_all_bullets=True)
        self.assertEqual(result,(False,'Strict rewrite incomplete'));verify.assert_not_called();modal.find_element.assert_not_called()

    def test_strict_upload_saved_flag_refuses_source_fallback(self):
        from modules.resume_upload import upload_verified_resume
        from modules import config_loader
        from unittest.mock import Mock
        with patch.object(config_loader,'peek_runtime_config',return_value={'rewrite_all_bullets':True}):
            result=upload_verified_resume(Mock(),'source.docx','JD','id',optimize=Mock(return_value=None))
        self.assertFalse(result[0])


class InventoryContractTest(unittest.TestCase):
    def test_explicit_styled_projects_academic_duplicate_ids_and_skills_exclusion(self):
        with TemporaryDirectory() as tmp:
            path=Path(tmp)/'resume.docx';doc=Document()
            doc.add_heading('SKILLS',1);doc.add_paragraph('Python, AWS')
            doc.add_heading('PROJECTS',1);doc.add_paragraph('Example project')
            doc.add_paragraph('Built Python service',style='List Bullet')
            doc.add_paragraph('• Built Python service')
            doc.add_heading('ACADEMIC DISTINCTIONS',1)
            doc.add_paragraph('Earned academic distinction',style='List Paragraph');doc.save(path)
            measurements=[dict(paragraph_index=i,matched=True,line_count=1,lines=[dict(x0=72,x1=250,width=178,page=1)]) for i in (4,5,7)]
            with patch.object(m,'paragraph_density',return_value=measurements):
                result=m.inventory(path,ResumeOptimizer.__new__(ResumeOptimizer),'unused.pdf')
            self.assertFalse(result['errors']);self.assertEqual(len(result['bullets']),2)
            self.assertEqual(len({b['bullet_id'] for b in result['bullets']}),2)
            self.assertEqual(len(result['groups']),1)
            self.assertEqual(len(result['preserved_bullets']),1)
            self.assertEqual(result['preserved_bullets'][0]['source_text'],'Earned academic distinction')

    def test_unmapped_rendered_source_is_blocking(self):
        with TemporaryDirectory() as tmp:
            path=Path(tmp)/'resume.docx';doc=Document();doc.add_paragraph('Unmapped source',style='List Paragraph');doc.save(path)
            with patch.object(m,'paragraph_density',return_value=[]):
                result=m.inventory(path,ResumeOptimizer.__new__(ResumeOptimizer),'unused.pdf')
            self.assertGreaterEqual(len(result['errors']),2)


class AttachmentBindingTest(unittest.TestCase):
    def test_strict_attachment_rejects_unbound_pdf_even_when_optimizer_returns_path(self):
        from modules.resume_upload import upload_verified_resume
        from unittest.mock import Mock
        with TemporaryDirectory() as tmp:
            pdf=Path(tmp)/'partial.pdf';pdf.write_bytes(b'partial draft')
            modal=Mock()
            result=upload_verified_resume(modal,'source.docx','JD','id',optimize=Mock(return_value=str(pdf)),rewrite_all_bullets=True)
            self.assertFalse(result[0]);modal.find_element.assert_not_called()


class NumericAttributionTest(unittest.TestCase):
    def test_metric_values_cannot_swap_even_with_identical_number_set(self):
        self.assertIn('attribution',m.source_claim_issue('Raised engagement by 30% and efficiency by 15%',
            'Raised engagement by 15% and efficiency by 30%'))


class ExpectedArtifactBindingTest(unittest.TestCase):
    def test_current_job_jd_and_source_are_required_for_attachment(self):
        with TemporaryDirectory() as tmp:
            root=Path(tmp);pdf=root/'optimized.pdf';source=root/'source.docx';source.write_bytes(b'source')
            manifest=dict(source_inventory_count=1,jd_sha256=m.sha('exact JD'),role_id='job',source_docx={'sha256':m.sha(b'source')})
            report=dict(status='complete',strict_attachment_eligible=True,rewrite_completed=True,ai_execution_kind='provider',
                tailoring_contract='all-bullets-skills-titles-v1',source_retention_contract=m.SOURCE_RETENTION_VERSION,
                skills_optimization={'completed':True},skills_grounding={'policy':'grounded'},
                effective_config={'content_policy':'grounded','source_checks_enabled':True},
                ai_activity={'live_provider_response_observed':True},format_verification={'passed':True,'checks':{'academic_distinctions_preserved':True}},
                coverage={'accepted':1,'inventoried':1,'failed':0,'retained_source':0},output_path=str(pdf),source_sha256=m.sha(b'source'))
            (root/'match-report.json').write_text(json.dumps(report));(root/'rewrite-coverage.json').write_text(json.dumps({'bullets':[{'bullet_id':'id','status':'accepted',
                'source_text':'Built Python services','rewritten_text':'Python services were developed'}]}))
            with patch.object(m,'verify_binding',return_value=manifest):
                self.assertTrue(m.verify_strict_attachment(pdf,expected_jd='exact JD',expected_job_id='job',expected_source_docx=source))
                self.assertFalse(m.verify_strict_attachment(pdf,expected_jd='other JD'))
                self.assertFalse(m.verify_strict_attachment(pdf,expected_job_id='other job'))
                unchecked=dict(report,format_verification={'passed':True})
                (root/'match-report.json').write_text(json.dumps(unchecked))
                self.assertFalse(m.verify_strict_attachment(pdf),'Academic Distinctions preservation must be verified')
                (root/'match-report.json').write_text(json.dumps(report))
                legacy=dict(report);legacy.pop('tailoring_contract')
                (root/'match-report.json').write_text(json.dumps(legacy))
                self.assertFalse(m.verify_strict_attachment(pdf),'old all-bullets-only artifacts lack the combined contract')
                (root/'match-report.json').write_text(json.dumps(report))
                legacy=dict(report);legacy.pop('source_retention_contract')
                (root/'match-report.json').write_text(json.dumps(legacy))
                self.assertFalse(m.verify_strict_attachment(pdf),'old acceptance missed material source omissions')
                (root/'match-report.json').write_text(json.dumps(report))
                source.write_bytes(b'changed source')
                self.assertFalse(m.verify_strict_attachment(pdf,expected_source_docx=source))


class ProviderBoundaryRepairTest(unittest.TestCase):
    def test_null_blocked_text_does_not_discard_valid_sibling(self):
        source=inv(2,1);payload={'experience_id':'0','requested_bullets':source['bullets']};value=response(payload)
        value['results'][1].update(status='blocked',rewritten_text=None,blocked_reason='Facts cannot fit',change_reason=None)
        result=m.validate_response(value,source['groups'][0],source['bullets'])
        self.assertEqual(result['0']['status'],'candidate');self.assertEqual(result['1']['status'],'blocked')
        self.assertEqual(result['1']['rewritten_text'],source['bullets'][1]['source_text'])
        self.assertEqual(result['1']['lexical_difference'],[])

    def test_malformed_diagnostics_remain_blocked_without_batch_exception(self):
        source=inv(2,1);payload={'experience_id':'0','requested_bullets':source['bullets']}
        for text in (None,42,[],{},True):
            value=response(payload);value['results'][1]['rewritten_text']=text
            result=m.validate_response(value,source['groups'][0],source['bullets'])
            self.assertEqual(result['0']['status'],'candidate');self.assertEqual(result['1']['status'],'blocked')
        value=response(payload);value['results'][1]['matched_jd_spans']=None
        self.assertEqual(m.validate_response(value,source['groups'][0],source['bullets'])['1']['status'],'blocked')

    def test_known_usage_releases_unused_capacity_and_archives_body(self):
        scheduler=m.Scheduler(config(rewrite_max_cost_usd=.10),2)
        body={'keywords':['Python']}
        reply=SimpleNamespace(id='offline-fixture-response',usage=SimpleNamespace(model_dump=lambda:{'prompt_tokens':100,'completion_tokens':50}),choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps(body)))])
        for _ in range(4):self.assertEqual(scheduler.call(lambda *args:reply,'keywords',{'job_description':'Python'}),body)
        self.assertEqual(len(scheduler.events),4);self.assertEqual(len(scheduler.responses),4)
        self.assertAlmostEqual(scheduler.known_usage_cost,.003)
        self.assertLess(scheduler.reserved,.10);self.assertGreater(scheduler.gross_reserved,.10)
        self.assertEqual(json.loads(scheduler.responses[0]['content']),body)


class FairFirstPassTest(unittest.TestCase):
    def test_all_experiences_first_attempt_before_retry_with_complete_original_facts(self):
        source=inv(8,4);calls=[]
        def provider(stage,payload,timeout):
            calls.append(deepcopy(payload));value=response(payload)
            if sum(p['experience_id']==payload['experience_id'] for p in calls)==1:
                value['results'][0]['rewritten_text']=None
            return value
        scheduler=m.Scheduler(config(),4)
        records=m.rewrite_inventory(source,'Python complete JD','role','grounded',scheduler,provider)
        self.assertEqual(len({p['experience_id'] for p in calls[:4]}),4)
        self.assertEqual(len(calls),8)
        self.assertTrue(all(r['status']=='candidate' for r in records))
        for payload in calls:
            group=next(g for g in source['groups'] if g['experience_id']==payload['experience_id'])
            self.assertEqual(payload['job_description'],'Python complete JD')
            self.assertEqual(len(payload['source_experience']['bullets']),len(group['bullets']))
            self.assertEqual([b['source_text'] for b in payload['source_experience']['bullets']],[b['source_text'] for b in group['bullets']])
            for b in payload['source_experience']['bullets']:self.assertNotIn('source_rendered',b)
