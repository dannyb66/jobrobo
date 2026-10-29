"""Combined Skills/all-section bullets/titles, with real rendering and no network."""
import json
import socket
import time
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from docx import Document
from docx.shared import Inches, Pt
from modules import resume_all_bullets as m
from modules.ai.resume_optimizer import ResumeOptimizer


class FullTailoringTest(unittest.TestCase):
    def test_repeated_complete_word_cannot_hide_truncated_name(self):
        source='Received DVC award; recipient of Institute Awards'
        self.assertEqual(m.missing_source_terms(source,'Received DVC award; recipient of Institute Awa'),[])
        self.assertIn('truncated',m.source_scope_issue(source,'Received DVC award; recipient of Institute Awa'))
        self.assertIsNone(m.source_scope_issue(source,'Recipient of Institute Awards; received DVC award'))

    def test_passive_drafting_reference_keeps_intact_fact_list(self):
        source='Designed auditable agent control plane: registry, discovery and scope allowlists'
        reference=m.source_passive_reference(source)
        self.assertEqual(reference,'Auditable agent control plane designed: registry, discovery and scope allowlists')
        self.assertEqual(m.missing_source_terms(source,reference),[])
        self.assertTrue(m.substantive(source,reference))
        self.assertIsNone(m.source_passive_reference('Own agent architecture: memory and observability'))

    def test_optional_tokenizer_changes_rate_estimate_not_cost_or_context_caps(self):
        from test_resume_all_bullets import config
        class Encoder:
            def encode(self,value,**kwargs):return [0]*(len(value)//4+1)
        payload={'job_description':'Python '*3000}
        options=config(rewrite_token_estimation='tiktoken',rewrite_tokens_per_minute=120000,
            rewrite_requests_per_minute=60)
        with patch.object(m,'_token_encoder',return_value=Encoder()),patch.object(m.time,'monotonic',return_value=10):
            scheduler=m.Scheduler(options,1)
            scheduler.call(lambda *args:{'keywords':['Python']},'keywords',payload)
        event=scheduler.events[0]
        self.assertLess(event['rate_input_tokens_estimate'],event['estimated_input_tokens'])
        self.assertEqual(scheduler.next_request,10+max(1,60*(event['rate_input_tokens_estimate']+6000)/120000))
        expected=(event['estimated_input_tokens']*2.5+6000*10)/1e6
        self.assertEqual(scheduler.reserved,expected)
        self.assertEqual(m.DEFAULTS['rewrite_token_estimation'],'byte_upper_bound')
        self.assertEqual(options['rewrite_context_tokens'],m.DEFAULTS['rewrite_context_tokens'])
        with self.assertRaises(ValueError):m.validate_options(config(rewrite_token_estimation='unknown'),'full')

    def test_underestimated_provider_token_usage_stops_new_dispatch(self):
        from types import SimpleNamespace
        from test_resume_all_bullets import config
        scheduler=m.Scheduler(config(),1)
        response=SimpleNamespace(id='provider-test',usage=SimpleNamespace(model_dump=lambda:{'prompt_tokens':20,'completion_tokens':0}),
            choices=[SimpleNamespace(message=SimpleNamespace(content='{"keywords":["Python"]}'))])
        with patch.object(m,'rate_input_estimate',return_value=10):
            with self.assertRaises(m.Stopped):scheduler.call(lambda *args:response,'keywords',{'job_description':'Python'})
        self.assertEqual(scheduler.reason,'reported input tokens exceeded rate estimate; new dispatch cancelled')
        with self.assertRaises(m.Stopped):scheduler.call(lambda *args:response,'keywords',{'job_description':'Python'})
        self.assertEqual(len(scheduler.events),1)

    def test_source_creation_action_and_timeframe_subject_cannot_disappear(self):
        self.assertIn('action omitted',m.source_scope_issue('Designed auditable agent control plane',
            'Auditable agent control plane'))
        self.assertIn('action omitted',m.source_scope_issue('Engineered pipeline; onboarded clients',
            'Onboarded clients with pipeline'))
        self.assertIn('action omitted',m.source_scope_issue('Engineered Python pipeline; enabled customer workflows',
            'Enabled customer workflows with Python pipeline'))
        self.assertIsNone(m.source_scope_issue('Designed auditable agent control plane',
            'Auditable agent control plane was designed'))
        self.assertIn('timeframe subject',m.source_scope_issue(
            'Deployed architecture; deliver prototype in 3 months',
            'For prototype, deployed architecture in 3 months'))
        self.assertIsNone(m.source_scope_issue('Deployed architecture; deliver prototype in 3 months',
            'Deliver prototype in 3 months; deployed architecture'))
        self.assertIsNone(m.source_scope_issue('Deliver prototype in 3 months',
            'Prototype delivered in 3 months'))

    def test_length_truncation_cannot_drop_closing_delimiters(self):
        source='Built backend (Python, AWS) with responsive frontends (React)'
        self.assertIn('unbalanced',m.source_scope_issue(source,
            'With responsive frontends (React), built backend (Python, AWS'))
        self.assertIsNone(m.source_scope_issue(source,
            'With responsive frontends (React), built backend (Python, AWS)'))

    def test_smaller_batches_keep_full_experience_and_recalculate_request_bound(self):
        from test_resume_all_bullets import inv,response,config
        source=inv(8,1);seen=[]
        def provider(stage,payload,timeout):seen.append(payload);return response(payload)
        options=config(rewrite_bullets_per_request=3)
        scheduler=m.Scheduler(options,3)
        values=m.rewrite_inventory(source,'Python','Engineer','grounded',scheduler,provider)
        self.assertEqual([len(p['requested_bullets']) for p in seen],[3,3,2])
        self.assertTrue(all(len(p['source_experience']['bullets'])==8 for p in seen))
        self.assertEqual(scheduler.limit,7)
        self.assertTrue(all(v['status']=='candidate' for v in values))
        for bad in (0,13,2.5,True):
            with self.assertRaises(ValueError):m.validate_options(config(rewrite_bullets_per_request=bad),'full')

    def test_material_source_omissions_cannot_pass_using_sibling_context(self):
        cases=[
            ('Own agent architecture: plan mode, memory, observability and per-query cost cap',
             'Own agent architecture: plan mode, observability and cost cap'),
            ('Sole developer of self-hosted inference with vLLM',
             'Developer of inference with vLLM'),
            ('Built Responses API with no public ingress',
             'Built Response API with public ingress'),
            ('Co-founder with a mission to democratize high-school education, aiming to reach 100M students',
             'Co-founder democratized education for 100M students'),
            ('Built runbooks for incident response, rollback and feature rollout',
             'Built runbooks for feature management')]
        for source,rewrite in cases:
            with self.subTest(source=source):
                self.assertIsNotNone(m.source_claim_issue(source,rewrite,technology_context=source))
        self.assertIsNone(m.source_claim_issue('Built Python API services for customer workflows',
            'For customer workflows, Python API services were built'))

    def test_negation_and_intent_cannot_be_relocated_to_fake_retention(self):
        self.assertIn('negation',m.source_scope_issue('Built proxy with no public ingress',
            'Built public ingress proxy with no other issues'))
        self.assertIn('intended',m.source_scope_issue('Mission to democratize education, aiming to reach students',
            'Democratize education and reach students; mission to improve, aiming to grow'))
        self.assertIn('transition',m.source_scope_issue('Built 0→1 platform','Built 1→0 platform'))
        self.assertIsNone(m.source_scope_issue('Built proxy with no public ingress; fixed errors',
            'Fixed errors; built proxy with no public ingress'))
        self.assertIn('quantity',m.source_scope_issue('Supported 20 tools and 40 skills',
            'Supported 40 tools and 20 skills'))
        self.assertIn('timeframe',m.source_scope_issue('Built MVP within 2 months',
            'Built MVP after 2 months'))
        self.assertIn('direction',m.source_scope_issue('Reduced latency by 50%',
            'Reduced latency to 50%'))
        self.assertIsNone(m.source_scope_issue('Built 50% of requests',
            'Built requests comprising 50%'))

    def test_technical_identity_is_not_singularized(self):
        terms=m.source_retention_terms('AWS ECS Redis Kubernetes SaaS Responses API C++ C#')
        for term in ('aws','ecs','redis','kubernetes','saas','responses','api','c++','c#'):
            self.assertIn(term,terms)
        self.assertIn('responses',m.missing_source_terms('Responses API','Response API'))
        self.assertIn('c++',m.missing_source_terms('C++','C'))
        self.assertIn('kubernetes',m.missing_source_terms('kubernetes','kubernete'))
        self.assertEqual(m.missing_source_terms('Built services','Services were built'),[])
        self.assertEqual(m.missing_source_terms('Managed Services','Managed services'),[])
        self.assertEqual(m.missing_source_terms('Built customer workflows','Customer workflow services were built'),[])
        self.assertEqual(m.missing_source_terms('Increasing efficiency, reducing manual effort',
            'Reduced manual effort; increased efficiency'),[])
        self.assertIn('aiming',m.missing_source_terms('Aiming to reach students','Achieved student reach'))

    def test_source_rotation_is_only_a_drafting_aid_and_keeps_all_terms(self):
        source='Own agent architecture: plan mode, memory and observability'
        reference=m.source_reordering_reference(source)
        self.assertEqual(reference,'plan mode, memory and observability: Own agent architecture')
        self.assertEqual(m.missing_source_terms(source,reference),[])
        self.assertTrue(m.substantive(source,reference))
        self.assertIsNone(m.source_reordering_reference('Built a single undivided clause'))
        self.assertEqual(m.source_reordering_reference('Built (Python, AWS) services; reduced toil'),
                         'reduced toil; Built (Python, AWS) services')

    def test_whole_clause_reorder_with_verb_edit_is_substantive(self):
        source='Built Python API services for customers; enabled reliable workflows'
        rewrite='Enabled reliable workflows; designed Python API services for customers'
        self.assertTrue(m.substantive(source,rewrite))
        self.assertFalse(m.substantive(source,source.replace('Built','Designed')))
        self.assertFalse(m.substantive(source,source.replace(';',',')))
        self.assertTrue(m.substantive(
            'Designed agent control plane with skills and allowlists across tools',
            'Across tools, designed agent control plane with skill and allowlists'))

    def test_metadata_cannot_substitute_for_actual_factual_checks(self):
        from test_resume_all_bullets import inv,response,config
        source=inv(1,1);group=source['groups'][0]
        def provider(stage,payload,timeout):
            raw=response(payload)
            raw['results'][0].update(preserved_fact_anchors=[],matched_jd_spans=['invented quotation'])
            return raw
        values=m.rewrite_inventory(source,'Python','Engineer','grounded',m.Scheduler(config(),1),provider)
        self.assertEqual(values[0]['status'],'candidate')
        self.assertEqual(values[0]['matched_jd_spans'],[])
        self.assertEqual(values[0]['discarded_jd_spans'],['invented quotation'])
        bad=response({'experience_id':group['experience_id'],'requested_bullets':source['bullets']})
        bad['results'][0]['rewritten_text']='Invented 999 Python API customers'
        self.assertEqual(m.validate_response(bad,group,source['bullets'])['0']['status'],'blocked')

    def test_selects_only_factual_fitting_alternative(self):
        from test_resume_all_bullets import inv,response
        source=inv(1,1);group=source['groups'][0]
        raw=response({'experience_id':group['experience_id'],'requested_bullets':source['bullets']})
        good=raw['results'][0]['rewritten_text']
        raw['results'][0].update(rewritten_text='Invented 999 customers',alternative_texts=[good,good+' extra words'*15])
        result=m.validate_response(raw,group,source['bullets'],fit_check=lambda r:'too wide' if len(r['rewritten_text'])>100 else None)['0']
        self.assertEqual(result['status'],'candidate');self.assertEqual(result['rewritten_text'],good)
        self.assertEqual([v['status'] for v in result['candidate_evaluations']],['blocked','candidate','blocked'])

    def test_reasoning_model_parameters_and_legacy_compatibility(self):
        legacy=m.completion_options(m.DEFAULTS)
        self.assertIn('max_tokens',legacy);self.assertNotIn('reasoning_effort',legacy)
        modern=m.completion_options(dict(m.DEFAULTS,rewrite_model='gpt-5.4'))
        self.assertNotIn('max_tokens',modern)
        self.assertEqual(modern['max_completion_tokens'],m.DEFAULTS['rewrite_max_output_tokens'])
        self.assertEqual(modern['extra_body']['reasoning_effort'],'low')

    def test_role_heading_exits_ignored_company_section(self):
        from modules.resume_matching import jd_requirements
        jd='About Us\nWe use Ruby at Example.\nWhat We Are Looking For\nRequired: Python and AWS.\nBenefits\nKubernetes training.'
        rows,_=jd_requirements(jd,['Ruby','Python','AWS','Kubernetes'])
        self.assertEqual({term for row in rows for term in row['terms']},{'Python','AWS'})

    def test_same_experience_technology_does_not_transfer_metrics(self):
        from test_resume_all_bullets import inv, response
        source=inv(2,1);group=source['groups'][0]
        group['bullets'][1]['source_text']='Built Kubernetes services with 20% lower latency'
        payload={'experience_id':group['experience_id'],'requested_bullets':[source['bullets'][0]]}
        value=response(payload)
        value['results'][0]['rewritten_text']='Built customer workflow services number 0 through Python API development on Kubernetes'
        self.assertEqual(m.validate_response(value,group,payload['requested_bullets'])['0']['status'],'candidate')
        foreign=inv(1,1)['groups'][0]
        self.assertEqual(m.validate_response(value,foreign,payload['requested_bullets'])['0']['status'],'blocked')
        value['results'][0]['rewritten_text']+=' with 20% lower latency'
        self.assertEqual(m.validate_response(value,group,payload['requested_bullets'])['0']['status'],'blocked')

    def test_request_schema_binds_ids_to_exact_owned_payload(self):
        payload={'experience_id':'exp-a','requested_bullets':[
            {'bullet_id':'owned-a','source_text_hash':'hash-a'},
            {'bullet_id':'owned-b','source_text_hash':'hash-b'}]}
        schema=m.provider_response_format('bullets',payload)['json_schema']['schema']['properties']
        self.assertEqual(schema['experience_id']['enum'],['exp-a'])
        fields=schema['results']['items']['properties']
        self.assertEqual(fields['bullet_id']['enum'],['owned-a','owned-b'])
        self.assertEqual(fields['source_text_hash']['enum'],['hash-a','hash-b'])

    def test_schema_bounds_drafting_without_confusing_char_count_with_glyph_fit(self):
        from test_resume_all_bullets import inv
        source=inv(2,1);group=source['groups'][0]
        payload=m.compact_payload(source,'Python','Engineer','grounded',group,group['bullets'])
        variants=m.provider_response_format('bullets',payload)['json_schema']['schema']['properties']['results']['items']['anyOf']
        for bullet,variant in zip(source['bullets'],variants):
            fields=variant['properties']
            self.assertEqual(fields['bullet_id']['enum'],[bullet['bullet_id']])
            self.assertEqual(fields['source_text_hash']['enum'],[bullet['source_text_hash']])
            limit=fields['rewritten_text']['maxLength']
            self.assertGreater(limit,len(bullet['source_text']))
            self.assertLessEqual(limit,2*max(len(bullet['source_text']),bullet['soft_character_estimate']))
            self.assertEqual(fields['alternative_texts']['items']['maxLength'],limit)

    def test_private_use_padding_cannot_count_as_valid_resume_text(self):
        from test_resume_all_bullets import inv,response
        source=inv(1,1);group=source['groups'][0]
        raw=response({'experience_id':group['experience_id'],'requested_bullets':source['bullets']})
        raw['results'][0]['rewritten_text']+='\U0010fc00'
        result=m.validate_response(raw,group,source['bullets'])['0']
        self.assertEqual(result['status'],'blocked')
        self.assertIn('private-use',result['reason'])

    def test_skills_all_sections_and_titles_render_together(self):
        self._render_case()

    def test_explicit_cancellation_still_blocks_after_local_render(self):
        self._render_case(cancel_during_render=True)

    def _render_case(self,cancel_during_render=False):
        with TemporaryDirectory() as tmp:
            root=Path(tmp);source=root/'original.docx';doc=Document()
            for sec in doc.sections:
                sec.top_margin=sec.bottom_margin=Inches(.5)
            doc.styles['Normal'].font.name='Arial';doc.styles['Normal'].font.size=Pt(10)
            doc.add_heading('SKILLS',level=1);doc.add_paragraph('AWS, Kubernetes, Python')
            replacements=[
                ('Built Python API services for customer workflows',
                 'Delivered customer workflows through Python API service development'),
                ('Designed AWS services for event processing',
                 'Implemented event processing through AWS service design'),
                ('Studied distributed systems through Python experiments',
                 'Explored Python experiments to study distributed systems')]
            for index,section in enumerate(('PROFESSIONAL EXPERIENCE','PROJECTS','EDUCATION')):
                doc.add_heading(section,level=1);doc.add_paragraph('Example company' if index==0 else 'Example project')
                if index==0:doc.add_paragraph('Software Engineer\tJanuary 2020 - Present')
                doc.add_paragraph(replacements[index][0],style='List Bullet')
            doc.save(source);original=source.read_bytes()
            optimizer=ResumeOptimizer.__new__(ResumeOptimizer)
            self.assertTrue(optimizer.convert_docx_to_pdf(str(source),str(source.with_suffix('.pdf'))))
            def provider(stage,payload,timeout):
                if stage=='keywords':return {'keywords':['Python','Python API','AWS','Rust']}
                return {'experience_id':payload['experience_id'],'results':[
                    dict(bullet_id=b['bullet_id'],source_text_hash=b['source_text_hash'],status='rewritten',
                         rewritten_text=replacements[b['ordinal']][1],matched_jd_spans=[],
                         preserved_fact_anchors=b['fact_anchors'],change_reason='Faithful action and structure rewrite',blocked_reason=None)
                    for b in payload['requested_bullets']]}
            config=dict(m.DEFAULTS,rewrite_all_bullets=True,expected_page_count=1,replace_job_title=True,
                rewrite_requests_per_minute=1000000,rewrite_tokens_per_minute=1000000000)
            out=root/'review'
            schedulers=[]
            class CapturedScheduler(m.Scheduler):
                def __init__(self,*args,**kwargs):
                    super().__init__(*args,**kwargs);schedulers.append(self)
            original_convert=ResumeOptimizer.convert_docx_to_pdf
            def render_after_provider_deadline(instance,*args):
                result=original_convert(instance,*args)
                for scheduler in schedulers:
                    scheduler.deadline=time.monotonic()-1
                    if cancel_during_render:scheduler.cancel.set()
                return result
            with patch.object(socket.socket,'connect',side_effect=AssertionError('network forbidden')), \
                 patch.object(m,'Scheduler',CapturedScheduler), \
                 patch.object(ResumeOptimizer,'convert_docx_to_pdf',render_after_provider_deadline):
                result=m.run_all_bullets('Required: Python API development. Python services. AWS. Rust.',source,
                    'synthetic-role','Principal Engineer',config,policy='review_draft',execution_kind='offline_fixture',provider=provider,output_root=out)
            self.assertIsNone(result,'fixtures must never become live attachments')
            report=json.loads((out/'match-report.json').read_text())
            # One keyword call plus first attempts and bounded fullness revisions.
            self.assertEqual(report['ai_activity']['request_attempts'],7)
            self.assertLessEqual(report['ai_activity']['request_attempts'],report['request_budget']['effective_max_requests'])
            if cancel_during_render:
                self.assertEqual(report['status'],'partial')
                self.assertFalse(report['strict_attachment_eligible'])
                return
            self.assertEqual(report['status'],'complete',str(report.get('errors'))+str(report.get('format_verification',{}).get('checks')))
            self.assertEqual(report['coverage']['accepted'],3)
            self.assertEqual(report['tailoring_contract'],'all-bullets-skills-titles-v1')
            self.assertTrue(report['skills_optimization']['completed'])
            self.assertTrue(report['skills_optimization']['changed'])
            self.assertEqual(report['skills_grounding']['policy'],'grounded')
            self.assertNotIn('Rust',report['skills_optimization']['output'])
            self.assertEqual(len(report['title_changes']),1)
            self.assertTrue(report['title_changes'][0]['output'].startswith('Principal Engineer\t'))
            final=Document(out/'optimized.docx');text='\n'.join(p.text for p in final.paragraphs)
            for before,after in replacements:
                self.assertIn(after,text);self.assertNotIn(before,text)
            self.assertEqual(source.read_bytes(),original)
            self.assertTrue(report['format_verification']['checks']['exact_source_paragraph_run_format'])
            self.assertIn('Skills and professional title changes',(out/'comparison.html').read_text())
            m.verify_binding(out)


if __name__=='__main__':unittest.main()
