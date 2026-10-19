"""Synthetic draft/title/requirement/fit contracts; no credentials or network."""
import json
import unittest
from copy import deepcopy
from pathlib import Path
from unittest.mock import Mock, patch

from docx import Document
from modules.ai.resume_optimizer import ResumeOptimizer, build_guided_keywords
from modules.resume_matching import jd_requirements, match_analysis
import test_resume_pipeline as pipeline
from test_resume_pipeline import api_response


class DraftHelpersTest(unittest.TestCase):
    def optimizer(self):
        optimizer = ResumeOptimizer.__new__(ResumeOptimizer)
        optimizer.client = Mock()
        return optimizer

    def test_late_existing_jd_skills_survive_capacity_clipping(self):
        optimizer = self.optimizer()
        original = ["Old Tool " + str(i) for i in range(60)] + ["AWS", "Docker", "Terraform"]
        selected = optimizer.optimize_skills(original, [], "Required: AWS, Docker, Terraform", max_count=3)
        self.assertEqual(selected, ["AWS", "Docker", "Terraform"])
        self.assertIn("Old Tool 0", optimizer.last_skills_grounding["removed_source_skills"])

    def test_off_jd_bank_and_model_terms_do_not_enter_keyword_denominator(self):
        bank = {"keywords": [{"keyword": "RAG", "tier": "Core"}, {"keyword": "Price Transparency (TiC)", "tier": "Core"}]}
        result = build_guided_keywords("Build storage systems with AWS", ["AWS", "Scala"], "Engineer", ["RAG", "AWS"], bank, "")
        self.assertEqual(result, ["AWS"])

    def test_related_existing_cloud_tools_rank_without_becoming_jd_requirements(self):
        optimizer = self.optimizer()
        original = ["Old Tool " + str(i) for i in range(60)] + ["AWS", "Docker", "Terraform"]
        self.assertEqual(optimizer.optimize_skills(original, [], "Build cloud infrastructure deployments", max_count=3), ["AWS", "Docker", "Terraform"])
        self.assertEqual(jd_requirements("Build cloud infrastructure deployments", ["AWS", "Docker", "Terraform"])[0], [])

    def test_annotated_source_tool_keeps_context_priority_without_requirement_credit(self):
        optimizer = self.optimizer()
        original = ["Old Tool " + str(i) for i in range(60)] + ["Terraform (IaC)"]
        self.assertEqual(optimizer.optimize_skills(original, [], "Build cloud infrastructure", max_count=1), ["Terraform (IaC)"])
        self.assertEqual(jd_requirements("Build cloud infrastructure", ["Terraform (IaC)"])[0], [])

    def test_opt_in_draft_adds_proposals_but_grounded_default_does_not(self):
        optimizer = self.optimizer()
        self.assertEqual(optimizer.optimize_skills(["Python", "AWS"], ["Scala"], "Scala"), ["Python", "AWS"])
        optimizer.content_policy = "review_draft"
        self.assertIn("Scala", optimizer.optimize_skills(["Python", "AWS"], ["Scala"], "Scala"))
        self.assertEqual(optimizer.last_skills_grounding["proposed_keywords"], ["Scala"])
        self.assertTrue(optimizer.last_skills_grounding["review_required"])

    def test_required_preferred_or_and_title_scores_are_separate_and_not_vendor_weights(self):
        jd = "Required: Python or Java\nPreferred: AWS\nMust have Kubernetes."
        report = match_analysis("Built Python APIs", "Python AWS Kubernetes", jd, ["Python", "Java", "AWS", "Kubernetes"],
                                ["Principal Engineer"], ["Software Engineer"], "Software Engineer")
        self.assertEqual(report["coverage"]["required"], {"total": 2, "source": 50.0, "output": 100.0})
        self.assertEqual(report["coverage"]["preferred"]["source"], 0)
        self.assertEqual(report["title_alignment"]["source"]["best_term_coverage"], 50)
        self.assertEqual(report["title_alignment"]["output"]["best_term_coverage"], 100)
        self.assertIn("not learned or vendor-calibrated", report["weight_basis"])
        self.assertEqual(next(r for r in report["requirements"] if r["terms"] == ["Kubernetes"])["status"], "draft_only_mention")

    def test_repeated_requirement_and_negated_source_do_not_inflate_source_evidence(self):
        requirements, _ = jd_requirements("Must know AWS.\nMust know AWS.", ["AWS"])
        self.assertEqual(len(requirements), 1)
        report = match_analysis("No AWS experience", "AWS", "AWS required", ["AWS"])
        self.assertEqual(report["coverage"]["required"]["source"], 0)
        self.assertTrue(report["requirements"][0]["evidence"]["source"]["negated_only"])

    def test_bare_and_flattened_headings_preserve_required_preferred_and_ignore_benefits(self):
        bare = jd_requirements("REQUIREMENTS\nPython\nNICE TO HAVE\nRust", ["Python", "Rust"])[0]
        self.assertEqual([(r["terms"], r["kind"]) for r in bare], [(["Python"], "required"), (["Rust"], "preferred")])
        flattened = "ABOUT EXAMPLE We build Scala. THE ROLE Build APIs. RESPONSIBILITIES - Use Kubernetes. REQUIREMENTS - Python. NICE TO HAVE - Rust. BENEFITS - ML startups exposure."
        rows = jd_requirements(flattened, ["Scala", "APIs", "Kubernetes", "Python", "Rust", "ML startups"])[0]
        self.assertEqual({r["terms"][0]: r["kind"] for r in rows}, {"APIs": "context", "Kubernetes": "context", "Python": "required", "Rust": "preferred"})

    def test_isnt_required_is_optional_even_inside_required_section(self):
        rows = jd_requirements("REQUIREMENTS Prior ML experience isn't required, though model serving is a plus.", ["ML", "model serving"])[0]
        self.assertTrue(rows)
        self.assertTrue(all(r["kind"] == "preferred" for r in rows))

    def test_invalid_weight_shape_or_nonpositive_values_are_rejected(self):
        for weights in ([], {"required": 0}, {"preferred": float("nan")}, {"unknown": 2}):
            with self.subTest(weights=weights), self.assertRaises(ValueError):
                jd_requirements("Python required", ["Python"], weights)

    def test_fit_budget_accepts_live_length_example_and_retry_preserves_valid_bullets(self):
        optimizer = self.optimizer()
        optimizer.rewrite_length_multiplier = 1.15
        optimizer.rewrite_response_attempts = 2
        original = [{"company": "Example", "job_title": "Engineer", "bullets": ["a"*121, "Used AWS"]}]
        first = {"rewritten_experience": [dict(original[0], bullets=["b"*180, "New AWS"])]}
        second = {"rewritten_experience": [dict(original[0], bullets=["b"*132, "New AWS"])]}
        optimizer.client.chat.completions.create.side_effect = [api_response(json.dumps(first)), api_response(json.dumps(second))]
        result = optimizer._call_openai_to_rewrite_bullets(original, "Engineer", "AWS")
        self.assertEqual(result[0]["bullets"], ["b"*132, "New AWS"])
        self.assertEqual(optimizer.client.chat.completions.create.call_count, 2)
        optimizer._call_openai_to_rewrite_bullets(original, "Engineer", "AWS")
        self.assertEqual(optimizer.client.chat.completions.create.call_count, 2, "layout retry must reuse a validated rewrite")
        self.assertIsNone(optimizer.last_rewrite_diagnostics[-1]["error"])

    def test_final_invalid_bullet_does_not_discard_other_valid_rewrites(self):
        optimizer = self.optimizer()
        source = [{"company": "Example", "job_title": "Engineer", "bullets": ["Used Python", "Used AWS"]}]
        response = api_response(json.dumps({"rewritten_experience": [dict(source[0], bullets=["x"*300, "Ran AWS"])]}))
        result = optimizer._parse_openai_response(response, source, allow_partial=True)
        self.assertEqual(result[0]["bullets"], ["Used Python", "Ran AWS"])
        self.assertIn("Retained individual source bullets", optimizer.last_bullet_rewrite_error)

    def test_preserving_source_bank_sections_reuses_validated_other_entry_rewrites(self):
        optimizer = self.optimizer()
        source = Document()
        source.add_paragraph("PROFESSIONAL EXPERIENCE", style="Heading 1")
        for company, bullet in (("TARGET", "Built source API"), ("OTHER", "Built source tools")):
            source.add_paragraph(company + "\tRemote")
            source.add_paragraph("Engineer\t2020 - Present")
            source.add_paragraph(bullet, style="List Paragraph")
        bank = {"bullets": [{"id": "A1", "text": "Built bank API", "length": 14, "section": "A", "metrics": [], "lead_verb": "Built"}]}
        spec = {"company": "TARGET", "bank": bank, "min_bullets": 1, "max_bullets": 1, "required_sections": ["A"]}
        response = {"rewritten_experience": [{"company": "OTHER", "job_title": "Engineer", "bullets": ["Built Python tools"]}]}
        optimizer.client.chat.completions.create.return_value = api_response(json.dumps(response))
        optimizer.rewrite_experience_bullets(deepcopy(source), "Engineer", "Python API", ["Python"], project_bullet_banks=[spec])
        retry = deepcopy(source)
        optimizer.rewrite_experience_bullets(retry, "Engineer", "Python API", ["Python"], project_bullet_banks=[{**spec, "preserve_source_bullets": True}])
        bullets = [p.text for p in retry.paragraphs if p.style.name == "List Paragraph"]
        self.assertEqual(bullets, ["Built source API", "Built Python tools"])
        self.assertEqual(optimizer.client.chat.completions.create.call_count, 1)
        self.assertEqual(optimizer.last_bullet_bank_selections, {})


class DraftPipelineTest(unittest.TestCase):
    def setUp(self):
        self.runner = pipeline.ResumePipelineTest(methodName="test_all_modes_change_only_intended_sections_and_record_reports")
        self.runner.setUp()
        self.addCleanup(self.runner.doCleanups)

    def test_explicit_title_flags_override_modes_in_both_directions(self):
        for mode, flag in (("full", False), ("bullets", True), ("skills", True)):
            with self.subTest(mode=mode, flag=flag):
                self.runner.config["replace_job_title"] = flag
                result = self.runner.run_pipeline(mode)
                text = "\n".join(p.text for p in Document(Path(result).with_suffix(".docx")).paragraphs)
                self.assertEqual("Principal Engineer\t" in text, flag)
                report = self.runner.report(result)
                self.assertEqual(report["title_replacement_applied"], flag)
                if flag:
                    self.assertTrue(report["review_required"])

    def test_draft_artifact_reports_all_changes_and_proposed_claims(self):
        before = self.runner.source.read_bytes()
        self.runner.config.update(content_policy="review_draft", replace_job_title=True)
        result = self.runner.run_pipeline("bullets")
        report = self.runner.report(result)
        self.assertEqual(report["artifact_kind"], "manual_review_draft")
        self.assertIn("Kubernetes", report["keyword_provenance"]["introduced_terms"])
        self.assertIn("Kubernetes", report["skills_grounding"]["proposed_keywords"])
        self.assertTrue(report["review_required"])
        self.assertFalse(report["auto_submission_authorized"])
        self.assertTrue(report["changes"])
        self.assertEqual(self.runner.source.read_bytes(), before)

    def test_layout_retry_preserves_rewrite_and_title_before_skills_fallback(self):
        self.runner.config.update(layout_retry=True, replace_job_title=True)
        rejected = {"passed": False, "checks": {"page_count": False}, "warnings": []}
        accepted = {"passed": True, "checks": {"page_count": True}, "warnings": []}
        with patch.object(self.runner.optimizer, "verify_format", side_effect=[rejected, accepted]):
            result = self.runner.run_pipeline("bullets")
        report = self.runner.report(result)
        self.assertEqual(report["effective_mode"], "bullets-fit-retry")
        self.assertEqual(len(report["attempts"]), 2)
        text = "\n".join(p.text for p in Document(Path(result).with_suffix(".docx")).paragraphs)
        self.assertIn("Built production services with Python", text)
        self.assertIn("Principal Engineer\t", text)

    def test_failed_draft_layout_still_rejects_artifact(self):
        self.runner.config.update(content_policy="review_draft", layout_retry=True)
        with patch.object(self.runner.optimizer, "verify_format", return_value={"passed": False, "checks": {}, "warnings": []}):
            self.assertIsNone(self.runner.run_pipeline("full"))
        self.assertEqual(self.runner.report()["status"], "rejected")

    def test_bank_count_failure_retains_source_sections_and_requested_titles(self):
        self.runner.config.update(layout_retry=True, replace_job_title=True)
        candidate = {"id": "A1", "text": "Built bank API", "length": 14, "section": "A", "metrics": [], "lead_verb": "Built"}
        spec = {"company": "EXAMPLE COMPANY", "max_bullets": 7, "min_bullets": 7,
                "required_sections": ["A"], "bank": {"bullets": [candidate]}}
        rejected = {"passed": False, "checks": {"bullet_count_stable": False}, "warnings": []}
        accepted = {"passed": True, "checks": {"bullet_count_stable": True}, "warnings": []}
        with patch.object(pipeline.module, "load_project_bullet_banks", return_value=[spec]), \
                patch.object(self.runner.optimizer, "verify_format", side_effect=[rejected, rejected, accepted]):
            result = self.runner.run_pipeline("bullets")
        self.assertIsNone(result)
        preview = next((self.runner.root / "temp").glob("*_source_banks.match_report.json"))
        report = json.loads(preview.read_text())
        self.assertEqual(report["status"], "not_rewritten")
        self.assertEqual(report["effective_mode"], "bullets-source-banks-retry")
        self.assertEqual(len(report["attempts"]), 3)
        self.assertEqual(report["project_bullet_bank_selections"], {})
        output = Document(Path(report["output_path"]).with_suffix(".docx"))
        original = Document(self.runner.source)
        self.assertEqual([p.text for p in output.paragraphs if p.style.name == "List Paragraph"],
                         [p.text for p in original.paragraphs if p.style.name == "List Paragraph"])
        self.assertTrue(report["title_replacement_applied"])

    def test_compact_retry_caps_bank_to_source_count_and_lengths(self):
        self.runner.config["layout_retry"] = True
        spec = {"company": "EXAMPLE COMPANY", "max_bullets": 7, "min_bullets": 7, "bank": {"bullets": []}}
        captured = []
        original_save = self.runner.optimizer.save_optimized_docx
        def save(*args):
            captured.append(deepcopy(args[-1]))
            return original_save(*args)
        with patch.object(pipeline.module, "load_project_bullet_banks", return_value=[spec]), \
                patch.object(self.runner.optimizer, "save_optimized_docx", side_effect=save), \
                patch.object(self.runner.optimizer, "verify_format", side_effect=[{"passed": False}, {"passed": True}]):
            self.assertIsNotNone(self.runner.run_pipeline("bullets"))
        self.assertEqual(captured[0][0]["max_bullets"], 7)
        self.assertEqual(captured[1][0]["max_bullets"], 2)
        self.assertEqual(captured[1][0]["min_bullets"], 2)
        self.assertEqual(captured[1][0]["bank"]["max_characters"], len("Built production services using Python"))


if __name__ == "__main__":
    unittest.main()
