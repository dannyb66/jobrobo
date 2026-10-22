"""Relevant skill expansion and fixed-line geometry, without provider calls."""
import unittest
from unittest.mock import Mock,patch
import json
import socket
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from docx import Document
from docx.shared import Pt
from modules import resume_all_bullets as m
from modules.resume_skills_layout import candidate_pool,skill_key,pack_rows,settings,verify_fill
from modules.ai.resume_optimizer import run_resume_optimization,ResumeOptimizer


class ExpandedSkillsTest(unittest.TestCase):
    def test_jd_first_expansion_aliases_and_non_skill_exclusions(self):
        bank={'keywords':[{'keyword':'API Design','tier':'Core'},
                          {'keyword':'Machine Embroidery','tier':'Common'},
                          {'keyword':'Docker','tier':'Core'}]}
        pool=candidate_pool('Backend Go APIs with cloud infrastructure',
            ['Go','Principal Backend Engineer','10+ years of experience','United States','authorized to work in the US'],
            ['API Performance Optimization','Golang','Profiling'],
            ['Python','Golang','Docker'],bank,'Principal Backend Engineer')
        terms=[v['term'] for v in pool]
        self.assertEqual(terms[:3],['Go','API Performance Optimization','Profiling'])
        self.assertEqual(sum(skill_key(v)=='go' for v in terms),1)
        self.assertIn('Python',terms)
        self.assertIn('API Design',terms)
        self.assertNotIn('Machine Embroidery',terms)
        self.assertFalse(any('authorized' in v or 'years' in v or 'Engineer' in v for v in terms))

    def test_packing_fills_exact_count_with_whole_unique_ranked_terms(self):
        candidates=[dict(term=f'skill-{i:02d}'+('x'*(i%8)),origin='jd_direct',priority=0) for i in range(80)]
        rows=pack_rows(candidates,[100]*4,len,target=.95,minimum=.9)
        self.assertEqual(len(rows),4)
        self.assertEqual(rows[0][0],candidates[0])
        selected=[v['term'] for row in rows for v in row]
        self.assertEqual(len(selected),len(set(selected)))
        self.assertTrue(set(selected)<={v['term'] for v in candidates})
        for i,row in enumerate(rows):
            occupied=len(', '.join(v['term'] for v in row)+(',' if i<3 else ''))
            self.assertGreaterEqual(occupied,90)
            self.assertLessEqual(occupied,100)

    def test_insufficient_content_does_not_silently_pass_sparse_rows(self):
        with self.assertRaisesRegex(ValueError,'cannot fill line'):
            pack_rows([dict(term='Go',origin='jd_direct',priority=0)],[100]*2,len)
        layout={'source_line_count':2,'minimum_fill_ratio':.92,'lines':[{'usable_width':100}]*2}
        density={'output':[{'kind':'skills','lines':[{'width':70},{'width':96}]}]}
        self.assertFalse(verify_fill(density,layout)['passed'])
        density['output'][0]['lines'][0]['width']=95
        self.assertTrue(verify_fill(density,layout)['passed'])
        density['output'][0]['lines'].pop()
        self.assertFalse(verify_fill(density,layout)['passed'])

    def test_settings_validate_explicit_expansion_and_fill_targets(self):
        for config in ({'skills_expand_to_fill':'true'},{'skills_min_fill_ratio':1.1},
                       {'skills_min_fill_ratio':.98,'skills_target_fill_ratio':.95},{'max_skills':0}):
            with self.assertRaises(ValueError):settings(config)
        self.assertFalse(settings({})['enabled'])
        self.assertTrue(settings({'skills_expand_to_fill':True})['enabled'])

    def test_shared_provider_prompt_separates_expanded_skills_from_jd_scoring(self):
        payload={'job_description':'Improve backend API latency','expand_skills':True}
        schema=m.provider_response_format('keywords',payload)['json_schema']['schema']
        self.assertEqual(set(schema['required']),{'keywords','skill_candidates'})
        self.assertIn('Relax exact phrase matching',m.keyword_system(payload))
        self.assertIn('for scoring',m.keyword_system(payload))
        optimizer=Mock()
        def invoke(*args,provider,**kwargs):provider('keywords',payload,5)
        with patch.object(m,'run_all_bullets',side_effect=invoke):
            run_resume_optimization('JD','source.docx','job',rewrite_all_bullets=True,
                optimizer_instance=optimizer,runtime_config_override={'optimization_mode':'full'})
        call=optimizer.client.with_options.return_value.chat.completions.create.call_args.kwargs
        self.assertEqual(call['messages'][0]['content'],m.keyword_system(payload))
        self.assertEqual(call['response_format'],m.provider_response_format('keywords',payload))

    def test_real_rendered_expansion_is_eligible_for_existing_verified_uploader(self):
        from modules.resume_upload import upload_verified_resume
        with TemporaryDirectory() as tmp:
            root=Path(tmp);source=root/'original.docx';doc=Document()
            doc.styles['Normal'].font.name='Arial';doc.styles['Normal'].font.size=Pt(10)
            doc.add_heading('SKILLS',1)
            doc.add_paragraph('Python, TypeScript, AWS, Docker, Kubernetes, APIs, Go, Terraform, Redis, Kafka\n'
                'PostgreSQL, SQL, NoSQL, Distributed Systems, Microservices, Observability, Testing')
            doc.add_heading('PROFESSIONAL EXPERIENCE',1);doc.add_paragraph('Example company')
            doc.add_paragraph('Engineer\tJanuary 2020 - Present')
            doc.add_paragraph('Built Python APIs for 100 clients',style='List Bullet')
            doc.add_heading('ACADEMIC DISTINCTIONS',1)
            doc.add_paragraph('Earned a university award with a 3.9 GPA',style='List Bullet')
            doc.save(source)
            optimizer=ResumeOptimizer.__new__(ResumeOptimizer)
            self.assertTrue(optimizer.convert_docx_to_pdf(str(source),str(source.with_suffix('.pdf'))))
            jd='Build Python backend APIs, distributed systems and reliable cloud infrastructure.'
            config=dict(m.DEFAULTS,rewrite_all_bullets=True,source_checks_enabled=False,
                content_policy='synthetic_jd_test',skills_expand_to_fill=True,optimization_mode='full',expected_page_count=1,
                rewrite_requests_per_minute=1000000,rewrite_tokens_per_minute=1000000000)
            def provider(stage,payload,timeout):
                if stage=='keywords':
                    self.assertTrue(payload['expand_skills'])
                    data={'keywords':['Python','APIs','Distributed Systems'],
                          'skill_candidates':['API Performance Optimization','Cloud Infrastructure Automation',
                            'Service Reliability Engineering','Backend Architecture','Production Observability']}
                else:
                    data={'experience_id':payload['experience_id'],'results':[
                        dict(bullet_id=b['bullet_id'],source_text_hash=b['source_text_hash'],status='rewritten',
                             rewritten_text='Built reliable Python backend APIs for 100 clients.',matched_jd_spans=[],
                             preserved_fact_anchors=[],change_reason='Fixture rewrite',blocked_reason=None)
                        for b in payload['requested_bullets']]}
                return SimpleNamespace(id='fixture-provider',usage=SimpleNamespace(model_dump=lambda:
                    {'prompt_tokens':0,'completion_tokens':0}),choices=[SimpleNamespace(message=
                    SimpleNamespace(content=json.dumps(data)))])
            with patch.object(socket.socket,'connect',side_effect=AssertionError('network forbidden')):
                result=m.run_all_bullets(jd,source,'test-job','Staff Engineer',config,
                    provider=provider,execution_kind='provider',output_root=root/'review')
            report=json.loads((root/'review/match-report.json').read_text())
            self.assertIsNotNone(result,str(report.get('errors'))+str(report.get('skills_layout')))
            self.assertTrue(report['format_verification']['checks']['skills_full_width'])
            self.assertEqual(report['skills_layout']['rendered']['line_count'],2)
            self.assertTrue(all(v>=.92 for v in report['skills_layout']['rendered']['fill_ratios']))
            self.assertEqual(json.loads((root/'review/scoring-manifest.json').read_text())['keywords'],
                             ['Python','APIs','Distributed Systems'])
            modal=Mock()
            self.assertTrue(upload_verified_resume(modal,str(source),jd,'test-job',
                optimize=Mock(return_value=result),runtime_config_override=config)[0])
            modal.find_element.return_value.send_keys.assert_called_once_with(str(Path(result).resolve()))


if __name__=='__main__':unittest.main()
