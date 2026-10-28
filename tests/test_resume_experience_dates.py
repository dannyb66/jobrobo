import tempfile
import unittest
from pathlib import Path

from docx import Document

from modules.resume_experience_dates import (
    ExperienceRange, best_experience_range, load_experience_ranges,
)


class ResumeExperienceDateTests(unittest.TestCase):
    def test_product_owner_maps_to_resume_product_management_role_dates(self):
        records = [
            ExperienceRange('Senior Software Engineer (Founding Engineer)', 'January', '2024', 'June', '2025'),
            ExperienceRange('Product Manager (AI & Strategy)', 'January', '2020', 'January', '2022'),
        ]
        match = best_experience_range('Product Owner, San Francisco, customer workflow and product strategy', records)
        self.assertEqual(match.title, 'Product Manager (AI & Strategy)')
        self.assertEqual((match.end_month, match.end_year), ('January', '2022'))

    def test_unmatched_role_uses_latest_completed_resume_role(self):
        records = [
            ExperienceRange('Product Manager', 'January', '2020', 'January', '2022'),
            ExperienceRange('Senior Software Engineer', 'January', '2024', 'June', '2025'),
        ]
        match = best_experience_range('Completely unrelated entry', records)
        self.assertEqual((match.end_month, match.end_year), ('June', '2025'))

    def test_extracts_dated_roles_and_ignores_education(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'resume.docx'
            doc = Document()
            doc.add_paragraph('University of Example August 2012 - May 2016')
            doc.add_paragraph('Product Manager (AI & Strategy) January 2020 - January 2022')
            doc.add_paragraph('Current Engineer June 2025 - Present')
            doc.save(path)
            records = load_experience_ranges(path)
        self.assertEqual([item.title for item in records], ['Product Manager (AI & Strategy)'])


if __name__ == '__main__':
    unittest.main()
