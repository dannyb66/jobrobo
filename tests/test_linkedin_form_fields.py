import tempfile
import unittest
import os
from pathlib import Path
from urllib.parse import quote

from docx import Document
from selenium.webdriver.common.by import By

from modules.linkedin_easy_apply import MissingApplicationAnswer
from modules.linkedin_form_fields import education_checkbox_answers, field_limit, fill_text, invalid_text, short_numeric_rejection, inapplicable_followup, bounded_text, checkbox_prompt, demographic_request
from test_linkedin_modern_search import DOMBrowser


class FieldConstraintsTests(unittest.TestCase):
    def test_application_location_prefers_posting_city_over_first_option(self):
        from modules.linkedin_form_fields import application_location_answer
        choices = [{'text': text} for text in ['Los Angeles, CA','Oakland, CA','New York, NY','Toronto']]
        self.assertEqual(application_location_answer('Which location are you applying for?', choices,
                         'Oakland, CA', 'San Francisco Bay Area'), 'Oakland, CA')

    def test_application_location_rejects_outside_region_or_ambiguous_choices(self):
        from modules.linkedin_form_fields import application_location_answer
        for options in [['Los Angeles, CA','Toronto'], ['Oakland, CA','San Francisco, CA']]:
            with self.assertRaisesRegex(MissingApplicationAnswer, 'No unambiguous application location'):
                application_location_answer('Which location are you applying for?',
                    [{'text': text} for text in options], 'United States', 'San Francisco Bay Area')

    def test_same_city_in_another_state_is_not_a_posting_match(self):
        from modules.linkedin_form_fields import application_location_answer
        with self.assertRaises(MissingApplicationAnswer):
            application_location_answer('Which location are you applying for?',
                [{'text':'Oakland, MI'}], 'Oakland, CA', 'San Francisco Bay Area')

    def test_application_location_does_not_answer_personal_address_or_boolean_questions(self):
        from modules.linkedin_form_fields import application_location_answer
        self.assertIsNone(application_location_answer('Current work location?', [{'text':'Oakland'}], 'Oakland'))
        self.assertIsNone(application_location_answer('Are you willing to work at this location?',
                         [{'text':'Yes'},{'text':'No'}], 'Oakland'))
    def test_availability_uses_explicit_immediate_and_calendar_today(self):
        from datetime import date
        from modules.linkedin_form_fields import availability_answer
        settings = {'available_start_date': 'Immediate', 'notice_period': '0'}
        self.assertEqual(availability_answer('What is your potential start date?', settings), 'Immediate')
        self.assertEqual(availability_answer('When could you start?', settings, 'date', date(2026,10,27)), '2026-10-27')
        self.assertEqual(availability_answer('Availability', {'available_start_date':'2026-11-01'}), '2026-11-01')
        self.assertIsNone(availability_answer('Employment start date', settings))

    def test_calendar_date_is_set_in_iso_format_without_locale_keyboard_parsing(self):
        browser, control = self.field('<input type="date">')
        def execute(script, *args):
            if 'jobrobo-verified-field-replacement' in script:
                args[0].node.set('value', args[1])
        browser.execute_script = execute
        fill_text(browser, control, '2026-10-27', 200)
        self.assertEqual(control.get_attribute('value'), '2026-10-27')

    def test_experience_bucket_uses_configured_years_without_affecting_other_choices(self):
        from modules.linkedin_form_fields import experience_range_answer
        choices = [{'text': text} for text in ['Select an option', '<1 year', '1-2 years', '3-6 years', '7-10 years', '10+ years']]
        self.assertEqual(experience_range_answer('How long have you worked in an industry role relevant for this position?', choices, '8'), '7-10 years')
        self.assertEqual(experience_range_answer('Years of experience?', choices, '0'), '<1 year')
        self.assertEqual(experience_range_answer('Years of experience?', choices, '12'), '10+ years')
        self.assertIsNone(experience_range_answer('Desired compensation?', choices, '8'))

    def test_experience_duration_dropdown_is_not_a_calendar_year(self):
        from modules.linkedin_form_fields import is_calendar_date_select
        browser = DOMBrowser('''<select required><option>Select an option</option><option>&lt;1 year</option>
            <option>1-2 years</option><option>3-6 years</option><option>10+ years</option></select>''')
        from test_linkedin_modern_search import DOMElement
        self.assertFalse(is_calendar_date_select(DOMElement(browser, browser.root)))

    def test_native_month_year_placeholders_and_explicit_labels_are_dates(self):
        from modules.linkedin_form_fields import is_calendar_date_select
        from test_linkedin_modern_search import DOMElement
        for markup in ['<select><option>Year</option><option>2025</option></select>',
                       '<select><option>Month</option><option>January</option></select>',
                       '<select aria-label="From Year"><option>Select an option</option><option>2025</option></select>']:
            browser = DOMBrowser(markup)
            self.assertTrue(is_calendar_date_select(DOMElement(browser, browser.root)))

    def test_validation_message_names_external_group_prompt_without_answers(self):
        from modules.linkedin_form_fields import validation_messages
        browser = DOMBrowser('''<dialog open><div componentkey="easyApplyFieldFocus_ea_validation_p2_0_1">
          <p>Would you require visa sponsorship?*</p><fieldset><input type="checkbox"><input type="checkbox"></fieldset>
          <p>This field is required</p></div></dialog>''')
        self.assertEqual(validation_messages(browser.find_element(By.CSS_SELECTOR, 'dialog')),
                         ['Would you require visa sponsorship?*: This field is required'])

    def test_employer_no_option_can_explicitly_mean_decline(self):
        label = 'Are you Hispanic or Latino? If you choose not to self-identify, please indicate "No".'
        self.assertEqual(demographic_request(label, [{'text': 'Yes'}, {'text': 'No'}], {'ethnicity': 'Decline'}), (True, 'No'))

    def test_no_optout_preserves_explicit_first_option_policy(self):
        self.assertEqual(demographic_request('Are you a veteran?', [{'text': 'Yes'}, {'text': 'No'}],
            {'veteran_status': 'Decline', 'unanswered_choice_policy': 'first_option'}), (False, None))
    def test_prose_shortening_uses_utf16_and_does_not_split_surrogates(self):
        value = bounded_text('Engineer \U0001f600 ' * 30, 127)
        self.assertLessEqual(len(value.encode('utf-16-le')) // 2, 127)
        self.assertFalse(value.endswith('\ud83d'))
    def field(self, markup):
        browser = DOMBrowser('<main>' + markup + '</main>')
        return browser, browser.find_element(By.CSS_SELECTOR, 'input,textarea')

    def test_counter_without_maxlength_limits_only_own_field(self):
        browser, control = self.field('<section><div><input><span>490/20 characters</span></div>'
                                      '<div><input><span>1/500</span></div></section>')
        self.assertEqual(field_limit(browser, control, 500), 20)
        second = browser.find_elements(By.CSS_SELECTOR, 'input')[1]
        self.assertEqual(field_limit(browser, second, 500), 500)

    def test_native_and_counter_use_stricter_limit(self):
        browser, control = self.field('<div><textarea maxlength="100"></textarea><p>0/20</p></div>')
        self.assertEqual(field_limit(browser, control, 1500), 20)

    def test_fallback_does_not_take_neighbour_counter(self):
        browser, control = self.field('<div><input></div><div><input><span>0/20</span></div>')
        self.assertEqual(field_limit(browser, control, 500), 500)

    def test_contact_page_progress_is_not_a_phone_character_limit(self):
        browser = DOMBrowser('<dialog open><div>1/5 pages</div><section><select><option>Email</option></select>'
            '<select><option>Phone country</option></select><div><input type="tel"></div></section></dialog>')
        field = browser.find_element(By.CSS_SELECTOR, 'input')
        self.assertEqual(field_limit(browser, field, 500), 500)

    def test_page_ratio_in_single_field_parent_is_ignored(self):
        browser, control = self.field('<div><input><p>1/5 pages</p></div>')
        self.assertEqual(field_limit(browser, control, 500), 500)

    def test_nonempty_invalid_previous_answer_is_repaired(self):
        browser, control = self.field('<div><input value="Overlong previous answer"><span>24/20</span></div>')
        self.assertTrue(invalid_text(control, field_limit(browser, control, 500)))
        fill_text(browser, control, 'Vector search', 20)
        self.assertEqual(control.get_attribute('value'), 'Vector search')
        self.assertFalse(invalid_text(control, 20))

    def test_deterministic_email_not_silently_truncated(self):
        browser, control = self.field('<input type="email">')
        with self.assertRaisesRegex(MissingApplicationAnswer, '20-character'):
            fill_text(browser, control, 'candidate@example.test', 20)
        self.assertIsNone(control.get_attribute('value'))

    def test_browser_format_validation_is_checked(self):
        browser, control = self.field('<input type="number">')
        browser.execute_script = lambda *args: {'valid': False, 'badInput': True}
        with self.assertRaisesRegex(MissingApplicationAnswer, 'format'):
            fill_text(browser, control, 'prose', 20)

    def test_utf16_browser_counter_is_respected(self):
        browser, control = self.field('<input>')
        with self.assertRaises(MissingApplicationAnswer):
            fill_text(browser, control, '\U0001f600' * 11, 20)

    def test_numeric_recovery_requires_short_field_and_its_local_rejection(self):
        browser, control = self.field('<input>')
        browser.execute_script = lambda *args: True
        self.assertTrue(short_numeric_rejection(browser, control, 20))
        self.assertFalse(short_numeric_rejection(browser, control, 500))
        browser.execute_script = lambda *args: False
        self.assertFalse(short_numeric_rejection(browser, control, 20))

    def test_numeric_zero_is_preserved(self):
        browser, control = self.field('<input type="number">')
        fill_text(browser, control, 0, 20)
        self.assertEqual(control.get_attribute('value'), '0')

    def test_dropped_keyboard_characters_replaced_and_verified(self):
        browser, control = self.field('<input>')
        original_type = browser.type_into
        browser.type_into = lambda node, text: original_type(node, text.replace('e', ''))
        def execute(script, *args):
            if 'jobrobo-verified-field-replacement' in script:
                args[0].node.set('value', args[1])
        browser.execute_script = execute
        fill_text(browser, control, 'Reliable service', 200)
        self.assertEqual(control.get_attribute('value'), 'Reliable service')

    def test_unsuccessful_replacement_stops_without_accepting_partial_value(self):
        browser, control = self.field('<input>')
        original_type = browser.type_into
        browser.type_into = lambda node, text: original_type(node, text.replace('e', ''))
        with self.assertRaisesRegex(MissingApplicationAnswer, 'after replacement'):
            fill_text(browser, control, 'Reliable service', 200)

    def test_optional_ai_failure_does_not_abort_questionnaire(self):
        from unittest.mock import Mock
        from test_linkedin_easy_apply import NavigationTests
        provider = Mock(side_effect=RuntimeError('Provider unavailable'))
        browser, questions = NavigationTests().answer_native_form(
            '<dialog open><div><label for="motivation">Motivation</label><input id="motivation" type="text"></div></dialog>',
            ai_descriptive_answers_enabled=True, _draft_answer=provider)
        provider.assert_called_once()
        self.assertEqual(len(questions), 1)

    def test_required_ai_failure_remains_a_retryable_blocker(self):
        from unittest.mock import Mock
        from test_linkedin_easy_apply import NavigationTests
        provider = Mock(side_effect=RuntimeError('Provider unavailable'))
        with self.assertRaisesRegex(RuntimeError, 'Provider unavailable'):
            NavigationTests().answer_native_form(
                '<dialog open><div><label for="motivation">Motivation</label><input id="motivation" type="text" required></div></dialog>',
                ai_descriptive_answers_enabled=True, _draft_answer=provider)


class RequiredEducationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / 'resume.docx'
        doc = Document()
        for text in ('EDUCATION', 'Bachelor, Computer Science', 'SKILLS', 'Electrical Engineering'):
            doc.add_paragraph(text)
        doc.save(self.path)

    def modal(self, prompt='Field of study *'):
        browser = DOMBrowser('<dialog open><fieldset><legend>' + prompt + '</legend>'
            '<div><label for="cs">Computer Science/Software Engineering</label><input id="cs" type="checkbox"></div>'
            '<div><label for="ee">Electrical Engineering</label><input id="ee" type="checkbox"></div>'
            '</fieldset></dialog>')
        return browser.find_element(By.CSS_SELECTOR, 'dialog')

    def test_required_group_matches_source_degree_only(self):
        answers = education_checkbox_answers(self.modal(), self.path, {})
        self.assertEqual(answers, {'computer science/software engineering': True})

    def test_optional_group_untouched(self):
        self.assertEqual(education_checkbox_answers(self.modal('Field of study'), self.path, {}), {})

    def test_configured_selection_takes_precedence(self):
        answers = education_checkbox_answers(self.modal(), self.path,
            {'computer science/software engineering': False, 'electrical engineering': True})
        self.assertEqual(answers, {'computer science/software engineering': False, 'electrical engineering': True})

    def test_configured_true_does_not_add_another_resume_choice(self):
        answers = education_checkbox_answers(self.modal(), self.path, {'electrical engineering': True})
        self.assertEqual(answers, {'electrical engineering': True})

    def test_unknown_degree_is_not_invented(self):
        with self.assertRaisesRegex(MissingApplicationAnswer, 'resume-backed'):
            education_checkbox_answers(self.modal(), self.path.with_name('absent.docx'), {})

    def test_required_agreement_group_not_answered_from_resume(self):
        self.assertEqual(education_checkbox_answers(self.modal('Required agreements *'), self.path, {}), {})


@unittest.skipUnless(os.environ.get('JOBROBO_BROWSER_REPLAY_TESTS') == '1', 'Opt-in local Chrome DOM replay')
class LocalBrowserReplayTests(unittest.TestCase):
    """No LinkedIn/network forms: verify real DOM geometry and custom validity."""
    @classmethod
    def setUpClass(cls):
        from selenium import webdriver
        options = webdriver.ChromeOptions()
        options.add_argument('--headless=new')
        options.add_argument('--window-size=1300,900')
        cls.driver = webdriver.Chrome(options=options)

    @classmethod
    def tearDownClass(cls):
        cls.driver.quit()

    def test_direct_job_page_chat_is_harmless_but_covering_dialog_blocks(self):
        from modules.linkedin_post_submit import _dialog_blocks_jobs
        for style, modal, expected in [
                ('right:0;bottom:0;width:280px;height:250px', '', False),
                ('inset:0', '', True),
                ('right:0;bottom:0;width:280px;height:250px', ' aria-modal="true"', True)]:
            with self.subTest(style=style, modal=modal):
                self.driver.get('data:text/html,' + quote(
                    '<main style="width:800px;height:700px">Job details</main>'
                    '<aside role="dialog"' + modal + ' style="position:fixed;background:white;' + style + '">Chat</aside>'))
                self.assertEqual(_dialog_blocks_jobs(self.driver, self.driver.find_element(By.CSS_SELECTOR, 'aside')), expected)

    def test_privacy_checkbox_required_heading_is_read_without_native_required(self):
        self.driver.get('data:text/html,' + quote('<dialog open><section>'
            '<p>You declare that you have read and agree to the privacy notice of Example.*</p>'
            '<div role="checkbox"><input id="consent" type="checkbox"><label for="consent"></label>'
            '<p>I consent</p></div></section></dialog>'))
        control = self.driver.find_element(By.ID, 'consent')
        self.assertFalse(control.get_property('required'))
        state = checkbox_prompt(self.driver, control)
        self.assertTrue(state['required'])
        self.assertIn('privacy notice of Example', state['text'])

    def test_education_native_selections_saved_before_editor_closes(self):
        from modules.linkedin_education import repair_education
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / 'source.docx'
            doc = Document()
            for text in ['EDUCATION', 'Example University August 2012 - May 2016',
                         "Bachelor's, Computer Science", 'SKILLS']:
                doc.add_paragraph(text)
            doc.save(source)
            self.driver.get('data:text/html,' + quote('<dialog open><h2>Apply to Example</h2>'
                '<input aria-label="School" value="Example University">'
                '<label for="degree">Education Degree*</label><select id="degree" required>'
                '<option value="">Select an option</option><option>Bachelor of Science (B.S)</option></select>'
                '<label for="study">Education Field of Study*</label><select id="study" required>'
                '<option value="">Select an option</option><option>Computer and Information Science</option></select>'
                '<button onclick="window.saved=[degree.value,study.value];this.parentElement.remove()">Save</button>'
                '<button>Next</button></dialog>'))
            self.assertTrue(repair_education(self.driver, source, {}, set()))
            self.assertEqual(self.driver.execute_script('return window.saved'),
                             ['Bachelor of Science (B.S)', 'Computer and Information Science'])

    def test_custom_numeric_error_is_local_and_counter_constrains_typing(self):
        self.driver.get('data:text/html,' + quote('<main><div><input id="skill"><p>Invalid input</p><span>0/20</span></div>'
            '<div><input id="other"><span>0/500</span></div></main>'))
        skill, other = [self.driver.find_element(By.ID, ident) for ident in ('skill', 'other')]
        self.assertEqual(field_limit(self.driver, skill, 500), 20)
        self.assertEqual(field_limit(self.driver, other, 500), 500)
        self.assertTrue(short_numeric_rejection(self.driver, skill, 20))
        self.assertFalse(short_numeric_rejection(self.driver, other, 20))
        fill_text(self.driver, skill, '2', 20)
        self.assertEqual(skill.get_attribute('value'), '2')

    def test_contact_page_counter_does_not_constrain_phone(self):
        self.driver.get('data:text/html,' + quote('<dialog open><div>1/5 pages</div><section>'
            '<select><option>Email</option></select><select><option>Country</option></select>'
            '<div><input type="tel"></div></section></dialog>'))
        control = self.driver.find_element(By.CSS_SELECTOR, 'input')
        self.assertEqual(field_limit(self.driver, control, 500), 500)
        fill_text(self.driver, control, '5551234567', 500)
        self.assertEqual(control.get_attribute('value'), '5551234567')

    def test_chat_does_not_prevent_discard_of_separate_application_modal(self):
        from modules.linkedin_post_submit import _dialog_blocks_jobs
        self.driver.get('data:text/html,' + quote(
            '<main style="width:800px;height:700px">Job details</main>'
            '<div id="app" role="dialog" aria-modal="true" style="position:fixed;inset:0;background:white">Application</div>'
            '<aside role="dialog" style="position:fixed;right:0;bottom:0;width:280px;height:250px;background:white">Chat</aside>'))
        self.assertTrue(_dialog_blocks_jobs(self.driver, self.driver.find_element(By.ID, 'app')))
        self.assertFalse(_dialog_blocks_jobs(self.driver, self.driver.find_element(By.CSS_SELECTOR, 'aside')))

    def test_conditional_followup_skipped_only_after_explicit_no(self):
        for selected, required, expected in [('No', '', True), ('Yes', '', False), ('No', ' required', False)]:
            with self.subTest(selected=selected, required=required):
                radios = ''.join('<input type="radio" name="g" id="' + answer + '"' +
                    (' checked' if answer == selected else '') + '><label for="' + answer + '">' + answer + '</label>'
                    for answer in ('Yes', 'No'))
                self.driver.get('data:text/html,' + quote('<dialog open><fieldset>' + radios +
                    '</fieldset><input type="text" id="detail"' + required + '></dialog>'))
                field = self.driver.find_element(By.ID, 'detail')
                self.assertEqual(inapplicable_followup(self.driver, field, 'If yes, provide details'), expected)
                self.assertFalse(inapplicable_followup(self.driver, field, 'Describe a project'))


if __name__ == '__main__':
    unittest.main()
