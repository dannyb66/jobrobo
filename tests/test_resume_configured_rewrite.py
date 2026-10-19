"""Configured production policy and real artifact verification; provider is mocked.

All fixtures stay temporary. No employer actions, credentials or provider calls.
"""
import json
import socket
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import Mock,patch

from docx import Document
from docx.shared import Pt
from modules import resume_all_bullets as m, resume_jd_experiment as jd
from modules.ai.resume_optimizer import ResumeOptimizer, run_resume_optimization


class ConfiguredRewriteTest(unittest.TestCase):
    def test_boolean_validation_and_explicit_override_precedence(self):
        for value in ('false',0,None):
            with self.assertRaises(ValueError):m.validate_options({'source_checks_enabled':value},'full')
        with self.assertRaises(ValueError):m.validate_options({'source_checks_enabled':False},'full')
        optimizer=Mock()
        with patch.object(m,'run_all_bullets',return_value=None) as run:
            run_resume_optimization('JD','source.docx','job',rewrite_all_bullets=True,
                source_checks_enabled=False,optimizer_instance=optimizer,
                runtime_config_override={'source_checks_enabled':True,'optimization_mode':'full'})
        self.assertFalse(run.call_args.args[4]['source_checks_enabled'])

    def test_production_payload_uses_optional_original_reference_and_metrics(self):
        from test_resume_all_bullets import inv
        inventory=inv(1,1);group=inventory['groups'][0]
        payload=m.compact_payload(inventory,'Build Rust agents','Engineer',jd.PRODUCTION_POLICY,group,group['bullets'])
        self.assertEqual(payload['content_policy'],'jd_rewrite')
        self.assertFalse(payload['synthetic_test'])
        self.assertEqual(payload['requested_bullets'][0]['original_bullet_reference'],group['bullets'][0]['source_text'])
        self.assertIn('original_metric_hints',payload['requested_bullets'][0])
        self.assertNotIn('protected_source_terms',payload['requested_bullets'][0])
        self.assertEqual(m.generation_system(jd.PRODUCTION_POLICY),m.SYSTEM+jd.PRODUCTION_OVERRIDE)
        self.assertNotIn('synthetic',payload['requested_bullets'][0]['rewrite_focus'])
        self.assertNotIn('fictional',payload['acceptance_criteria']['facts'])

    def test_metric_preference_is_soft_and_selects_quantified_candidate(self):
        from test_resume_all_bullets import inv,response
        inventory=inv(1,1);group=inventory['groups'][0];source=group['bullets'][0]
        source.update(source_text='Built APIs for 100+ clients',source_text_hash='hash')
        raw=response({'experience_id':group['experience_id'],'requested_bullets':[source]})
        raw['results'][0].update(rewritten_text='Built Rust agents.',alternative_texts=['Built Rust agents for 100+ clients.'])
        out=m.validate_response(raw,group,[source],content_policy=jd.PRODUCTION_POLICY,ranking_keywords=['Rust'])
        self.assertIn('100+',out[source['bullet_id']]['rewritten_text'])

    def test_complete_production_artifact_binds_policy_job_and_configuration(self):
        with TemporaryDirectory() as tmp:
            root=Path(tmp);source=root/'original.docx';doc=Document()
            doc.styles['Normal'].font.name='Arial';doc.styles['Normal'].font.size=Pt(10)
            doc.add_heading('SKILLS',1);doc.add_paragraph('Python, AWS, React')
            doc.add_heading('PROFESSIONAL EXPERIENCE',1);doc.add_paragraph('Example company')
            doc.add_paragraph('Engineer\tJanuary 2020 - Present')
            doc.add_paragraph('Built Python services for 100 clients',style='List Bullet')
            doc.add_paragraph('Reduced API latency by 50%',style='List Bullet')
            doc.add_heading('ACADEMIC DISTINCTIONS',1)
            doc.add_paragraph('Earned a university award with a 3.9 GPA',style='List Bullet')
            doc.save(source);before=source.read_bytes()
            optimizer=ResumeOptimizer.__new__(ResumeOptimizer)
            self.assertTrue(optimizer.convert_docx_to_pdf(str(source),str(source.with_suffix('.pdf'))))
            jd_text='Required: Build Rust agents with AG-UI and MCP.'
            config=dict(m.DEFAULTS,rewrite_all_bullets=True,source_checks_enabled=False,optimization_mode='full',
                content_policy='grounded',expected_page_count=1,rewrite_requests_per_minute=1000000,
                rewrite_tokens_per_minute=1000000000)
            def provider(stage,payload,timeout):
                if stage=='keywords':data={'keywords':['Rust','AG-UI','MCP']}
                else:data={'experience_id':payload['experience_id'],'results':[
                    {'bullet_id':b['bullet_id'],'source_text_hash':b['source_text_hash'],'status':'rewritten',
                     'rewritten_text':'Built Rust agents for 100 clients.' if b['ordinal']==0 else 'Shipped AG-UI tools with MCP.',
                     'matched_jd_spans':[],'preserved_fact_anchors':[],
                     'change_reason':'Mock provider JD rewrite','blocked_reason':None} for b in payload['requested_bullets']]}
                return SimpleNamespace(id='unit-test-provider-response',
                    usage=SimpleNamespace(model_dump=lambda:{'prompt_tokens':0,'completion_tokens':0}),
                    choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps(data)))])
            progress_events=[]
            with patch.object(socket.socket,'connect',side_effect=AssertionError('network forbidden')):
                result=m.run_all_bullets(jd_text,source,'job-1','Staff Engineer',config,
                    provider=provider,execution_kind='provider',output_root=root/'production',progress=progress_events.append)
            stages=[event['event'] for event in progress_events]
            self.assertLess(stages.index('inventory_started'),stages.index('inventory_ready'))
            self.assertLess(stages.index('keywords_ready'),stages.index('bullets_started'))
            self.assertLess(stages.index('rendering'),stages.index('verifying'))
            self.assertLess(stages.index('verifying'),stages.index('saving'))
            self.assertEqual(stages[-1],'finished')
            self.assertEqual(progress_events[-1]['accepted'],2)
            self.assertTrue(progress_events[-1]['eligible'])
            self.assertTrue(Path(result).name.startswith('resume-job-1-'))
            self.assertTrue(Path(result).with_suffix('.docx').is_file())
            self.assertEqual(Path(result).read_bytes(),(root/'production/optimized.pdf').read_bytes())
            expected=dict(expected_jd=jd_text,expected_job_id='job-1',expected_source_docx=source,
                expected_source_checks_enabled=False,expected_config=config)
            self.assertTrue(m.verify_strict_attachment(result,**expected))
            self.assertFalse(m.verify_strict_attachment(result,**{**expected,'expected_source_checks_enabled':True}))
            self.assertFalse(m.verify_strict_attachment(result,**{**expected,'expected_job_id':'different-job'}))
            self.assertFalse(m.verify_strict_attachment(result,**{**expected,'expected_jd':'Different JD'}))
            self.assertFalse(m.verify_strict_attachment(result,**{**expected,'expected_config':{**config,'rewrite_model':'changed'}}))
            self.assertFalse(m.verify_strict_attachment(result),'Disabled checks must be explicitly expected by the caller')
            report=json.loads((root/'production/match-report.json').read_text())
            self.assertEqual(report['effective_config']['content_policy'],'jd_rewrite')
            self.assertFalse(report['source_content_checks_enabled'])
            self.assertEqual(report['coverage']['accepted'],2)
            self.assertEqual(report['preserved_sections']['Academic Distinctions']['bullet_count'],1)
            self.assertTrue(report['format_verification']['checks']['academic_distinctions_preserved'])
            self.assertEqual(source.read_bytes(),before)
            # A synthetic override under another runtime policy remains a preview.
            with patch.object(socket.socket,'connect',side_effect=AssertionError('network forbidden')):
                preview=m.run_all_bullets(jd_text,source,'job-1','Staff Engineer',config,
                    policy=jd.POLICY,provider=provider,execution_kind='provider',output_root=root/'synthetic')
            self.assertIsNone(preview)
            self.assertFalse(m.verify_strict_attachment(root/'synthetic/optimized.pdf',**expected))
            # Explicit runtime selection joins the existing verified upload path.
            synthetic_config={**config,'content_policy':jd.POLICY}
            with patch.object(socket.socket,'connect',side_effect=AssertionError('network forbidden')):
                synthetic=m.run_all_bullets(jd_text,source,'job-1','Staff Engineer',synthetic_config,
                    provider=provider,execution_kind='provider',output_root=root/'synthetic-upload')
            self.assertIsNotNone(synthetic)
            synthetic_expected={**expected,'expected_config':synthetic_config}
            self.assertTrue(m.verify_strict_attachment(synthetic,**synthetic_expected))
            self.assertFalse(m.verify_strict_attachment(synthetic,**expected))
            self.assertFalse(m.verify_strict_attachment(synthetic,**{**synthetic_expected,'expected_config':None}))
            self.assertFalse(m.verify_strict_attachment(synthetic,**{**synthetic_expected,'expected_source_checks_enabled':True}))
            self.assertFalse(m.verify_strict_attachment(synthetic,**{**synthetic_expected,'expected_job_id':'wrong-job'}))
            self.assertFalse(m.verify_strict_attachment(synthetic,**{**synthetic_expected,'expected_jd':'wrong-JD'}))
            synthetic_report=json.loads((root/'synthetic-upload/match-report.json').read_text())
            self.assertTrue(synthetic_report['synthetic_test'])
            self.assertTrue(synthetic_report['strict_attachment_eligible'])
            self.assertFalse(synthetic_report['source_content_checks_enabled'])
            from modules.resume_upload import upload_verified_resume
            modal=Mock()
            self.assertTrue(upload_verified_resume(modal,str(source),jd_text,'job-1',
                optimize=Mock(return_value=synthetic),runtime_config_override=synthetic_config)[0])
            modal.find_element.return_value.send_keys.assert_called_once_with(str(Path(synthetic).resolve()))
            # Mutating the attached artifact is still detected independently.
            Path(synthetic).write_bytes(Path(synthetic).read_bytes()+b'changed')
            self.assertFalse(m.verify_strict_attachment(synthetic,**synthetic_expected))


if __name__=='__main__':unittest.main()
