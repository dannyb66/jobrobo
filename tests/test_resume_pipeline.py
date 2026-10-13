"""Pipeline tests use synthetic documents and mocked APIs/conversion only."""
import json
import os
import shutil
import subprocess
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from docx import Document
from docx.shared import Pt
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter

import modules.ai.resume_optimizer as module


def synthetic_source(path):
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading("SKILLS", level=1)
    paragraph = doc.add_paragraph("Python, Legacy Tool, AWS")
    paragraph.runs[0].font.size = Pt(10)
    doc.add_heading("PROFESSIONAL EXPERIENCE", level=1)
    doc.add_paragraph("EXAMPLE COMPANY\tRemote")
    doc.add_paragraph("Software Engineer\tJanuary 2020 - Present")
    for text in ["Built production services using Python", "Improved API reliability on AWS"]:
        paragraph = doc.add_paragraph(text, style="List Paragraph")
        paragraph.runs[0].font.size = Pt(10)
    doc.save(path)


def pdf_from_lines(path, lines):
    writer = canvas.Canvas(str(path), pagesize=letter)
    writer.setFont("Helvetica", 10)
    for index, line in enumerate(lines):
        writer.drawString(72, 730 - index * 20, line)
    writer.showPage()
    writer.save()


def mocked_convert(docx_path, pdf_path):
    lines = [paragraph.text.replace("\t", " ") for paragraph in Document(docx_path).paragraphs if paragraph.text]
    pdf_from_lines(pdf_path, lines)
    return True


def api_response(content):
    return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=content))])


class ResumePipelineTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / "default" / "resume.docx"
        self.source.parent.mkdir()
        synthetic_source(self.source)
        mocked_convert(str(self.source), str(self.source.with_suffix(".pdf")))
        self.config = {"optimization_mode": "full", "expected_page_count": 1, "project_bullet_banks": [], "max_skills": 50}
        self.optimizer = module.ResumeOptimizer.__new__(module.ResumeOptimizer)
        self.optimizer.client = Mock()
        self.optimizer.client.chat.completions.create.return_value = api_response('{"keywords": ["Python", "AWS", "Kubernetes"]}')
        self.addCleanup(patch.stopall)
        patch.dict(os.environ, {"OPENAI_API_KEY": "test-key"}).start()
        patch.object(module, "load_runtime_config", return_value=self.config).start()
        patch.object(module, "load_dotenv").start()
        patch.object(module, "ResumeOptimizer", return_value=self.optimizer).start()
        patch.object(self.optimizer, "convert_docx_to_pdf", side_effect=mocked_convert).start()
        self.rewrite = patch.object(self.optimizer, "_call_openai_to_rewrite_bullets", side_effect=self.rewrite_entries).start()

    def rewrite_entries(self, entries, *args):
        return [dict(entry, bullets=["Built Kubernetes service" for _ in entry["bullets"]]) for entry in entries]

    def run_pipeline(self, mode="full", source=None, job_id="test"):
        return module.run_resume_optimization("Python AWS Kubernetes", str(source or self.source), job_id,
                                              "Principal Engineer", "Test", "Candidate", mode)

    def report(self, output=None):
        path = Path(output).with_suffix(".match_report.json") if output else self.root / "temp" / "test_candidate_principal_engineer_test.match_report.json"
        return json.loads(path.read_text())

    def test_all_modes_change_only_intended_sections_and_record_reports(self):
        for mode in ("skills", "bullets", "full"):
            with self.subTest(mode=mode):
                self.rewrite.reset_mock()
                result = self.run_pipeline(mode)
                self.assertIsNotNone(result)
                doc = Document(Path(result).with_suffix(".docx"))
                text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
                self.assertIn("Kubernetes", text)
                self.assertEqual("Built Kubernetes service" in text, mode != "skills")
                self.assertEqual("Principal Engineer\t" in text, mode == "full")
                self.assertEqual(self.rewrite.call_count, int(mode != "skills"))
                report = self.report(result)
                self.assertEqual(report["status"], "accepted")
                self.assertEqual(report["effective_mode"], mode)
                self.assertTrue(report["format_verification"]["passed"])
                self.assertGreaterEqual(report["post_optimization"]["score"], report["pre_optimization"]["score"])

    def test_invalid_mode_is_normalized_and_original_value_retained(self):
        result = self.run_pipeline("invalid")
        report = self.report(result)
        self.assertEqual(report["requested_mode"], "invalid")
        self.assertEqual(report["normalized_mode"], "full")
        self.assertEqual(report["effective_mode"], "full")

    def test_skills_fallback_retains_original_rejection_and_clears_bank_selection(self):
        def save(doc, skills, path, title, description, mode, *args):
            self.optimizer._update_skills_section(doc, *self.optimizer._find_skills_section(doc), skills)
            if mode != "skills":
                self.optimizer.last_bullet_bank_selection = [{"id": "A1"}]
                self.optimizer.last_bullet_bank_selections = {"Example": [{"id": "A1"}]}
            doc.save(path)
        patch.object(self.optimizer, "save_optimized_docx", side_effect=save).start()
        rejected = {"passed": False, "checks": {"page_count": False}, "warnings": []}
        accepted = {"passed": True, "checks": {"page_count": True}, "warnings": []}
        patch.object(self.optimizer, "verify_format", side_effect=[rejected, accepted]).start()
        result = self.run_pipeline()
        self.assertIn("_layout_safe.pdf", result)
        report = self.report(result)
        self.assertEqual(report["status"], "accepted")
        self.assertEqual(report["effective_mode"], "skills-fallback")
        self.assertEqual(report["selected_bullet_bank_ids"], [])
        self.assertEqual(report["project_bullet_bank_selections"], {})
        self.assertFalse(report["attempts"][0]["format_verification"]["passed"])
        self.assertEqual(report["attempts"][0]["project_bullet_bank_selections"], {"Example": ["A1"]})
        self.assertEqual(report["attempts"][1]["selected_bullet_bank_ids"], [])

    def test_both_attempts_rejected_returns_no_resume_and_saves_report(self):
        patch.object(self.optimizer, "verify_format", return_value={"passed": False, "checks": {"page_count": False}, "warnings": []}).start()
        self.assertIsNone(self.run_pipeline())
        report = self.report()
        self.assertEqual(report["status"], "rejected")
        self.assertEqual(len(report["attempts"]), 2)
        self.assertTrue(all(not attempt["format_verification"]["passed"] for attempt in report["attempts"]))

    def test_skills_rejection_does_not_retry(self):
        verifier = patch.object(self.optimizer, "verify_format", return_value={"passed": False, "checks": {}, "warnings": []}).start()
        self.assertIsNone(self.run_pipeline("skills"))
        self.assertEqual(verifier.call_count, 1)
        self.assertEqual(len(self.report()["attempts"]), 1)

    def test_conversion_failure_saves_report_and_cannot_accept_stale_output(self):
        destination = self.root / "temp" / "test_candidate_principal_engineer_test.pdf"
        destination.parent.mkdir()
        shutil.copyfile(self.source.with_suffix(".pdf"), destination)
        self.optimizer.convert_docx_to_pdf.side_effect = None
        self.optimizer.convert_docx_to_pdf.return_value = False
        self.optimizer.last_conversion_error = "conversion failed"
        verifier = patch.object(self.optimizer, "verify_format").start()
        self.assertIsNone(self.run_pipeline())
        self.assertFalse(destination.exists())
        verifier.assert_not_called()
        report = self.report()
        self.assertEqual(report["status"], "failed")
        self.assertIn("conversion failed", report["error"])
        self.assertFalse(report["attempts"][0]["conversion_succeeded"])

    def test_save_failure_cannot_accept_stale_docx_or_pdf(self):
        destination = self.root / "temp" / "test_candidate_principal_engineer_test.pdf"
        destination.parent.mkdir()
        shutil.copyfile(self.source, destination.with_suffix(".docx"))
        shutil.copyfile(self.source.with_suffix(".pdf"), destination)
        patch.object(self.optimizer, "save_optimized_docx", return_value=None).start()
        self.assertIsNone(self.run_pipeline())
        self.assertFalse(destination.exists())
        self.assertFalse(destination.with_suffix(".docx").exists())
        self.assertIn("DOCX was not created", self.report()["error"])

    def test_failed_fallback_conversion_preserves_initial_rejection(self):
        # Use a callable sequence so first conversion creates the requested file.
        calls = []
        def convert(docx_path, pdf_path):
            calls.append(pdf_path)
            return mocked_convert(docx_path, pdf_path) if len(calls) == 1 else False
        self.optimizer.convert_docx_to_pdf.side_effect = convert
        patch.object(self.optimizer, "verify_format", return_value={"passed": False, "checks": {"page_count": False}, "warnings": []}).start()
        self.assertIsNone(self.run_pipeline())
        report = self.report()
        self.assertEqual(report["status"], "rejected")
        self.assertFalse(report["format_verification"]["passed"])
        self.assertIn("Skills-only retry failed", report["error"])
        self.assertEqual(len(report["attempts"]), 2)

    def test_pdf_input_is_effectively_skills_only(self):
        result = self.run_pipeline("full", source=self.source.with_suffix(".pdf"))
        self.assertIsNotNone(result)
        report = self.report(result)
        self.assertEqual(report["requested_mode"], "full")
        self.assertEqual(report["effective_mode"], "skills-pdf")
        self.rewrite.assert_not_called()

    def test_job_id_cannot_escape_output_directory(self):
        result = self.run_pipeline("skills", job_id="../../outside")
        self.assertEqual(Path(result).parent, self.root / "temp")
        self.assertNotIn("..", Path(result).name)

    def test_missing_key_saves_failure_report_without_api_calls(self):
        with patch.dict(os.environ, {}, clear=True):
            self.assertIsNone(self.run_pipeline())
        report = self.report()
        self.assertEqual(report["status"], "failed")
        self.assertIn("OPENAI_API_KEY", report["error"])
        self.optimizer.client.chat.completions.create.assert_not_called()

    def test_missing_source_saves_failure_report(self):
        self.source.unlink()
        self.assertIsNone(self.run_pipeline())
        self.assertIn("Resume file not found", self.report()["error"])
        self.assertEqual(self.report()["attempts"], [])

    def test_unsupported_source_saves_failure_report(self):
        source = self.source.with_suffix(".txt")
        source.write_text("unsupported")
        self.assertIsNone(self.run_pipeline(source=source))
        self.assertIn("Unsupported file format", self.report()["error"])

    def test_docx_save_exception_saves_failure_report(self):
        patch.object(self.optimizer, "save_optimized_docx", side_effect=OSError("disk full")).start()
        self.assertIsNone(self.run_pipeline())
        self.assertIn("disk full", self.report()["error"])
        self.assertIn("disk full", self.report()["attempts"][0]["error"])

    def test_converter_claiming_success_without_requested_pdf_is_rejected(self):
        self.optimizer.convert_docx_to_pdf.side_effect = None
        self.optimizer.convert_docx_to_pdf.return_value = True
        self.assertIsNone(self.run_pipeline())
        self.assertIn("did not create the requested PDF", self.report()["error"])

    def test_unrelated_converter_output_is_rejected_including_retry(self):
        def convert(docx_path, pdf_path):
            pdf_from_lines(pdf_path, ["Unrelated document"])
            return True
        self.optimizer.convert_docx_to_pdf.side_effect = convert
        self.assertIsNone(self.run_pipeline())
        report = self.report()
        self.assertEqual(report["status"], "rejected")
        self.assertEqual(len(report["attempts"]), 2)
        self.assertFalse(report["format_verification"]["checks"]["docx_pdf_content_consistent"])

    def test_keyword_api_failure_is_reported_without_crashing(self):
        self.optimizer.client.chat.completions.create.side_effect = RuntimeError("API unavailable")
        result = self.run_pipeline("skills")
        self.assertIsNotNone(result)
        self.assertEqual(self.report(result)["status"], "accepted")
        self.assertIn("API unavailable", self.report(result)["api_warnings"])


class LibreOfficeConverterTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / "source.docx"
        synthetic_source(self.source)
        self.destination = self.root / "different-name.pdf"
        self.optimizer = module.ResumeOptimizer.__new__(module.ResumeOptimizer)

    def test_discovers_configured_executable_and_moves_requested_basename(self):
        def command(args, **kwargs):
            self.assertEqual(args[0], "/custom/soffice")
            self.assertEqual(kwargs["timeout"], 60)
            output_dir = Path(args[args.index("--outdir") + 1])
            pdf_from_lines(output_dir / "source.pdf", ["Fresh generated PDF"])
            return SimpleNamespace(stdout="converted", stderr="")
        with patch.dict(os.environ, {"LIBREOFFICE_PATH": "/custom/soffice"}), patch.object(module.subprocess, "run", side_effect=command):
            self.assertTrue(self.optimizer.convert_docx_to_pdf(str(self.source), str(self.destination)))
        self.assertTrue(self.destination.exists())

    def test_existing_destination_is_not_accepted_when_no_fresh_output_generated(self):
        pdf_from_lines(self.destination, ["Stale PDF"])
        with patch.dict(os.environ, {"LIBREOFFICE_PATH": "/custom/soffice"}), patch.object(module.subprocess, "run", return_value=SimpleNamespace(stdout="", stderr="no output")):
            self.assertFalse(self.optimizer.convert_docx_to_pdf(str(self.source), str(self.destination)))
        self.assertIn("did not create", self.optimizer.last_conversion_error)

    def test_timeout_returns_failure(self):
        with patch.dict(os.environ, {"LIBREOFFICE_PATH": "/custom/soffice"}), patch.object(module.subprocess, "run", side_effect=subprocess.TimeoutExpired("soffice", 60)):
            self.assertFalse(self.optimizer.convert_docx_to_pdf(str(self.source), str(self.destination)))
        self.assertFalse(self.destination.exists())

    def test_invalid_generated_pdf_returns_failure(self):
        def command(args, **kwargs):
            output_dir = Path(args[args.index("--outdir") + 1])
            (output_dir / "source.pdf").write_text("not a PDF")
            return SimpleNamespace(stdout="", stderr="")
        with patch.dict(os.environ, {"LIBREOFFICE_PATH": "/custom/soffice"}), patch.object(module.subprocess, "run", side_effect=command):
            self.assertFalse(self.optimizer.convert_docx_to_pdf(str(self.source), str(self.destination)))
        self.assertFalse(self.destination.exists())


if __name__ == "__main__":
    unittest.main()
